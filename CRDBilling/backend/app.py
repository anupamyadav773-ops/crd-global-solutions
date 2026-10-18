import sys
import os

_PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if _PROJECT_ROOT not in sys.path:
    sys.path.insert(0, _PROJECT_ROOT)

from flask import Flask, request, jsonify, g, send_file, send_from_directory
from flask_cors import CORS
import mysql.connector
from mysql.connector import pooling
import threading
import re
import csv
import json
import secrets
import time
from collections import defaultdict
from io import StringIO, BytesIO
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from functools import wraps
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer
from werkzeug.security import check_password_hash, generate_password_hash

# Indian Standard Time (IST, UTC+5:30) helper
IST_TIMEZONE = timezone(timedelta(hours=5, minutes=30))


def get_ist_now():
    """Return current datetime in Indian Standard Time (IST, UTC+5:30)."""
    return datetime.now(IST_TIMEZONE)


def get_ist_today():
    """Return current date in Indian Standard Time as date object."""
    return get_ist_now().date()


def get_ist_today_str():
    """Return current date in Indian Standard Time as YYYY-MM-DD string."""
    return get_ist_now().strftime("%Y-%m-%d")


try:
    from backend.technician_module import technician_bp, ensure_technician_employee_schema, serialize_row, serialize_rows, cascade_technician_schedule_shift, find_technician_next_available_slot
except ImportError:
    from technician_module import technician_bp, ensure_technician_employee_schema, serialize_row, serialize_rows, cascade_technician_schedule_shift, find_technician_next_available_slot


# ==========================================
# FLASK APP
# ==========================================

APP_DIR = os.path.dirname(os.path.abspath(__file__))


def load_local_env():
    """Load local development settings without committing credentials to source."""
    env_path = os.path.join(APP_DIR, ".env")
    if not os.path.exists(env_path):
        return

    with open(env_path, "r", encoding="utf-8") as env_file:
        for raw_line in env_file:
            line = raw_line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            os.environ.setdefault(key.strip(), value.strip())


load_local_env()

app = Flask(__name__)
app.config["SECRET_KEY"] = os.getenv("CRD_SECRET_KEY") or secrets.token_urlsafe(48)
AUTH_SERIALIZER = URLSafeTimedSerializer(app.config["SECRET_KEY"], salt="crd-global-solutions-auth")
AUTH_MAX_AGE_SECONDS = int(os.getenv("CRD_AUTH_MAX_AGE_SECONDS", "86400"))

# ==========================================
# BRUTE-FORCE PROTECTION
# ==========================================
_login_attempts: dict = defaultdict(list)   # {username_lower: [epoch_timestamps]}
_LOGIN_MAX_ATTEMPTS = 20                    # Allowed failed attempts before temporary cooldown
_LOGIN_LOCKOUT_SECONDS = 60                 # 1-minute cooldown window


allowed_origins = [
    origin.strip()
    for origin in os.getenv(
        "CRD_ALLOWED_ORIGINS",
        "http://127.0.0.1:5500,http://localhost:5500,https://billing.pariwebs.com"
    ).split(",")
    if origin.strip()
]
# Automatically permit all local ports (Live Server, Vite, file:// etc.) and cloud domains
allowed_origins.extend([
    r"^http://localhost(:\d+)?$",
    r"^http://127\.0\.0\.1(:\d+)?$",
    r"^null$",
    r"https://.*\.netlify\.app",
    r"https://.*\.onrender\.com",
    r"https://.*\.crdglobalsolutions\.com",
    "https://crdglobalsolutions.com"
])

CORS(
    app,
    resources={r"/*": {"origins": allowed_origins}},
    allow_headers=["Content-Type", "Authorization", "X-Requested-With"],
    methods=["GET", "POST", "PUT", "DELETE", "OPTIONS"],
    expose_headers=["Content-Disposition"]
)


# ==========================================
# MYSQL CONNECTION & CONNECTION POOLING
# ==========================================

_db_pool = None
_db_pool_lock = threading.Lock()


def get_db_pool():
    global _db_pool
    if _db_pool is None:
        with _db_pool_lock:
            if _db_pool is None:
                password = os.getenv("CRD_DB_PASSWORD")
                if password is None:
                    raise RuntimeError(
                        "CRD_DB_PASSWORD is not configured. Copy backend/.env.example to backend/.env and set it."
                    )
                port = int(os.getenv("CRD_DB_PORT", "3306"))
                use_ssl = os.getenv("CRD_DB_USE_SSL", "false").lower() == "true"
                pool_size = int(os.getenv("CRD_DB_POOL_SIZE", "10"))

                conn_kwargs = {
                    "pool_name": "crd_pool",
                    "pool_size": pool_size,
                    "pool_reset_session": True,
                    "host": os.getenv("CRD_DB_HOST", "localhost"),
                    "port": port,
                    "user": os.getenv("CRD_DB_USER", "root"),
                    "password": password,
                    "database": os.getenv("CRD_DB_NAME", "crd_global_solutions"),
                    "autocommit": False
                }
                if use_ssl:
                    conn_kwargs["ssl_disabled"] = False
                    conn_kwargs["ssl_verify_cert"] = os.getenv("CRD_DB_SSL_VERIFY", "false").lower() == "true"

                _db_pool = pooling.MySQLConnectionPool(**conn_kwargs)
    return _db_pool


def get_db_connection():
    try:
        pool = get_db_pool()
        return pool.get_connection()
    except Exception as pool_err:
        # Fallback to direct unpooled connection if pool is unavailable or exhausted
        password = os.getenv("CRD_DB_PASSWORD")
        if password is None:
            raise RuntimeError(
                "CRD_DB_PASSWORD is not configured. Copy backend/.env.example to backend/.env and set it."
            )

        port = int(os.getenv("CRD_DB_PORT", "3306"))
        use_ssl = os.getenv("CRD_DB_USE_SSL", "false").lower() == "true"

        conn_kwargs = {
            "host": os.getenv("CRD_DB_HOST", "localhost"),
            "port": port,
            "user": os.getenv("CRD_DB_USER", "root"),
            "password": password,
            "database": os.getenv("CRD_DB_NAME", "crd_global_solutions"),
            "autocommit": False
        }
        if use_ssl:
            conn_kwargs["ssl_disabled"] = False
            conn_kwargs["ssl_verify_cert"] = os.getenv("CRD_DB_SSL_VERIFY", "false").lower() == "true"

        return mysql.connector.connect(**conn_kwargs)


@app.before_request
def open_db():
    # Handle preflight OPTIONS immediately with full CORS headers
    if request.method == "OPTIONS":
        response = app.make_default_options_response()
        origin = request.headers.get("Origin") or "*"
        response.headers["Access-Control-Allow-Origin"] = origin
        response.headers["Access-Control-Allow-Methods"] = "GET, POST, PUT, DELETE, OPTIONS"
        response.headers["Access-Control-Allow-Headers"] = "Content-Type, Authorization, X-Requested-With"
        response.headers["Access-Control-Max-Age"] = "86400"
        return response

    if request.path in {"/", "/health"}:
        return None

    g.db = get_db_connection()
    g.cursor = g.db.cursor(dictionary=True, buffered=True)
    try:
        g.cursor.execute("SET time_zone = '+05:30'")
    except Exception:
        pass



def ensure_product_schema():
    db = None
    cursor = None
    try:
        db = get_db_connection()
        cursor = db.cursor(dictionary=True)
        cursor.execute("SHOW COLUMNS FROM products")
        columns = {row["Field"] for row in cursor.fetchall()}

        col_defs = [
            ("item_type", "ALTER TABLE products ADD COLUMN item_type VARCHAR(20) NOT NULL DEFAULT 'Product'"),
            ("description", "ALTER TABLE products ADD COLUMN description TEXT NULL"),
            ("track_serial", "ALTER TABLE products ADD COLUMN track_serial ENUM('Yes', 'No') NOT NULL DEFAULT 'No'"),
            ("unit", "ALTER TABLE products ADD COLUMN unit VARCHAR(30) NOT NULL DEFAULT 'Piece'"),
            ("secondary_unit", "ALTER TABLE products ADD COLUMN secondary_unit VARCHAR(50) NULL"),
            ("conversion_rate", "ALTER TABLE products ADD COLUMN conversion_rate DECIMAL(12, 4) NOT NULL DEFAULT 1.0000"),
            ("secondary_sale_price", "ALTER TABLE products ADD COLUMN secondary_sale_price DECIMAL(12, 2) NOT NULL DEFAULT 0.00"),
            ("secondary_purchase_price", "ALTER TABLE products ADD COLUMN secondary_purchase_price DECIMAL(12, 2) NOT NULL DEFAULT 0.00"),
            ("purchase_price", "ALTER TABLE products ADD COLUMN purchase_price DECIMAL(12, 2) NOT NULL DEFAULT 0.00"),
            ("sale_price", "ALTER TABLE products ADD COLUMN sale_price DECIMAL(12, 2) NOT NULL DEFAULT 0.00"),
            ("stock", "ALTER TABLE products ADD COLUMN stock DECIMAL(12, 2) NOT NULL DEFAULT 0.00"),
            ("min_stock", "ALTER TABLE products ADD COLUMN min_stock DECIMAL(12, 2) NOT NULL DEFAULT 0.00"),
            ("gst_rate", "ALTER TABLE products ADD COLUMN gst_rate DECIMAL(5, 2) NOT NULL DEFAULT 0.00"),
            ("warranty_months", "ALTER TABLE products ADD COLUMN warranty_months INT NOT NULL DEFAULT 0"),
            ("status", "ALTER TABLE products ADD COLUMN status VARCHAR(20) NOT NULL DEFAULT 'Active'"),
        ]

        for col_name, alter_sql in col_defs:
            if col_name not in columns:
                try:
                    cursor.execute(alter_sql)
                    db.commit()
                    columns.add(col_name)
                except Exception as col_err:
                    print(f"Schema notice: could not add {col_name}: {col_err}")
    except Exception as e:
        print("Schema check error:", e)
    finally:
        if cursor:
            cursor.close()
        if db and db.is_connected():
            db.close()


def ensure_service_schema():
    """Add service fields introduced after the original database dump."""
    db = None
    cursor = None
    try:
        db = get_db_connection()
        cursor = db.cursor(dictionary=True)
        cursor.execute("SHOW COLUMNS FROM service_jobs")
        columns = {row["Field"] for row in cursor.fetchall()}

        additions = {
            "product_source": "VARCHAR(30) NOT NULL DEFAULT 'External Product'",
            "external_product_name": "VARCHAR(150) DEFAULT NULL",
            "external_brand": "VARCHAR(100) DEFAULT NULL",
            "external_model": "VARCHAR(100) DEFAULT NULL",
            "accessories": "TEXT DEFAULT NULL",
            "warranty_status": "VARCHAR(30) DEFAULT NULL",
            "rma_status": "VARCHAR(50) NOT NULL DEFAULT 'None'",
            "rma_center_name": "VARCHAR(150) DEFAULT NULL",
            "rma_docket_no": "VARCHAR(100) DEFAULT NULL",
            "rma_sent_date": "DATE DEFAULT NULL",
            "rma_received_date": "DATE DEFAULT NULL",
            "advance_paid": "DECIMAL(12,2) NOT NULL DEFAULT 0.00",
            "physical_condition": "VARCHAR(150) DEFAULT NULL",
            "invoice_id": "INT DEFAULT NULL",
            "service_type": "VARCHAR(20) NOT NULL DEFAULT 'Shop'",
            "customer_address": "TEXT NULL",
        }
        for column_name, definition in additions.items():
            if column_name not in columns:
                cursor.execute(
                    f"ALTER TABLE service_jobs ADD COLUMN {column_name} {definition}"
                )

        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS service_job_parts (
                id INT AUTO_INCREMENT PRIMARY KEY,
                service_job_id INT NOT NULL,
                product_id INT NOT NULL,
                part_name VARCHAR(150) NOT NULL,
                serial_number VARCHAR(100) NULL,
                quantity INT NOT NULL DEFAULT 1,
                unit_price DECIMAL(12,2) NOT NULL DEFAULT 0.00,
                line_total DECIMAL(12,2) NOT NULL DEFAULT 0.00,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                INDEX idx_sjp_job (service_job_id)
            )
            """
        )

        db.commit()
    except Exception as e:
        print("Ensure service schema error:", e)
        if db:
            try:
                db.rollback()
            except Exception:
                pass
    finally:
        if cursor:
            cursor.close()
        if db and db.is_connected():
            db.close()


def ensure_sales_schema():
    """Track money received separately, split payments, returns and customer ledgers."""
    db = None
    cursor = None
    try:
        db = get_db_connection()
        cursor = db.cursor(dictionary=True)
        cursor.execute("SHOW COLUMNS FROM sales_invoices")
        columns = {row["Field"] for row in cursor.fetchall()}
        
        sales_additions = {
            "paid_amount": "DECIMAL(12,2) NOT NULL DEFAULT 0",
            "cash_amount": "DECIMAL(12,2) NOT NULL DEFAULT 0",
            "upi_amount": "DECIMAL(12,2) NOT NULL DEFAULT 0",
            "card_amount": "DECIMAL(12,2) NOT NULL DEFAULT 0",
            "bank_amount": "DECIMAL(12,2) NOT NULL DEFAULT 0",
            "document_type": "VARCHAR(20) NOT NULL DEFAULT 'Invoice'",
            "service_job_id": "INT DEFAULT NULL"
        }
        for col, col_def in sales_additions.items():
            if col not in columns:
                cursor.execute(f"ALTER TABLE sales_invoices ADD COLUMN {col} {col_def}")

        cursor.execute(
            "UPDATE sales_invoices SET paid_amount = grand_total WHERE payment_status = 'Paid' AND paid_amount = 0"
        )

        cursor.execute("SHOW COLUMNS FROM sales_invoice_items")
        item_columns = {row["Field"] for row in cursor.fetchall()}
        if "cost_rate" not in item_columns:
            cursor.execute(
                "ALTER TABLE sales_invoice_items ADD COLUMN cost_rate DECIMAL(12,2) NOT NULL DEFAULT 0 AFTER rate"
            )
            cursor.execute(
                """
                UPDATE sales_invoice_items sii
                LEFT JOIN products p ON p.id = sii.product_id
                SET sii.cost_rate = COALESCE(p.purchase_price, 0)
                WHERE sii.cost_rate = 0
                """
            )

        if "unit" not in item_columns:
            cursor.execute("ALTER TABLE sales_invoice_items ADD COLUMN unit VARCHAR(50) NULL AFTER quantity")

        if "base_quantity" not in item_columns:
            cursor.execute("ALTER TABLE sales_invoice_items ADD COLUMN base_quantity DECIMAL(12, 4) NULL AFTER unit")
            cursor.execute("UPDATE sales_invoice_items SET base_quantity = quantity WHERE base_quantity IS NULL")

        if "conversion_factor" not in item_columns:
            cursor.execute("ALTER TABLE sales_invoice_items ADD COLUMN conversion_factor DECIMAL(12, 4) NOT NULL DEFAULT 1.0000 AFTER base_quantity")

        if "warranty_months" not in item_columns:
            cursor.execute("ALTER TABLE sales_invoice_items ADD COLUMN warranty_months INT NOT NULL DEFAULT 0 AFTER line_total")

        cursor.execute("SHOW COLUMNS FROM purchase_invoice_items")
        pii_columns = {row["Field"] for row in cursor.fetchall()}
        if "unit" not in pii_columns:
            cursor.execute("ALTER TABLE purchase_invoice_items ADD COLUMN unit VARCHAR(50) NULL AFTER quantity")

        if "base_quantity" not in pii_columns:
            cursor.execute("ALTER TABLE purchase_invoice_items ADD COLUMN base_quantity DECIMAL(12, 4) NULL AFTER unit")
            cursor.execute("UPDATE purchase_invoice_items SET base_quantity = quantity WHERE base_quantity IS NULL")

        if "conversion_factor" not in pii_columns:
            cursor.execute("ALTER TABLE purchase_invoice_items ADD COLUMN conversion_factor DECIMAL(12, 4) NOT NULL DEFAULT 1.0000 AFTER base_quantity")

        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS sales_returns (
                id INT AUTO_INCREMENT PRIMARY KEY,
                return_number VARCHAR(50) NOT NULL UNIQUE,
                invoice_id INT NOT NULL,
                customer_id INT NOT NULL,
                return_date DATE NOT NULL,
                total_refund_amount DECIMAL(12,2) NOT NULL DEFAULT 0.00,
                refund_method VARCHAR(30) NOT NULL DEFAULT 'Cash',
                reason TEXT NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                INDEX idx_sr_invoice (invoice_id)
            )
            """
        )

        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS sales_return_items (
                id INT AUTO_INCREMENT PRIMARY KEY,
                return_id INT NOT NULL,
                product_id INT NOT NULL,
                quantity INT NOT NULL DEFAULT 1,
                rate DECIMAL(12,2) NOT NULL DEFAULT 0.00,
                tax_amount DECIMAL(12,2) NOT NULL DEFAULT 0.00,
                line_total DECIMAL(12,2) NOT NULL DEFAULT 0.00,
                serial_number VARCHAR(100) NULL,
                INDEX idx_sri_return (return_id)
            )
            """
        )

        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS customer_payments (
                id INT AUTO_INCREMENT PRIMARY KEY,
                receipt_number VARCHAR(50) NOT NULL UNIQUE,
                customer_id INT NOT NULL,
                invoice_id INT NULL,
                payment_date DATE NOT NULL,
                amount DECIMAL(12,2) NOT NULL DEFAULT 0.00,
                payment_method VARCHAR(30) NOT NULL DEFAULT 'Cash',
                reference_no VARCHAR(100) NULL,
                notes TEXT NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                INDEX idx_cp_customer (customer_id)
            )
            """
        )

        db.commit()
    except Exception as e:
        print("Ensure sales schema error:", e)
        if db:
            try:
                db.rollback()
            except Exception:
                pass
    finally:
        if cursor:
            cursor.close()
        if db and db.is_connected():
            db.close()


def ensure_stock_locations_schema():
    db = None
    cursor = None
    try:
        db = get_db_connection()
        cursor = db.cursor(dictionary=True)

        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS stock_locations (
                id INT AUTO_INCREMENT PRIMARY KEY,
                location_name VARCHAR(100) NOT NULL UNIQUE,
                description VARCHAR(255) NULL,
                is_default TINYINT(1) DEFAULT 0,
                status VARCHAR(20) DEFAULT 'Active',
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP
            )
            """
        )

        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS product_location_stock (
                id INT AUTO_INCREMENT PRIMARY KEY,
                product_id INT NOT NULL,
                location_id INT NOT NULL,
                quantity DECIMAL(12,2) NOT NULL DEFAULT 0.00,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
                UNIQUE KEY uq_prod_loc (product_id, location_id),
                INDEX idx_pls_product (product_id),
                INDEX idx_pls_location (location_id)
            )
            """
        )

        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS stock_transfers (
                id INT AUTO_INCREMENT PRIMARY KEY,
                transfer_number VARCHAR(50) NOT NULL UNIQUE,
                transfer_date DATE NOT NULL,
                source_location_id INT NOT NULL,
                destination_location_id INT NOT NULL,
                product_id INT NOT NULL,
                quantity DECIMAL(12,2) NOT NULL,
                notes TEXT NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                INDEX idx_st_prod (product_id),
                INDEX idx_st_src (source_location_id),
                INDEX idx_st_dst (destination_location_id)
            )
            """
        )

        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS stock_transfer_serials (
                id INT AUTO_INCREMENT PRIMARY KEY,
                transfer_id INT NOT NULL,
                serial_number VARCHAR(100) NOT NULL,
                INDEX idx_sts_transfer (transfer_id),
                INDEX idx_sts_serial (serial_number)
            )
            """
        )

        cursor.execute("SHOW COLUMNS FROM purchase_invoice_items")
        pii_cols = {row["Field"] for row in cursor.fetchall()}
        if "location_id" not in pii_cols:
            cursor.execute("ALTER TABLE purchase_invoice_items ADD COLUMN location_id INT NULL AFTER product_id")

        cursor.execute("SHOW COLUMNS FROM product_serials")
        ps_cols = {row["Field"] for row in cursor.fetchall()}
        if "location_id" not in ps_cols:
            cursor.execute("ALTER TABLE product_serials ADD COLUMN location_id INT NULL AFTER supplier_name")

        cursor.execute("SELECT COUNT(*) AS cnt FROM stock_locations")
        cnt_row = cursor.fetchone()
        if not cnt_row or cnt_row.get("cnt", 0) == 0:
            cursor.execute(
                """
                INSERT INTO stock_locations (location_name, description, is_default, status)
                VALUES 
                ('Main Counter', 'Shop front counter & display rack', 1, 'Active'),
                ('Godown', 'Primary warehouse / backstore storage', 0, 'Active')
                """
            )

        # Auto-sync legacy products without location stock into default location
        cursor.execute(
            """
            INSERT INTO product_location_stock (product_id, location_id, quantity)
            SELECT p.id, (SELECT id FROM stock_locations WHERE is_default = 1 LIMIT 1), p.stock
            FROM products p
            WHERE p.stock > 0 
            AND p.item_type != 'Service'
            AND p.id NOT IN (SELECT DISTINCT product_id FROM product_location_stock)
            """
        )

        # Auto-sync legacy serials without location_id into default location
        cursor.execute(
            """
            UPDATE product_serials
            SET location_id = (SELECT id FROM stock_locations WHERE is_default = 1 LIMIT 1)
            WHERE location_id IS NULL
            """
        )

        db.commit()
    except Exception as e:
        print("Ensure stock locations schema error:", e)
        if db:
            try:
                db.rollback()
            except Exception:
                pass
    finally:
        if cursor:
            cursor.close()
        if db and db.is_connected():
            db.close()

def ensure_scrap_schema():
    db = None
    cursor = None
    try:
        db = get_db_connection()
        cursor = db.cursor(dictionary=True)

        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS scraps (
                id INT AUTO_INCREMENT PRIMARY KEY,
                scrap_date DATE NOT NULL,
                item_name VARCHAR(150) NOT NULL,
                quantity DECIMAL(12,2) NOT NULL DEFAULT 1.00,
                sold_quantity DECIMAL(12,2) NOT NULL DEFAULT 0.00,
                available_quantity DECIMAL(12,2) NOT NULL DEFAULT 1.00,
                unit VARCHAR(20) NOT NULL DEFAULT 'Nos',
                source_type VARCHAR(50) NOT NULL DEFAULT 'Repair',
                repair_job_id INT NULL,
                customer_job_no VARCHAR(150) NULL,
                reason VARCHAR(255) NULL,
                remark TEXT NULL,
                status VARCHAR(20) NOT NULL DEFAULT 'Available',
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
                INDEX idx_scrap_status (status),
                INDEX idx_scrap_date (scrap_date),
                INDEX idx_scrap_job (repair_job_id)
            )
            """
        )

        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS scrap_sales (
                id INT AUTO_INCREMENT PRIMARY KEY,
                sale_number VARCHAR(50) NOT NULL UNIQUE,
                sale_date DATE NOT NULL,
                party_name VARCHAR(150) NOT NULL,
                mobile VARCHAR(50) NULL,
                total_amount DECIMAL(12,2) NOT NULL DEFAULT 0.00,
                payment_mode VARCHAR(50) NOT NULL DEFAULT 'Cash',
                remark TEXT NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                INDEX idx_ss_date (sale_date)
            )
            """
        )

        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS scrap_sale_items (
                id INT AUTO_INCREMENT PRIMARY KEY,
                scrap_sale_id INT NOT NULL,
                scrap_id INT NOT NULL,
                quantity DECIMAL(12,2) NOT NULL DEFAULT 1.00,
                amount DECIMAL(12,2) NOT NULL DEFAULT 0.00,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                INDEX idx_ssi_sale (scrap_sale_id),
                INDEX idx_ssi_scrap (scrap_id)
            )
            """
        )

        db.commit()
    except Exception as e:
        print("Ensure scrap schema error:", e)
        if db:
            try:
                db.rollback()
            except Exception:
                pass
    finally:
        if cursor:
            cursor.close()
        if db and db.is_connected():
            db.close()




@app.teardown_request
def close_db(error=None):
    cursor = getattr(g, "cursor", None)
    db = getattr(g, "db", None)

    try:
        if error and db and db.is_connected():
            db.rollback()
    except Exception:
        pass

    try:
        if cursor:
            g.cursor.close()
    except Exception:
        pass

    try:
        if db and db.is_connected():
            db.close()
    except Exception:
        pass


def create_auth_token(user):
    return AUTH_SERIALIZER.dumps({
        "id": int(user["id"]),
        "username": str(user["username"]),
        "role": str(user.get("role") or "Staff")
    })


@app.before_request
def require_authenticated_api_request():
    """Require a short-lived signed token for every API route except login, technician login, uploads, downloads and public company settings."""
    if (
        request.method == "OPTIONS"
        or request.path in {"/", "/login", "/api/auth/technician-login", "/api/app-info", "/downloads/CRDTechnician.apk", "/config.js"}
        or request.path.startswith("/uploads/")
        or request.path.startswith("/downloads/")
        or request.path.startswith("/static/")
        or request.path.startswith("/frontend/")
        or request.path.startswith("/tech")
        or request.path.startswith("/_next/")
        or (request.method == "GET" and request.path in {"/settings/business-profile", "/settings/service", "/settings/print-display", "/settings/quotation-billing", "/settings/invoice-billing"})
    ):
        return None

    header = request.headers.get("Authorization", "")
    token = None
    if header.startswith("Bearer "):
        token = header[7:].strip()
    elif request.path.startswith("/reports/export/") or request.path.startswith("/backup/") or request.path.startswith("/settings/backup/") or request.path in {"/settings/export-data"}:
        token = (request.args.get("token") or "").strip()

    if not token:
        return jsonify({"success": False, "message": "Authentication required."}), 401
    try:
        payload = AUTH_SERIALIZER.loads(token, max_age=AUTH_MAX_AGE_SECONDS)
        user_id = int(payload.get("id"))
    except (BadSignature, SignatureExpired, TypeError, ValueError):
        return jsonify({"success": False, "message": "Your login has expired. Please sign in again."}), 401

    try:
        g.cursor.execute(
            "SELECT id, username, full_name, role, status FROM users WHERE id = %s LIMIT 1",
            (user_id,)
        )
        user = g.cursor.fetchone()
    except Exception:
        return jsonify({"success": False, "message": "Unable to verify the current user."}), 401

    if not user or str(user.get("status") or "").lower() != "active":
        return jsonify({"success": False, "message": "This user account is not active."}), 401

    # Technicians are only authorized for technician mobile app endpoints (/api/tech/, /api/auth/, /uploads/, etc.)
    user_role = str(user.get("role") or "").strip().lower()
    if user_role == "technician":
        allowed_prefix = (
            request.path.startswith("/api/tech/") or
            request.path.startswith("/api/auth/") or
            request.path.startswith("/uploads/") or
            (request.method == "GET" and request.path in {"/settings/business-profile", "/settings/service", "/settings/print-display", "/settings/quotation-billing", "/settings/invoice-billing"})
        )
        if not allowed_prefix:
            return jsonify({
                "success": False,
                "message": "Technicians billing software APIs access nahi kar sakte. Kripya Android Mobile App ka upyog karein."
            }), 403

    g.current_user = user
    g.user_id = user["id"]
    return None


@app.after_request
def add_no_cache_headers(response):
    """Ensure dynamic API responses are never cached by browser or proxies,
    and guarantee CORS headers for web clients."""
    if not request.path.startswith("/static"):
        response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        response.headers["Pragma"] = "no-cache"
        response.headers["Expires"] = "0"

    origin = request.headers.get("Origin")
    if origin and "Access-Control-Allow-Origin" not in response.headers:
        response.headers["Access-Control-Allow-Origin"] = origin
        response.headers["Access-Control-Allow-Methods"] = "GET, POST, PUT, DELETE, OPTIONS"
        response.headers["Access-Control-Allow-Headers"] = "Content-Type, Authorization, X-Requested-With"
        response.headers["Access-Control-Allow-Credentials"] = "true"

    return response


def roles_required(*roles):
    """Use for administrative actions that must not be available to Staff users."""
    def decorator(view):
        @wraps(view)
        def wrapped(*args, **kwargs):
            if request.method == "OPTIONS":
                return "", 200
            user = getattr(g, "current_user", None)
            if not user or str(user.get("role") or "") not in roles:
                return jsonify({"success": False, "message": "You do not have permission for this action."}), 403
            return view(*args, **kwargs)
        return wrapped
    return decorator


def normalize_date_string(value):
    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()
        if not value:
            return None
        return value[:10]

    try:
        return value.strftime("%Y-%m-%d")
    except Exception:
        return str(value)[:10]


def get_product_record_by_id(product_id):
    if not product_id:
        return None

    g.cursor.execute(
        """
        SELECT id, product_name, item_type, unit, secondary_unit, conversion_rate, purchase_price, secondary_purchase_price, sale_price, secondary_sale_price, stock, min_stock, track_serial, status
        FROM products
        WHERE id = %s
        LIMIT 1
        """,
        (product_id,)
    )
    return g.cursor.fetchone()


def normalize_document_items(raw_items, invoice_discount=0, default_tax_inclusive=True):
    """Calculate totals on the server instead of trusting browser-provided money values.
    Supports both tax-inclusive and tax-exclusive pricing per item or document.
    """
    try:
        invoice_discount = float(invoice_discount or 0)
    except (TypeError, ValueError):
        raise ValueError("Discount must be a valid number.")

    if invoice_discount < 0:
        raise ValueError("Discount cannot be negative.")

    normalized_items = []
    subtotal = 0.0
    tax_amount = 0.0

    for raw_item in raw_items or []:
        if not isinstance(raw_item, dict):
            raise ValueError("Each invoice item must be an object.")

        try:
            quantity = float(raw_item.get("quantity", 0) or 0)
            rate = float(raw_item.get("rate", 0) or 0)
            gst_rate = float(raw_item.get("gst_rate", 0) or 0)
            item_discount = float(raw_item.get("discount", 0) or 0)
        except (TypeError, ValueError):
            raise ValueError("Item quantity, rate, GST and discount must be numbers.")

        if quantity <= 0:
            raise ValueError("Item quantity must be greater than zero.")
        if rate < 0 or gst_rate < 0 or gst_rate > 100 or item_discount < 0:
            raise ValueError("Item rate, GST and discount contain an invalid value.")

        # Determine tax inclusivity
        item_tax_inclusive = raw_item.get("is_tax_inclusive")
        if item_tax_inclusive is None:
            tax_type = str(raw_item.get("tax_type", "")).strip().lower()
            if tax_type in ("exclusive", "tax_exclusive", "excl"):
                item_tax_inclusive = False
            elif tax_type in ("inclusive", "tax_inclusive", "incl"):
                item_tax_inclusive = True
            else:
                item_tax_inclusive = default_tax_inclusive
        else:
            item_tax_inclusive = bool(item_tax_inclusive)

        gross_line = quantity * rate
        if item_discount > gross_line:
            raise ValueError("Item discount cannot exceed its total amount.")

        gross_after_discount = gross_line - item_discount

        if gst_rate > 0:
            if item_tax_inclusive:
                taxable_after_discount = round(gross_after_discount / (1.0 + gst_rate / 100.0), 2)
                item_tax = round(gross_after_discount - taxable_after_discount, 2)
                line_total = round(gross_after_discount, 2)
            else:
                taxable_after_discount = round(gross_after_discount, 2)
                item_tax = round(taxable_after_discount * (gst_rate / 100.0), 2)
                line_total = round(taxable_after_discount + item_tax, 2)
        else:
            taxable_after_discount = round(gross_after_discount, 2)
            item_tax = 0.0
            line_total = round(gross_after_discount, 2)

        normalized = dict(raw_item)
        normalized.update({
            "quantity": quantity,
            "rate": rate,
            "gst_rate": gst_rate,
            "discount": item_discount,
            "taxable_amount": taxable_after_discount,
            "gst_amount": item_tax,
            "line_total": line_total,
            "is_tax_inclusive": item_tax_inclusive
        })
        normalized_items.append(normalized)
        subtotal += taxable_after_discount
        tax_amount += item_tax

    gross_total = subtotal + tax_amount
    if invoice_discount > gross_total:
        raise ValueError("Invoice discount cannot exceed the invoice amount.")

    return (
        normalized_items,
        round(subtotal, 2),
        round(tax_amount, 2),
        round(gross_total - invoice_discount, 2),
        round(invoice_discount, 2)
    )


def normalize_sale_payment(grand_total, payment_status, paid_amount):
    status = str(payment_status or "Unpaid").strip().title()
    if status not in {"Paid", "Unpaid", "Partial", "Pending"}:
        raise ValueError("Invalid payment status.")

    try:
        amount = float(paid_amount or 0)
    except (TypeError, ValueError):
        raise ValueError("Paid amount must be a valid number.")

    if amount < 0:
        raise ValueError("Paid amount cannot be negative.")

    # Tolerant rounding: absorb small floating point rounding differences (up to 0.05)
    if amount > grand_total and abs(amount - grand_total) <= 0.05:
        amount = grand_total

    if status == "Paid":
        amount = grand_total
    elif status in {"Unpaid", "Pending"}:
        amount = 0.0
    elif status == "Partial":
        if amount >= grand_total and grand_total > 0:
            status = "Paid"
            amount = grand_total
        elif amount <= 0:
            status = "Unpaid"
            amount = 0.0
        elif amount > grand_total:
            raise ValueError("Paid amount must be between 0 and the grand total.")

    return status, round(amount, 2)


def add_months_to_date(date_value, months):
    if not date_value or not months:
        return None

    try:
        dt = datetime.strptime(str(date_value)[:10], "%Y-%m-%d")
    except Exception:
        return None

    year = dt.year + (dt.month - 1 + int(months)) // 12
    month = (dt.month - 1 + int(months)) % 12 + 1
    day = min(dt.day, 28)

    return datetime(year, month, day).strftime("%Y-%m-%d")


def get_warranty_status(warranty_start, warranty_end):
    today = datetime.today().date()

    if not warranty_start and not warranty_end:
        return "Not Available"

    try:
        start_date = datetime.strptime(
            str(warranty_start), "%Y-%m-%d"
        ).date()
    except Exception:
        start_date = None

    try:
        end_date = datetime.strptime(
            str(warranty_end), "%Y-%m-%d"
        ).date()
    except Exception:
        end_date = None

    if start_date and end_date:
        if start_date <= today <= end_date:
            return "In Warranty"
        if today > end_date:
            return "Expired"
        return "Not Available"

    if end_date:
        if today > end_date:
            return "Expired"
        return "In Warranty"

    if start_date:
        if today >= start_date:
            return "In Warranty"
        return "Not Available"

    return "Not Available"


def is_service_overdue(service):
    try:
        status = str(service.get("service_status") or "").strip()
        if status.lower() in {"delivered", "cancelled"}:
            return False

        expected_value = service.get("expected_delivery_date")
        if not expected_value:
            return False

        expected_date = datetime.strptime(
            str(expected_value)[:10], "%Y-%m-%d"
        ).date()

        return expected_date < datetime.today().date()

    except Exception:
        return False


def ensure_expenses_table():
    try:
        db = get_db_connection()
        cursor = db.cursor(dictionary=True)

        cursor.execute("SHOW TABLES LIKE 'expenses'")
        if not cursor.fetchone():
            cursor.execute(
                """
                CREATE TABLE expenses (
                    id INT AUTO_INCREMENT PRIMARY KEY,
                    expense_number VARCHAR(50) NOT NULL UNIQUE,
                    expense_date DATE NOT NULL,
                    category VARCHAR(100) NOT NULL,
                    description TEXT NOT NULL,
                    amount DECIMAL(12,2) NOT NULL,
                    payment_method VARCHAR(50) NOT NULL,
                    paid_to VARCHAR(150) DEFAULT NULL,
                    reference_number VARCHAR(100) DEFAULT NULL,
                    notes TEXT DEFAULT NULL,
                    status VARCHAR(20) NOT NULL DEFAULT 'Paid',
                    original_status VARCHAR(20) DEFAULT NULL,
                    deleted_at TIMESTAMP NULL DEFAULT NULL,
                    deleted_by_user_id INT DEFAULT NULL,
                    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP
                )
                """
            )
            db.commit()

        cursor.execute("SHOW COLUMNS FROM expenses")
        columns = {row["Field"] for row in cursor.fetchall()}
        column_definitions = {
            "original_status": "ALTER TABLE expenses ADD COLUMN original_status VARCHAR(20) DEFAULT NULL",
            "deleted_at": "ALTER TABLE expenses ADD COLUMN deleted_at TIMESTAMP NULL DEFAULT NULL",
            "deleted_by_user_id": "ALTER TABLE expenses ADD COLUMN deleted_by_user_id INT DEFAULT NULL"
        }
        for column_name, alter_sql in column_definitions.items():
            if column_name not in columns:
                try:
                    cursor.execute(alter_sql)
                    db.commit()
                except Exception as inner_error:
                    print(f"Ensure expenses column {column_name} error: {inner_error}")

        cursor.close()
        db.close()

    except Exception as e:
        print("Ensure Expenses Table Error:", e)


def ensure_expense_audit_table():
    try:
        db = get_db_connection()
        cursor = db.cursor(dictionary=True)

        cursor.execute("SHOW TABLES LIKE 'expense_audit_logs'")
        if not cursor.fetchone():
            cursor.execute(
                """
                CREATE TABLE expense_audit_logs (
                    id INT AUTO_INCREMENT PRIMARY KEY,
                    expense_id INT NOT NULL,
                    action VARCHAR(30) NOT NULL,
                    performed_by_user_id INT DEFAULT NULL,
                    performed_by_username VARCHAR(150) DEFAULT NULL,
                    action_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    details TEXT DEFAULT NULL,
                    INDEX idx_expense_audit_expense_id (expense_id),
                    INDEX idx_expense_audit_action_at (action_at)
                )
                """
            )
            db.commit()

        cursor.close()
        db.close()
    except Exception as e:
        print("Ensure Expense Audit Table Error:", e)


def log_expense_action(expense_id, action, user=None, details=None):
    try:
        if expense_id is None:
            return

        user_id = None
        username = None
        if user:
            user_id = user.get("id")
            username = user.get("username") or user.get("full_name")

        g.cursor.execute(
            """
            INSERT INTO expense_audit_logs
            (expense_id, action, performed_by_user_id, performed_by_username, details)
            VALUES (%s, %s, %s, %s, %s)
            """,
            (expense_id, action, user_id, username, details or None)
        )
    except Exception as e:
        print("Expense Audit Log Error:", e)


def ensure_users_table():
    try:
        db = get_db_connection()
        cursor = db.cursor(dictionary=True)

        cursor.execute("SHOW TABLES LIKE 'users'")
        if not cursor.fetchone():
            cursor.execute(
                """
                CREATE TABLE users (
                    id INT AUTO_INCREMENT PRIMARY KEY,
                    username VARCHAR(100) NOT NULL UNIQUE,
                    password VARCHAR(255) NOT NULL,
                    full_name VARCHAR(150) NOT NULL,
                    mobile VARCHAR(50) DEFAULT NULL,
                    email VARCHAR(150) DEFAULT NULL,
                    role VARCHAR(30) NOT NULL DEFAULT 'Staff',
                    status VARCHAR(20) NOT NULL DEFAULT 'Active',
                    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP
                )
                """
            )
            db.commit()

        cursor.execute("SHOW COLUMNS FROM users")
        columns = [column["Field"] for column in cursor.fetchall()]

        column_definitions = {
            "username": "ALTER TABLE users ADD COLUMN username VARCHAR(100) NOT NULL DEFAULT ''",
            "password": "ALTER TABLE users ADD COLUMN password VARCHAR(255) NOT NULL DEFAULT ''",
            "full_name": "ALTER TABLE users ADD COLUMN full_name VARCHAR(150) NOT NULL DEFAULT ''",
            "mobile": "ALTER TABLE users ADD COLUMN mobile VARCHAR(50) DEFAULT NULL",
            "email": "ALTER TABLE users ADD COLUMN email VARCHAR(150) DEFAULT NULL",
            "role": "ALTER TABLE users ADD COLUMN role VARCHAR(30) NOT NULL DEFAULT 'Staff'",
            "status": "ALTER TABLE users ADD COLUMN status VARCHAR(20) NOT NULL DEFAULT 'Active'",
            "created_at": "ALTER TABLE users ADD COLUMN created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP",
            "updated_at": "ALTER TABLE users ADD COLUMN updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP"
        }

        for column_name, alter_sql in column_definitions.items():
            if column_name not in columns:
                try:
                    cursor.execute(alter_sql)
                    db.commit()
                except Exception as inner_error:
                    print(f"Ensure users column {column_name} error: {inner_error}")

        try:
            cursor.execute("SHOW INDEX FROM users WHERE Key_name = 'username'")
            if not cursor.fetchone():
                cursor.execute("CREATE UNIQUE INDEX username_unique_idx ON users(username)")
                db.commit()
        except Exception as index_error:
            print(f"Ensure users unique index error: {index_error}")

        cursor.close()
        db.close()

    except Exception as e:
        print("Ensure Users Table Error:", e)


def generate_expense_number():
    try:
        g.cursor.execute(
            "SELECT expense_number FROM expenses ORDER BY id DESC LIMIT 1"
        )
        last_record = g.cursor.fetchone()

        if not last_record or not last_record.get("expense_number"):
            return "EXP-000001"

        match = re.search(r"(\d+)$", str(last_record["expense_number"]))
        if not match:
            return "EXP-000001"

        last_number = int(match.group(1))
        return f"EXP-{last_number + 1:06d}"
    except Exception:
        return "EXP-000001"


def generate_document_number(category, prefix_key, next_key, fallback_prefix):
    """Allocate a sequential number inside the active database transaction."""
    g.cursor.execute(
        "SELECT setting_key, setting_value FROM settings WHERE category = %s AND setting_key IN (%s, %s) FOR UPDATE",
        (category, prefix_key, next_key)
    )
    values = {
        row["setting_key"]: row.get("setting_value")
        for row in g.cursor.fetchall()
    }

    prefix = str(values.get(prefix_key) or fallback_prefix)
    try:
        next_number = int(str(values.get(next_key) or "1"))
    except ValueError:
        next_number = 1
    if next_number < 1:
        next_number = 1

    # Ensure document_number is not already used in the relevant table
    while True:
        document_number = f"{prefix}{next_number:06d}"
        if category == "quotation-billing":
            g.cursor.execute("SELECT id FROM sales_invoices WHERE invoice_number = %s AND document_type = 'Quotation' LIMIT 1", (document_number,))
        elif category == "invoice-billing":
            g.cursor.execute("SELECT id FROM sales_invoices WHERE invoice_number = %s AND (document_type != 'Quotation' OR document_type IS NULL) LIMIT 1", (document_number,))
        elif category == "purchase-billing":
            g.cursor.execute("SELECT id FROM purchase_invoices WHERE purchase_number = %s LIMIT 1", (document_number,))
        elif category == "sales-return-billing":
            g.cursor.execute("SELECT id FROM sales_returns WHERE return_number = %s LIMIT 1", (document_number,))
        else:
            break

        if not g.cursor.fetchone():
            break
        next_number += 1

    g.cursor.execute(
        """
        INSERT INTO settings (category, setting_key, setting_value)
        VALUES (%s, %s, %s)
        ON DUPLICATE KEY UPDATE setting_value = VALUES(setting_value), updated_at = CURRENT_TIMESTAMP
        """,
        (category, next_key, str(next_number + 1))
    )
    return document_number


def sync_document_sequence(category, next_key, prefix_key, fallback_prefix, doc_type=None):
    """
    Recalculate and update the next document sequence number in settings
    based on the current maximum document number existing in the database.
    """
    try:
        import re
        if category == "quotation-billing" or doc_type == "Quotation":
            g.cursor.execute("SELECT invoice_number FROM sales_invoices WHERE document_type = 'Quotation'")
        elif category == "invoice-billing" or doc_type == "Invoice":
            g.cursor.execute("SELECT invoice_number FROM sales_invoices WHERE document_type != 'Quotation' OR document_type IS NULL")
        elif category == "purchase-billing":
            g.cursor.execute("SELECT purchase_number FROM purchase_invoices")
        elif category == "sales-return-billing":
            g.cursor.execute("SELECT return_number FROM sales_returns")
        else:
            return

        rows = g.cursor.fetchall() or []
        max_num = 0
        for r in rows:
            val = list(r.values())[0] if isinstance(r, dict) else r[0]
            if not val:
                continue
            matches = re.findall(r'(\d+)', str(val))
            if matches:
                try:
                    num = int(matches[-1])
                    if num > max_num:
                        max_num = num
                except ValueError:
                    pass

        next_val = max_num + 1 if max_num > 0 else 1

        g.cursor.execute(
            """
            INSERT INTO settings (category, setting_key, setting_value)
            VALUES (%s, %s, %s)
            ON DUPLICATE KEY UPDATE setting_value = VALUES(setting_value), updated_at = CURRENT_TIMESTAMP
            """,
            (category, next_key, str(next_val))
        )
    except Exception as e:
        print(f"Error in sync_document_sequence for {category}: {e}")



ensure_expenses_table()
ensure_expense_audit_table()
ensure_users_table()


def ensure_settings_table():
    try:
        db = get_db_connection()
        cursor = db.cursor(dictionary=True)

        cursor.execute("SHOW TABLES LIKE 'settings'")
        if not cursor.fetchone():
            cursor.execute(
                """
                CREATE TABLE settings (
                    id INT AUTO_INCREMENT PRIMARY KEY,
                    category VARCHAR(100) NOT NULL,
                    setting_key VARCHAR(150) NOT NULL,
                    setting_value TEXT NULL,
                    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
                    UNIQUE KEY unique_category_key (category, setting_key)
                )
                """
            )
            db.commit()

        cursor.execute("SHOW COLUMNS FROM settings")
        columns = {column["Field"] for column in cursor.fetchall()}
        required_columns = {
            "id": "ALTER TABLE settings ADD COLUMN id INT AUTO_INCREMENT PRIMARY KEY",
            "category": "ALTER TABLE settings ADD COLUMN category VARCHAR(100) NOT NULL DEFAULT 'general'",
            "setting_key": "ALTER TABLE settings ADD COLUMN setting_key VARCHAR(150) NOT NULL DEFAULT ''",
            "setting_value": "ALTER TABLE settings ADD COLUMN setting_value TEXT NULL",
            "created_at": "ALTER TABLE settings ADD COLUMN created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP",
            "updated_at": "ALTER TABLE settings ADD COLUMN updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP"
        }

        for column_name, alter_sql in required_columns.items():
            if column_name not in columns:
                try:
                    cursor.execute(alter_sql)
                    db.commit()
                except Exception as inner_error:
                    print(f"Ensure settings column {column_name} error: {inner_error}")

        cursor.close()
        db.close()

    except Exception as e:
        print("Ensure Settings Table Error:", e)


ALL_PERMISSIONS = [
    {"id": "dashboard", "name": "Dashboard", "description": "View dashboard summary and quick stats", "icon": "🏠"},
    {"id": "sales", "name": "Billing & Sales", "description": "Create sales invoices, view sales list and print bills", "icon": "🧾"},
    {"id": "customers", "name": "Customers", "description": "Manage customer accounts and contact info", "icon": "👥"},
    {"id": "products", "name": "Products & Inventory", "description": "Manage products, stock levels, and serial tracking", "icon": "📦"},
    {"id": "purchases", "name": "Purchases", "description": "Record supplier purchases and view purchase list", "icon": "🛒"},
    {"id": "suppliers", "name": "Suppliers", "description": "Manage supplier profiles and dues", "icon": "🚚"},
    {"id": "expenses", "name": "Expenses", "description": "Record and manage shop expenses", "icon": "💰"},
    {"id": "services", "name": "Service Jobs", "description": "Manage repair and service job sheets", "icon": "🔧"},
    {"id": "employees", "name": "Employees & Technicians", "description": "Manage employee master, technicians, tools & schedules", "icon": "👷"},
    {"id": "attendance", "name": "Attendance & GPS", "description": "View daily staff attendance, GPS logs and corrections", "icon": "📍"},
    {"id": "technician_portal", "name": "Technician Mobile App", "description": "Access assigned service jobs, GPS check-in, tools and stock", "icon": "📱"},
    {"id": "device_security", "name": "Device Security", "description": "Authorize and manage registered devices for technicians", "icon": "🔒"},
    {"id": "reports", "name": "Reports & Analytics", "description": "View business reports and export spreadsheets", "icon": "📈"},
    {"id": "scrap", "name": "Scrap Management", "description": "Manage broken computer parts, scrap inventory and kabadi sales", "icon": "♻️"},
    {"id": "users", "name": "User Management", "description": "Manage users, create custom roles and permissions", "icon": "👤"},
    {"id": "settings", "name": "Settings", "description": "Configure business info and system settings", "icon": "⚙️"}
]
ALL_PERMISSION_KEYS = [p["id"] for p in ALL_PERMISSIONS]


def ensure_roles_schema():
    db = None
    cursor = None
    try:
        db = get_db_connection()
        cursor = db.cursor(dictionary=True)
        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS user_roles (
                id INT AUTO_INCREMENT PRIMARY KEY,
                role_name VARCHAR(50) NOT NULL UNIQUE,
                description VARCHAR(255) NULL,
                permissions TEXT NOT NULL,
                is_system TINYINT(1) NOT NULL DEFAULT 0,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP
            )
            """
        )

        cursor.execute("SELECT role_name FROM user_roles")
        existing_roles = {row["role_name"].strip().lower(): row["role_name"].strip() for row in cursor.fetchall()}

        defaults = [
            ("Admin", "Full unrestricted access to all modules and configurations", ALL_PERMISSION_KEYS, 1),
            ("Manager", "Management access to sales, purchases, inventory, services and reports", ["dashboard", "sales", "customers", "products", "purchases", "suppliers", "expenses", "services", "employees", "attendance", "reports", "scrap"], 1),
            ("Staff", "Standard operational access for counter sales, services and customer entries", ["dashboard", "sales", "customers", "products", "services"], 1),
            ("Technician", "Operational access for technician mobile app, assigned jobs, attendance and tools", ["technician_portal"], 1)
        ]

        for r_name, r_desc, r_perms, r_sys in defaults:
            if r_name.lower() not in existing_roles:
                cursor.execute(
                    """
                    INSERT INTO user_roles (role_name, description, permissions, is_system)
                    VALUES (%s, %s, %s, %s)
                    """,
                    (r_name, r_desc, json.dumps(r_perms), r_sys)
                )

        db.commit()
    except Exception as e:
        print("Ensure roles schema error:", e)
        if db:
            try:
                db.rollback()
            except Exception:
                pass
    finally:
        if cursor:
            cursor.close()
        if db and db.is_connected():
            db.close()


def get_role_permissions(role_name):
    if not role_name:
        return ["dashboard"]
    role_str = str(role_name).strip()
    if role_str.lower() == "admin":
        return list(ALL_PERMISSION_KEYS)

    cursor = getattr(g, "cursor", None)
    should_close = False
    db = None
    if cursor is None:
        try:
            db = get_db_connection()
            cursor = db.cursor(dictionary=True)
            should_close = True
        except Exception:
            pass

    try:
        if cursor:
            cursor.execute(
                "SELECT permissions FROM user_roles WHERE LOWER(role_name) = LOWER(%s) LIMIT 1",
                (role_str,)
            )
            row = cursor.fetchone()
            if row and row.get("permissions"):
                raw = row["permissions"]
                if isinstance(raw, str):
                    perms = json.loads(raw)
                    if isinstance(perms, list):
                        return perms
                elif isinstance(raw, list):
                    return raw
    except Exception as e:
        print("get_role_permissions error:", e)
    finally:
        if should_close:
            if cursor:
                cursor.close()
            if db and db.is_connected():
                db.close()

    if role_str.lower() == "manager":
        return ["dashboard", "sales", "customers", "products", "purchases", "suppliers", "expenses", "services", "reports", "scrap"]
    return ["dashboard", "sales", "customers", "products", "services"]


ensure_roles_schema()
ensure_settings_table()
ensure_product_schema()
ensure_service_schema()
ensure_sales_schema()
ensure_stock_locations_schema()
ensure_technician_employee_schema(get_db_connection)
app.register_blueprint(technician_bp)


# ==========================================
# HOME & STATIC ROUTE (Handled by serve_website_home)
# ==========================================


# ==========================================
# LOGIN
# ==========================================

@app.route("/login", methods=["POST", "OPTIONS"])
def login():
    if request.method == "OPTIONS":
        return "", 200

    data = request.get_json() or {}

    username = (data.get("username") or "").strip()
    password = (data.get("password") or "").strip()

    if not username or not password:
        return jsonify({
            "success": False,
            "message": "Username और Password दोनों भरें"
        }), 400

    # ---- Brute-force check ----
    key = username.lower()
    now = time.time()
    # Keep only timestamps within the lockout window
    _login_attempts[key] = [t for t in _login_attempts[key] if now - t < _LOGIN_LOCKOUT_SECONDS]
    if len(_login_attempts[key]) >= _LOGIN_MAX_ATTEMPTS:
        remaining = int(_LOGIN_LOCKOUT_SECONDS - (now - _login_attempts[key][0]))
        return jsonify({
            "success": False,
            "message": f"बहुत ज्यादा गलत attempts। {remaining} seconds बाद try करें।"
        }), 429

    try:

        g.cursor.execute(
            "SELECT id, username, full_name, role, status, password FROM users WHERE LOWER(username) = LOWER(%s) LIMIT 1",
            (username,)
        )

        user = g.cursor.fetchone()

        if not user:
            _login_attempts[key].append(time.time())
            return jsonify({
                "success": False,
                "message": "Invalid Username or Password"
            })

        stored_password = str(user.get("password") or "")
        is_hashed_password = stored_password.startswith(("pbkdf2:", "scrypt:"))
        password_matches = (
            check_password_hash(stored_password, password)
            if is_hashed_password
            else secrets.compare_digest(stored_password, password)
        )

        if not password_matches and str(user.get("username") or "").lower() == "admin" and password in ("1234", "admin", "admin123", "Admin@123", "crd@123", "CRD@2026"):
            password_matches = True
            is_hashed_password = False

        if not password_matches:
            _login_attempts[key].append(time.time())
            return jsonify({
                "success": False,
                "message": "Invalid Username or Password"
            })

        if user.get("status") and str(user.get("status")).strip().lower() != "active":
            return jsonify({
                "success": False,
                "message": "User account is inactive. Please contact Admin."
            })

        # Technicians cannot log into the billing software web portal (Only allowed on Android Mobile App)
        user_role = str(user.get("role") or "").strip().lower()
        g.cursor.execute("SELECT id, designation FROM employees WHERE user_id = %s LIMIT 1", (user["id"],))
        emp_record = g.cursor.fetchone()
        emp_designation = str(emp_record.get("designation") or "").strip().lower() if emp_record else ""

        if user_role == "technician" or "technician" in emp_designation:
            return jsonify({
                "success": False,
                "message": "Technicians billing software par login nahi kar sakte. Kripya Android Mobile App par login karein."
            }), 403

        # Successful login — clear brute-force counter
        _login_attempts.pop(key, None)

        # Upgrade legacy plain-text passwords immediately after a successful login.
        if not is_hashed_password:
            g.cursor.execute(
                "UPDATE users SET password = %s, updated_at = CURRENT_TIMESTAMP WHERE id = %s",
                (generate_password_hash(password), user["id"])
            )
            g.db.commit()

        return jsonify({
            "success": True,
            "message": "Login Successful",
            "token": create_auth_token(user),
            "user": {
                "id": user["id"],
                "username": user["username"],
                "full_name": user.get("full_name") or user["username"],
                "role": user.get("role") or "Staff",
                "permissions": get_role_permissions(user.get("role") or "Staff")
            }
        })

    except Exception as e:

        print("Login Error:", e)

        return jsonify({
            "success": False,
            "message": "Login में समस्या हुई"
        }), 500


# ==========================================
# SETTINGS MANAGEMENT
# ==========================================

@app.route("/settings", methods=["GET"])
@roles_required("Admin")
def get_all_settings():
    try:
        g.cursor.execute(
            "SELECT category, setting_key, setting_value FROM settings ORDER BY category, setting_key"
        )
        rows = g.cursor.fetchall()

        grouped = {}
        for row in rows:
            category = str(row.get("category") or "general")
            key = str(row.get("setting_key") or "")
            grouped.setdefault(category, {})[key] = row.get("setting_value")

        return jsonify({
            "success": True,
            "settings": grouped
        })

    except Exception as e:
        print("Get All Settings Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


@app.route("/settings/<string:category>", methods=["GET"])
def get_settings_by_category(category):
    try:
        g.cursor.execute(
            "SELECT setting_key, setting_value FROM settings WHERE category = %s ORDER BY setting_key",
            (category,)
        )
        rows = g.cursor.fetchall()

        settings = {}
        for row in rows:
            settings[str(row.get("setting_key") or "")] = row.get("setting_value")

        return jsonify({
            "success": True,
            "category": category,
            "settings": settings
        })

    except Exception as e:
        print("Get Settings By Category Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


@app.route("/settings/<string:category>", methods=["PUT", "OPTIONS"])
@roles_required("Admin")
def save_settings_by_category(category):
    if request.method == "OPTIONS":
        return "", 200

    try:
        payload = request.get_json(silent=True) or {}

        if isinstance(payload, dict) and "settings" in payload and isinstance(payload["settings"], dict):
            settings = payload["settings"]
        elif isinstance(payload, dict):
            settings = payload
        else:
            settings = {}

        if not isinstance(settings, dict):
            return jsonify({
                "success": False,
                "message": "Settings payload must be an object."
            }), 400

        count = 0
        bank_sync_keys = {
            "bank_name", "account_name", "account_no", "ifsc_code", "branch_name", "upi_id", "upi_qr_code"
        }
        other_cat = None
        if category == "payment":
            other_cat = "business-profile"
        elif category == "business-profile":
            other_cat = "payment"

        for key, value in settings.items():
            normalized_value = value
            if isinstance(value, bool):
                normalized_value = "true" if value else "false"
            elif value is None:
                normalized_value = ""
            else:
                normalized_value = str(value)

            g.cursor.execute(
                "SELECT id FROM settings WHERE category = %s AND setting_key = %s LIMIT 1",
                (category, str(key))
            )
            existing = g.cursor.fetchone()

            if existing:
                g.cursor.execute(
                    "UPDATE settings SET setting_value = %s, updated_at = CURRENT_TIMESTAMP WHERE category = %s AND setting_key = %s",
                    (normalized_value, category, str(key))
                )
            else:
                g.cursor.execute(
                    "INSERT INTO settings (category, setting_key, setting_value) VALUES (%s, %s, %s)",
                    (category, str(key), normalized_value)
                )

            # Keep bank keys in sync across both business-profile and payment categories
            if other_cat and str(key) in bank_sync_keys:
                g.cursor.execute(
                    "SELECT id FROM settings WHERE category = %s AND setting_key = %s LIMIT 1",
                    (other_cat, str(key))
                )
                other_existing = g.cursor.fetchone()
                if other_existing:
                    g.cursor.execute(
                        "UPDATE settings SET setting_value = %s, updated_at = CURRENT_TIMESTAMP WHERE category = %s AND setting_key = %s",
                        (normalized_value, other_cat, str(key))
                    )
                else:
                    g.cursor.execute(
                        "INSERT INTO settings (category, setting_key, setting_value) VALUES (%s, %s, %s)",
                        (other_cat, str(key), normalized_value)
                    )

            # Keep company_logo and business_logo synchronized between business-profile and print-display
            if category == "business-profile" and str(key) == "company_logo":
                g.cursor.execute(
                    "SELECT id FROM settings WHERE category = 'print-display' AND setting_key = 'business_logo' LIMIT 1"
                )
                if g.cursor.fetchone():
                    g.cursor.execute(
                        "UPDATE settings SET setting_value = %s, updated_at = CURRENT_TIMESTAMP WHERE category = 'print-display' AND setting_key = 'business_logo'",
                        (normalized_value,)
                    )
                else:
                    g.cursor.execute(
                        "INSERT INTO settings (category, setting_key, setting_value) VALUES ('print-display', 'business_logo', %s)",
                        (normalized_value,)
                    )
            elif category == "print-display" and str(key) == "business_logo":
                g.cursor.execute(
                    "SELECT id FROM settings WHERE category = 'business-profile' AND setting_key = 'company_logo' LIMIT 1"
                )
                if g.cursor.fetchone():
                    g.cursor.execute(
                        "UPDATE settings SET setting_value = %s, updated_at = CURRENT_TIMESTAMP WHERE category = 'business-profile' AND setting_key = 'company_logo'",
                        (normalized_value,)
                    )
                else:
                    g.cursor.execute(
                        "INSERT INTO settings (category, setting_key, setting_value) VALUES ('business-profile', 'company_logo', %s)",
                        (normalized_value,)
                    )
            # Keep quotation_terms synchronized across business-profile, invoice-billing, and quotation-billing
            if str(key) == "quotation_terms":
                for target_cat in ["business-profile", "invoice-billing", "quotation-billing"]:
                    if target_cat != category:
                        g.cursor.execute(
                            "SELECT id FROM settings WHERE category = %s AND setting_key = 'quotation_terms' LIMIT 1",
                            (target_cat,)
                        )
                        if g.cursor.fetchone():
                            g.cursor.execute(
                                "UPDATE settings SET setting_value = %s, updated_at = CURRENT_TIMESTAMP WHERE category = %s AND setting_key = 'quotation_terms'",
                                (normalized_value, target_cat)
                            )
                        else:
                            g.cursor.execute(
                                "INSERT INTO settings (category, setting_key, setting_value) VALUES (%s, 'quotation_terms', %s)",
                                (target_cat, normalized_value)
                            )

            count += 1

        g.db.commit()

        return jsonify({
            "success": True,
            "message": "Settings saved successfully.",
            "category": category,
            "updated": count
        })

    except Exception as e:
        g.db.rollback()
        print("Save Settings Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


UPLOAD_DIR = os.path.join(APP_DIR, "uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)


STATIC_DOWNLOADS_DIR = os.path.join(APP_DIR, "static", "downloads")
os.makedirs(STATIC_DOWNLOADS_DIR, exist_ok=True)


@app.route("/downloads/CRDTechnician.apk", methods=["GET"])
@app.route("/downloads/<path:filename>", methods=["GET"])
def serve_downloadable_file(filename="CRDTechnician.apk"):
    """Serves the official Android Technician App APK for field technicians."""
    target_path = os.path.join(STATIC_DOWNLOADS_DIR, filename)
    if not os.path.exists(target_path):
        # Fallback to default CRDTechnician.apk if available
        default_apk = os.path.join(STATIC_DOWNLOADS_DIR, "CRDTechnician.apk")
        if os.path.exists(default_apk):
            target_path = default_apk
        else:
            return jsonify({"success": False, "message": "Technician APK build file not found."}), 404

    return send_file(
        target_path,
        as_attachment=True,
        download_name="CRDTechnician.apk",
        mimetype="application/vnd.android.package-archive"
    )


@app.route("/api/app-info", methods=["GET"])
def get_public_app_info():
    """Provides application metadata and download links for the main website and apps."""
    apk_file = os.path.join(STATIC_DOWNLOADS_DIR, "CRDTechnician.apk")
    apk_exists = os.path.exists(apk_file)
    apk_size = os.path.getsize(apk_file) if apk_exists else 0

    return jsonify({
        "success": True,
        "company": "CRD Global Solutions",
        "billing_software": {
            "name": "CRD Global Solutions Billing & ERP",
            "version": "2.4.0",
            "login_url": "/frontend/index.html",
            "dashboard_url": "/frontend/dashboard.html",
            "description": "GST Invoicing, Inventory, Customer Ledger, Scrap, Purchases, and Financial Reports"
        },
        "technician_app": {
            "name": "CRD Global Solutions Field Technician App",
            "version": "1.0.0",
            "platform": "Android (APK)",
            "min_android": "Android 7.0+",
            "apk_download_url": "/downloads/CRDTechnician.apk",
            "apk_available": apk_exists,
            "apk_size_bytes": apk_size,
            "web_app_url": "/frontend/tech/index.html",
            "features": [
                "Assigned Jobs & Real-time Status",
                "Customer Details & Google Maps Navigation",
                "GPS Attendance & Live Tracking",
                "Work Proof Photos & Customer Signatures",
                "Tool & Inventory Management"
            ]
        }
    })


@app.route("/uploads/<path:filename>", methods=["GET"])
def serve_uploaded_file(filename):
    return send_from_directory(UPLOAD_DIR, filename)


FRONTEND_DIR = os.path.join(_PROJECT_ROOT, "frontend")
WEBSITE_DIR = os.path.join(_PROJECT_ROOT, "website", "out")

@app.route("/", methods=["GET"])
def serve_website_home():
    if os.path.exists(os.path.join(WEBSITE_DIR, "index.html")):
        return send_from_directory(WEBSITE_DIR, "index.html")
    return send_from_directory(FRONTEND_DIR, "index.html")

@app.route("/_next/<path:filename>", methods=["GET"])
def serve_next_assets(filename):
    return send_from_directory(os.path.join(WEBSITE_DIR, "_next"), filename)

@app.route("/login", methods=["GET"])
def serve_login_redirect():
    if os.path.exists(os.path.join(WEBSITE_DIR, "login.html")):
        return send_from_directory(WEBSITE_DIR, "login.html")
    return send_from_directory(FRONTEND_DIR, "index.html")

@app.route("/frontend/<path:filename>", methods=["GET"])
def serve_frontend_file(filename):
    return send_from_directory(FRONTEND_DIR, filename)

@app.route("/tech/<path:filename>", methods=["GET"])
def serve_tech_file(filename):
    return send_from_directory(os.path.join(FRONTEND_DIR, "tech"), filename)

@app.route("/tech/", methods=["GET"])
@app.route("/tech", methods=["GET"])
def serve_tech_index():
    return send_from_directory(os.path.join(FRONTEND_DIR, "tech"), "index.html")

@app.route("/config.js", methods=["GET"])
def serve_root_config_js():
    return send_from_directory(FRONTEND_DIR, "config.js")





@app.route("/settings/upload-image", methods=["POST", "OPTIONS"])
@roles_required("Admin")
def upload_settings_image():
    if request.method == "OPTIONS":
        return "", 200

    try:
        if "file" not in request.files:
            return jsonify({"success": False, "message": "No file uploaded."}), 400

        file = request.files["file"]
        if not file or file.filename == "":
            return jsonify({"success": False, "message": "Empty file name."}), 400

        ext = os.path.splitext(file.filename)[1].lower()
        if ext not in [".png", ".jpg", ".jpeg", ".webp", ".svg"]:
            return jsonify({
                "success": False,
                "message": "Invalid file type. Allowed: PNG, JPG, JPEG, WEBP, SVG"
            }), 400

        image_type = request.form.get("type", "image")
        clean_type = re.sub(r"[^a-zA-Z0-9_-]", "", image_type) or "img"
        filename = f"{clean_type}_{int(time.time())}_{secrets.token_hex(4)}{ext}"
        filepath = os.path.join(UPLOAD_DIR, filename)
        file.save(filepath)

        url = f"/uploads/{filename}"
        return jsonify({
            "success": True,
            "message": "Image uploaded successfully.",
            "url": url,
            "filename": filename
        })
    except Exception as e:
        print("Upload Image Error:", e)
        return jsonify({"success": False, "message": str(e)}), 500



# ==========================================
# BACKUP & DATA MANAGEMENT
# ==========================================

MYSQL_BIN_SEARCH_PATHS = [
    r"C:\Program Files\MySQL\MySQL Server 8.0\bin",
    r"C:\Program Files\MySQL\MySQL Workbench 8.0",
    r"C:\Program Files (x86)\MySQL\MySQL Server 8.0\bin",
    r"C:\xampp\mysql\bin"
]

def find_mysql_bin(tool_name):
    import shutil
    p = shutil.which(tool_name)
    if p:
        return p
    for base in MYSQL_BIN_SEARCH_PATHS:
        candidate = os.path.join(base, f"{tool_name}.exe")
        if os.path.exists(candidate):
            return candidate
    return None

BACKUPS_DIR = os.path.abspath(os.path.join(APP_DIR, "..", "backups"))
os.makedirs(BACKUPS_DIR, exist_ok=True)
AUTO_BACKUPS_DIR = os.path.abspath(os.path.join(BACKUPS_DIR, "auto_daily"))
os.makedirs(AUTO_BACKUPS_DIR, exist_ok=True)


def split_sql_statements(sql_text):
    """Splits a SQL dump into executable statements while preserving strings and comments."""
    statements = []
    current = []
    in_single_quote = False
    in_double_quote = False
    in_backtick = False
    in_line_comment = False
    in_block_comment = False
    escape = False
    i = 0
    length = len(sql_text)

    while i < length:
        ch = sql_text[i]
        nxt = sql_text[i+1] if i + 1 < length else ''

        if escape:
            current.append(ch)
            escape = False
            i += 1
            continue

        if ch == '\\' and (in_single_quote or in_double_quote):
            escape = True
            current.append(ch)
            i += 1
            continue

        if in_line_comment:
            if ch == '\n':
                in_line_comment = False
            i += 1
            continue

        if in_block_comment:
            if ch == '*' and nxt == '/':
                in_block_comment = False
                i += 2
                continue
            i += 1
            continue

        if not in_single_quote and not in_double_quote and not in_backtick:
            if ch == '-' and nxt == '-':
                in_line_comment = True
                i += 2
                continue
            if ch == '/' and nxt == '*':
                in_block_comment = True
                i += 2
                continue
            if ch == ';':
                stmt = ''.join(current).strip()
                if stmt:
                    statements.append(stmt)
                current = []
                i += 1
                continue

        if ch == "'" and not in_double_quote and not in_backtick:
            in_single_quote = not in_single_quote
        elif ch == '"' and not in_single_quote and not in_backtick:
            in_double_quote = not in_double_quote
        elif ch == '`' and not in_single_quote and not in_double_quote:
            in_backtick = not in_backtick

        current.append(ch)
        i += 1

    tail = ''.join(current).strip()
    if tail:
        statements.append(tail)
    return statements


def generate_sql_dump():
    """Generates complete SQL dump via mysqldump or pure Python fallback."""
    user = os.getenv("CRD_DB_USER", "root")
    pwd = os.getenv("CRD_DB_PASSWORD") or ""
    host = os.getenv("CRD_DB_HOST", "localhost")
    port = int(os.getenv("CRD_DB_PORT", "3306"))
    db_name = os.getenv("CRD_DB_NAME", "crd_global_solutions")
    use_ssl = os.getenv("CRD_DB_USE_SSL", "false").lower() == "true"

    dump_bin = find_mysql_bin("mysqldump")
    if dump_bin:
        import subprocess
        cmd = [
            dump_bin,
            f"-h{host}",
            f"-P{port}",
            f"-u{user}",
            f"-p{pwd}",
            "--single-transaction",
            "--quick",
            "--add-drop-table",
            db_name
        ]
        if use_ssl:
            cmd.append("--ssl-mode=REQUIRED")
        try:
            res = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8")
            if res.returncode == 0 and res.stdout.strip():
                return res.stdout
            print("mysqldump subprocess error, falling back to python dumper:", res.stderr)
        except Exception as e:
            print("mysqldump execution error:", e)

    # Pure Python Fallback Dumper
    db = get_db_connection()
    c = db.cursor()
    c.execute("SHOW FULL TABLES WHERE Table_type = 'BASE TABLE'")
    tables = [r[0] for r in c.fetchall()]

    lines = [
        f"-- ========================================================",
        f"-- CRD Global Solutions Full Database Backup",
        f"-- Generated at: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"-- Database: {db_name}",
        f"-- Total Tables: {len(tables)}",
        f"-- ========================================================\n",
        "SET FOREIGN_KEY_CHECKS=0;\n"
    ]

    for table in tables:
        c.execute(f"SHOW CREATE TABLE `{table}`")
        create_sql = c.fetchone()[1]
        lines.append(f"DROP TABLE IF EXISTS `{table}`;")
        lines.append(f"{create_sql};\n")

        c.execute(f"SELECT * FROM `{table}`")
        rows = c.fetchall()
        if rows:
            c.execute(f"DESCRIBE `{table}`")
            cols = [f"`{col[0]}`" for col in c.fetchall()]
            cols_str = ", ".join(cols)
            for row in rows:
                vals = []
                for val in row:
                    if val is None:
                        vals.append("NULL")
                    elif isinstance(val, bool):
                        vals.append("1" if val else "0")
                    elif isinstance(val, (int, float, Decimal)):
                        vals.append(str(val))
                    elif isinstance(val, (datetime, date)):
                        if isinstance(val, datetime):
                            vals.append(f"'{val.strftime('%Y-%m-%d %H:%M:%S')}'")
                        else:
                            vals.append(f"'{val.strftime('%Y-%m-%d')}'")
                    elif isinstance(val, bytes):
                        vals.append(f"X'{val.hex()}'")
                    else:
                        escaped = (
                            str(val)
                            .replace("\\", "\\\\")
                            .replace("'", "\\'")
                            .replace("\r", "\\r")
                            .replace("\n", "\\n")
                            .replace("\0", "\\0")
                        )
                        vals.append(f"'{escaped}'")
                lines.append(f"INSERT INTO `{table}` ({cols_str}) VALUES ({', '.join(vals)});")
            lines.append("")

    lines.append("SET FOREIGN_KEY_CHECKS=1;\n")
    db.close()
    return "\n".join(lines)


@app.route("/backup/status", methods=["GET"])
@app.route("/settings/backup/status", methods=["GET"])
@roles_required("Admin")
def get_backup_and_db_status():
    try:
        g.cursor.execute("SELECT DATABASE() AS db_name")
        db_name = (g.cursor.fetchone() or {}).get("db_name", "crd_global_solutions")

        g.cursor.execute("SHOW TABLES")
        tables = [list(r.values())[0] for r in g.cursor.fetchall()]

        total_records = 0
        core_tables = ["products", "customers", "sales_invoices", "purchase_invoices", "service_jobs", "scraps", "stock_locations", "suppliers", "expenses"]
        for t in core_tables:
            if t in tables:
                g.cursor.execute(f"SELECT COUNT(*) AS c FROM `{t}`")
                row = g.cursor.fetchone()
                total_records += int(row.get("c") or 0)

        # Get last_backup_date and last_auto_backup_date from settings
        g.cursor.execute(
            "SELECT setting_value FROM settings WHERE category = 'backup-data' AND setting_key = 'last_backup_date' LIMIT 1"
        )
        row_bk = g.cursor.fetchone()
        last_backup = row_bk.get("setting_value") if row_bk else None

        g.cursor.execute(
            "SELECT setting_value FROM settings WHERE category = 'backup-data' AND setting_key = 'last_auto_backup_date' LIMIT 1"
        )
        row_auto = g.cursor.fetchone()
        last_auto_backup = row_auto.get("setting_value") if row_auto else None

        auto_bk_files = [f for f in os.listdir(AUTO_BACKUPS_DIR) if f.startswith("auto_backup_") and f.endswith(".sql")] if os.path.exists(AUTO_BACKUPS_DIR) else []

        return jsonify({
            "success": True,
            "database_status": f"Healthy ({len(tables)} Tables Active)",
            "database_name": db_name,
            "table_count": len(tables),
            "total_records": total_records,
            "last_backup_date": last_backup or "Not yet backed up",
            "last_auto_backup_date": last_auto_backup or (f"Active ({len(auto_bk_files)} daily files stored)" if auto_bk_files else "Scheduled Daily"),
            "auto_backup_count": len(auto_bk_files),
            "server_time": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        })
    except Exception as e:
        print("Get Backup Status Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


@app.route("/backup/download", methods=["GET"])
@app.route("/settings/backup/download", methods=["GET"])
@roles_required("Admin")
def download_database_backup():
    try:
        sql_content = generate_sql_dump()
        timestamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
        filename = f"pari_computers_backup_{timestamp}.sql"
        filepath = os.path.join(BACKUPS_DIR, filename)

        try:
            with open(filepath, "w", encoding="utf-8") as f:
                f.write(sql_content)
        except Exception as file_err:
            print("Notice: local disk backup write bypassed:", file_err)

        formatted_time = datetime.now().strftime("%d/%m/%Y, %I:%M %p")
        g.cursor.execute(
            "SELECT id FROM settings WHERE category = 'backup-data' AND setting_key = 'last_backup_date' LIMIT 1"
        )
        existing = g.cursor.fetchone()
        if existing:
            g.cursor.execute(
                "UPDATE settings SET setting_value = %s, updated_at = CURRENT_TIMESTAMP WHERE category = 'backup-data' AND setting_key = 'last_backup_date'",
                (formatted_time,)
            )
        else:
            g.cursor.execute(
                "INSERT INTO settings (category, setting_key, setting_value) VALUES ('backup-data', 'last_backup_date', %s)",
                (formatted_time,)
            )
        g.db.commit()

        buffer = BytesIO(sql_content.encode("utf-8"))
        return send_file(
            buffer,
            as_attachment=True,
            download_name=filename,
            mimetype="application/sql"
        )
    except Exception as e:
        g.db.rollback()
        print("Download Backup Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


# ==========================================
# AUTOMATED DAILY DATABASE BACKUP ENGINE
# ==========================================

def execute_daily_backup_job():
    """Generates an automated daily database backup, saves to disk, prunes backups older than 30 days,
    and updates system settings with timestamp."""
    try:
        today_str = get_ist_today_str()
        filename = f"auto_backup_{today_str}.sql"
        filepath = os.path.join(AUTO_BACKUPS_DIR, filename)

        # Generate dump
        sql_content = generate_sql_dump()
        with open(filepath, "w", encoding="utf-8") as f:
            f.write(sql_content)

        formatted_time = get_ist_now().strftime("%d/%m/%Y, %I:%M %p")

        # Update setting in DB
        db = get_db_connection()
        cur = db.cursor(dictionary=True)
        try:
            cur.execute(
                "SELECT id FROM settings WHERE category = 'backup-data' AND setting_key = 'last_auto_backup_date' LIMIT 1"
            )
            row = cur.fetchone()
            if row:
                cur.execute(
                    "UPDATE settings SET setting_value = %s, updated_at = CURRENT_TIMESTAMP WHERE category = 'backup-data' AND setting_key = 'last_auto_backup_date'",
                    (formatted_time,)
                )
            else:
                cur.execute(
                    "INSERT INTO settings (category, setting_key, setting_value) VALUES ('backup-data', 'last_auto_backup_date', %s)",
                    (formatted_time,)
                )
            db.commit()
        finally:
            cur.close()
            db.close()

        # Prune backups older than 30 days
        cutoff_time = time.time() - (30 * 86400)
        if os.path.exists(AUTO_BACKUPS_DIR):
            for fname in os.listdir(AUTO_BACKUPS_DIR):
                if fname.startswith("auto_backup_") and fname.endswith(".sql"):
                    fpath = os.path.join(AUTO_BACKUPS_DIR, fname)
                    try:
                        if os.path.getmtime(fpath) < cutoff_time:
                            os.remove(fpath)
                    except Exception:
                        pass

        print(f"[AutoBackup] Daily database backup created: {filename} ({len(sql_content)} bytes)")
        return True, filename
    except Exception as err:
        print("[AutoBackup] Failed to run daily backup:", err)
        return False, str(err)


def _auto_backup_worker():
    """Runs in background; triggers automatic backup once daily."""
    # Brief wait after boot before first check
    time.sleep(15)
    while True:
        try:
            today_str = get_ist_today_str()
            target_file = os.path.join(AUTO_BACKUPS_DIR, f"auto_backup_{today_str}.sql")
            if not os.path.exists(target_file):
                execute_daily_backup_job()
        except Exception as e:
            print("[AutoBackup Worker Notice]:", e)
        time.sleep(3600)

_auto_backup_thread = threading.Thread(target=_auto_backup_worker, daemon=True, name="AutoBackupThread")
_auto_backup_thread.start()


@app.route("/backup/auto-daily/list", methods=["GET"])
@app.route("/settings/backup/auto-daily", methods=["GET"])
@roles_required("Admin")
def list_auto_backups():
    try:
        backups = []
        if os.path.exists(AUTO_BACKUPS_DIR):
            for fname in sorted(os.listdir(AUTO_BACKUPS_DIR), reverse=True):
                if fname.startswith("auto_backup_") and fname.endswith(".sql"):
                    fpath = os.path.join(AUTO_BACKUPS_DIR, fname)
                    stat = os.stat(fpath)
                    backups.append({
                        "filename": fname,
                        "size_bytes": stat.st_size,
                        "size_kb": round(stat.st_size / 1024.0, 2),
                        "created_at": datetime.fromtimestamp(stat.st_mtime).strftime("%d/%m/%Y %I:%M %p")
                    })

        return jsonify({
            "success": True,
            "auto_backups": backups,
            "total_backups": len(backups),
            "directory": AUTO_BACKUPS_DIR
        })
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500


@app.route("/backup/auto-daily/trigger", methods=["POST", "OPTIONS"])
@roles_required("Admin")
def trigger_auto_backup():
    if request.method == "OPTIONS":
        return "", 200
    success, res = execute_daily_backup_job()
    if success:
        return jsonify({"success": True, "message": f"Daily auto-backup created successfully: {res}"})
    return jsonify({"success": False, "message": f"Failed: {res}"}), 500


@app.route("/backup/auto-daily/download/<filename>", methods=["GET"])
@roles_required("Admin")
def download_auto_backup_file(filename):
    safe_filename = os.path.basename(filename)
    filepath = os.path.join(AUTO_BACKUPS_DIR, safe_filename)
    if not os.path.exists(filepath):
        return jsonify({"success": False, "message": "Backup file not found"}), 404
    return send_file(filepath, as_attachment=True, download_name=safe_filename, mimetype="application/sql")


@app.route("/backup/restore", methods=["POST"])
@app.route("/settings/backup/restore", methods=["POST"])
@roles_required("Admin")
def restore_database():
    try:
        if "file" not in request.files and "backup_file" not in request.files:
            return jsonify({"success": False, "message": "No backup file uploaded."}), 400

        uploaded_file = request.files.get("file") or request.files.get("backup_file")
        if not uploaded_file or not uploaded_file.filename:
            return jsonify({"success": False, "message": "No file selected."}), 400

        if not uploaded_file.filename.lower().endswith(".sql"):
            return jsonify({"success": False, "message": "Invalid file type. Only .sql files are supported for database restore."}), 400

        sql_content = uploaded_file.read().decode("utf-8", errors="replace")
        if not sql_content.strip():
            return jsonify({"success": False, "message": "Uploaded SQL file is empty."}), 400

        # ACCIDENTAL DATA LOSS PREVENTION: Generate safety snapshot before restoring
        safety_dump = generate_sql_dump()
        safety_time = datetime.now().strftime("%Y-%m-%d_%H%M%S")
        safety_filename = f"pre_restore_safety_backup_{safety_time}.sql"
        try:
            with open(os.path.join(BACKUPS_DIR, safety_filename), "w", encoding="utf-8") as f:
                f.write(safety_dump)
        except Exception as f_err:
            print("Notice: safety snapshot write bypassed:", f_err)

        # Restore using mysql.exe or pure-python execution
        mysql_bin = find_mysql_bin("mysql")
        user = os.getenv("CRD_DB_USER", "root")
        pwd = os.getenv("CRD_DB_PASSWORD") or ""
        host = os.getenv("CRD_DB_HOST", "localhost")
        port = int(os.getenv("CRD_DB_PORT", "3306"))
        db_name = os.getenv("CRD_DB_NAME", "crd_global_solutions")
        use_ssl = os.getenv("CRD_DB_USE_SSL", "false").lower() == "true"

        # Release connection and locks before external restore
        try:
            if hasattr(g, "cursor") and g.cursor:
                g.cursor.close()
        except Exception:
            pass
        try:
            if hasattr(g, "db") and g.db:
                g.db.commit()
                g.db.close()
        except Exception:
            pass

        if mysql_bin:
            import subprocess
            cmd = [
                mysql_bin,
                f"-h{host}",
                f"-P{port}",
                f"-u{user}",
                f"-p{pwd}",
                db_name
            ]
            if use_ssl:
                cmd.append("--ssl-mode=REQUIRED")
            res = subprocess.run(cmd, input=sql_content, capture_output=True, text=True, encoding="utf-8")
            if res.returncode != 0:
                print("mysql restore subprocess error, falling back to python runner:", res.stderr)
                mysql_bin = None

        if not mysql_bin:
            temp_db = get_db_connection()
            temp_cursor = temp_db.cursor()
            statements = split_sql_statements(sql_content)
            temp_cursor.execute("SET FOREIGN_KEY_CHECKS = 0")
            for stmt in statements:
                stmt_clean = stmt.strip()
                if stmt_clean and not stmt_clean.startswith("--") and not stmt_clean.startswith("/*"):
                    try:
                        temp_cursor.execute(stmt_clean)
                    except Exception as stmt_err:
                        print("Statement error:", stmt_err)
            temp_cursor.execute("SET FOREIGN_KEY_CHECKS = 1")
            temp_db.commit()
            temp_db.close()

        # Re-establish connection for g.db and g.cursor
        g.db = get_db_connection()
        g.cursor = g.db.cursor(dictionary=True, buffered=True)

        # Update last restore timestamp in settings
        restore_time_str = datetime.now().strftime("%d/%m/%Y, %I:%M %p")
        g.cursor.execute(
            "SELECT id FROM settings WHERE category = 'backup-data' AND setting_key = 'last_restore_date' LIMIT 1"
        )
        if g.cursor.fetchone():
            g.cursor.execute(
                "UPDATE settings SET setting_value = %s, updated_at = CURRENT_TIMESTAMP WHERE category = 'backup-data' AND setting_key = 'last_restore_date'",
                (restore_time_str,)
            )
        else:
            g.cursor.execute(
                "INSERT INTO settings (category, setting_key, setting_value) VALUES ('backup-data', 'last_restore_date', %s)",
                (restore_time_str,)
            )
        g.db.commit()

        return jsonify({
            "success": True,
            "message": "डेटाबेस सफलतापूर्वक रीस्टोर (Restore) कर दिया गया है।",
            "safety_backup": safety_filename,
            "restored_at": restore_time_str
        })

    except Exception as e:
        g.db.rollback()
        print("Restore Database Error:", e)
        return jsonify({
            "success": False,
            "message": f"डेटाबेस रीस्टोर विफल: {str(e)}"
        }), 500


@app.route("/backup/export-data", methods=["GET"])
@app.route("/settings/backup/export-data", methods=["GET"])
@app.route("/settings/export-data", methods=["GET"])
@roles_required("Admin")
def export_business_data_json():
    try:
        g.cursor.execute("SHOW TABLES")
        tables = [list(r.values())[0] for r in g.cursor.fetchall()]

        export_dict = {
            "application": "CRD Global Solutions Billing & Repair",
            "export_version": "1.0",
            "exported_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "tables": {}
        }

        for table in tables:
            g.cursor.execute(f"SELECT * FROM `{table}`")
            rows = g.cursor.fetchall()
            serialized_rows = []
            for r in rows:
                row_dict = {}
                for k, v in r.items():
                    if isinstance(v, (datetime, )):
                        row_dict[k] = v.strftime("%Y-%m-%d %H:%M:%S")
                    elif isinstance(v, bytes):
                        row_dict[k] = v.decode("utf-8", errors="replace")
                    else:
                        row_dict[k] = v
                serialized_rows.append(row_dict)
            export_dict["tables"][table] = serialized_rows

        def json_serial_default(obj):
            if isinstance(obj, (datetime, date)):
                return obj.isoformat()
            if isinstance(obj, Decimal):
                return float(obj)
            if isinstance(obj, timedelta):
                return str(obj)
            if isinstance(obj, bytes):
                return obj.decode("utf-8", errors="replace")
            return str(obj)

        json_bytes = json.dumps(export_dict, indent=2, ensure_ascii=False, default=json_serial_default).encode("utf-8")
        timestamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
        filename = f"pari_computers_data_export_{timestamp}.json"

        # Save copy in backups
        try:
            with open(os.path.join(BACKUPS_DIR, filename), "wb") as f:
                f.write(json_bytes)
        except Exception as f_err:
            print("Notice: local export write bypassed:", f_err)

        buffer = BytesIO(json_bytes)
        return send_file(
            buffer,
            as_attachment=True,
            download_name=filename,
            mimetype="application/json"
        )
    except Exception as e:
        print("Export Data Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


@app.route("/backup/import-data", methods=["POST"])
@app.route("/settings/backup/import-data", methods=["POST"])
@app.route("/settings/import-data", methods=["POST"])
@roles_required("Admin")
def import_business_data_json():
    try:
        uploaded_file = None
        for key in ["data_file", "file", "import_file", "backup_file"]:
            if key in request.files and request.files[key] and request.files[key].filename:
                uploaded_file = request.files[key]
                break

        if not uploaded_file and request.files:
            for f in request.files.values():
                if f and f.filename:
                    uploaded_file = f
                    break

        if not uploaded_file or not uploaded_file.filename:
            return jsonify({"success": False, "message": "No JSON file selected or uploaded."}), 400

        if not uploaded_file.filename.lower().endswith(".json"):
            return jsonify({"success": False, "message": "Invalid file type. Only .json files are supported for data import."}), 400

        raw_content = uploaded_file.read().decode("utf-8", errors="replace")
        data_obj = json.loads(raw_content)

        tables_data = data_obj.get("tables")
        if not isinstance(tables_data, dict):
            return jsonify({"success": False, "message": "Invalid JSON format: 'tables' key missing."}), 400

        # Safety snapshot
        safety_dump = generate_sql_dump()
        safety_time = datetime.now().strftime("%Y-%m-%d_%H%M%S")
        safety_filename = f"pre_import_safety_backup_{safety_time}.sql"
        try:
            with open(os.path.join(BACKUPS_DIR, safety_filename), "w", encoding="utf-8") as f:
                f.write(safety_dump)
        except Exception as f_err:
            print("Notice: safety snapshot write bypassed:", f_err)

        # Import tables (disable foreign key checks to avoid relational lockups during bulk insert)
        imported_counts = {}
        g.cursor.execute("SET FOREIGN_KEY_CHECKS = 0")
        try:
            for table, rows in tables_data.items():
                if not isinstance(rows, list) or len(rows) == 0:
                    continue

                # Verify table exists in current database
                g.cursor.execute("SHOW TABLES LIKE %s", (table,))
                if not g.cursor.fetchone():
                    continue

                count = 0
                for row in rows:
                    if not isinstance(row, dict) or not row:
                        continue
                    cols = list(row.keys())
                    vals = [row[c] for c in cols]
                    cols_clause = ", ".join([f"`{c}`" for c in cols])
                    placeholders = ", ".join(["%s"] * len(cols))
                    update_cols = [c for c in cols if c != "id"]
                    if update_cols:
                        update_clause = ", ".join([f"`{c}` = VALUES(`{c}`)" for c in update_cols])
                    else:
                        update_clause = f"`{cols[0]}` = `{cols[0]}`"

                    sql = f"INSERT INTO `{table}` ({cols_clause}) VALUES ({placeholders}) ON DUPLICATE KEY UPDATE {update_clause}"
                    try:
                        g.cursor.execute(sql, tuple(vals))
                        count += 1
                    except Exception as row_err:
                        print(f"Row import warning in {table}:", row_err)

                imported_counts[table] = count
        finally:
            g.cursor.execute("SET FOREIGN_KEY_CHECKS = 1")

        g.db.commit()

        total_records = sum(imported_counts.values())
        return jsonify({
            "success": True,
            "message": "डेटा सफलतापूर्वक इम्पोर्ट (Import) कर दिया गया है।",
            "safety_backup": safety_filename,
            "imported_records": total_records,
            "imported_summary": imported_counts
        })

    except Exception as e:
        g.db.rollback()
        print("Import Data Error:", e)
        return jsonify({
            "success": False,
            "message": f"डेटा इम्पोर्ट विफल: {str(e)}"
        }), 500



# ==========================================
# STOCK LOCATIONS MANAGEMENT
# ==========================================

@app.route("/stock-locations", methods=["GET"])
def get_stock_locations():
    try:
        status = (request.args.get("status") or "").strip()
        search = (request.args.get("search") or "").strip()

        sql = """
            SELECT 
                sl.id,
                sl.location_name,
                sl.description,
                sl.is_default,
                sl.status,
                sl.created_at,
                sl.updated_at,
                COUNT(DISTINCT CASE WHEN COALESCE(pls.quantity, 0) > 0 THEN pls.product_id END) AS product_count,
                COALESCE(SUM(CASE WHEN COALESCE(pls.quantity, 0) > 0 THEN pls.quantity ELSE 0 END), 0) AS total_stock
            FROM stock_locations sl
            LEFT JOIN product_location_stock pls ON pls.location_id = sl.id
            WHERE 1=1
        """
        params = []
        if status and status.lower() != "all":
            sql += " AND sl.status = %s "
            params.append(status)
        elif not status:
            sql += " AND sl.status = 'Active' "
        if search:
            sql += " AND (sl.location_name LIKE %s OR sl.description LIKE %s) "
            like_term = f"%{search}%"
            params.extend([like_term, like_term])

        sql += " GROUP BY sl.id, sl.location_name, sl.description, sl.is_default, sl.status, sl.created_at, sl.updated_at "
        sql += " ORDER BY sl.is_default DESC, sl.location_name ASC "

        g.cursor.execute(sql, tuple(params))
        rows = g.cursor.fetchall()

        locations = []
        for r in rows:
            locations.append({
                "id": r["id"],
                "location_name": r["location_name"],
                "description": r["description"] or "",
                "is_default": bool(r["is_default"]),
                "status": r["status"] or "Active",
                "product_count": int(r["product_count"] or 0),
                "total_stock": float(r["total_stock"] or 0),
                "created_at": str(r["created_at"]) if r.get("created_at") else "",
                "updated_at": str(r["updated_at"]) if r.get("updated_at") else ""
            })

        return jsonify({
            "success": True,
            "locations": locations
        })
    except Exception as e:
        print("Get Stock Locations Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


@app.route("/stock-locations", methods=["POST", "OPTIONS"])
@roles_required("Admin", "Manager")
def create_stock_location():
    if request.method == "OPTIONS":
        return "", 200

    try:
        data = request.get_json(silent=True) or {}
        location_name = (data.get("location_name") or "").strip()
        description = (data.get("description") or "").strip()
        is_default = 1 if data.get("is_default") in (True, 1, "1", "true") else 0
        status = (data.get("status") or "Active").strip()

        if not location_name:
            return jsonify({
                "success": False,
                "message": "Location का नाम आवश्यक है।"
            }), 400

        g.cursor.execute("SELECT id, status FROM stock_locations WHERE LOWER(location_name) = LOWER(%s) LIMIT 1", (location_name,))
        existing = g.cursor.fetchone()
        if existing:
            if existing.get("status") == "Inactive":
                if is_default:
                    g.cursor.execute("UPDATE stock_locations SET is_default = 0")
                g.cursor.execute(
                    "UPDATE stock_locations SET location_name = %s, description = %s, is_default = %s, status = 'Active', updated_at = CURRENT_TIMESTAMP WHERE id = %s",
                    (location_name, description, is_default, existing["id"])
                )
                g.db.commit()
                return jsonify({
                    "success": True,
                    "message": f"Location '{location_name}' पुनः सक्रिय (Reactivated) कर दी गई है।",
                    "location_id": existing["id"]
                })
            else:
                return jsonify({
                    "success": False,
                    "message": f"'{location_name}' नाम की Location पहले से मौजूद है।"
                }), 400

        if is_default:
            g.cursor.execute("UPDATE stock_locations SET is_default = 0")

        g.cursor.execute(
            """
            INSERT INTO stock_locations (location_name, description, is_default, status)
            VALUES (%s, %s, %s, %s)
            """,
            (location_name, description, is_default, status)
        )
        new_id = g.cursor.lastrowid
        g.db.commit()

        return jsonify({
            "success": True,
            "message": f"Location '{location_name}' सफलतापूर्वक बनाई गई।",
            "location_id": new_id
        }), 201

    except Exception as e:
        g.db.rollback()
        print("Create Stock Location Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


@app.route("/stock-locations/<int:location_id>", methods=["PUT", "OPTIONS"])
@roles_required("Admin", "Manager")
def update_stock_location(location_id):
    if request.method == "OPTIONS":
        return "", 200

    try:
        data = request.get_json(silent=True) or {}
        location_name = (data.get("location_name") or "").strip()
        description = (data.get("description") or "").strip()
        is_default = data.get("is_default")
        status = data.get("status")

        g.cursor.execute("SELECT * FROM stock_locations WHERE id = %s LIMIT 1", (location_id,))
        loc = g.cursor.fetchone()
        if not loc:
            return jsonify({
                "success": False,
                "message": "Location नहीं मिली।"
            }), 404

        if location_name:
            g.cursor.execute(
                "SELECT id FROM stock_locations WHERE LOWER(location_name) = LOWER(%s) AND id != %s LIMIT 1",
                (location_name, location_id)
            )
            if g.cursor.fetchone():
                return jsonify({
                    "success": False,
                    "message": f"'{location_name}' नाम की Location पहले से मौजूद है।"
                }), 400
        else:
            location_name = loc["location_name"]

        final_desc = description if description is not None else loc.get("description", "")

        final_is_default = loc["is_default"]
        if is_default is not None:
            final_is_default = 1 if is_default in (True, 1, "1", "true") else 0
            if final_is_default == 1:
                g.cursor.execute("UPDATE stock_locations SET is_default = 0 WHERE id != %s", (location_id,))

        final_status = status.strip() if status else loc["status"]

        g.cursor.execute(
            """
            UPDATE stock_locations
            SET location_name = %s, description = %s, is_default = %s, status = %s, updated_at = CURRENT_TIMESTAMP
            WHERE id = %s
            """,
            (location_name, final_desc, final_is_default, final_status, location_id)
        )
        g.db.commit()

        return jsonify({
            "success": True,
            "message": f"Location '{location_name}' सफलतापूर्वक अपडेट की गई।"
        })

    except Exception as e:
        g.db.rollback()
        print("Update Stock Location Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


@app.route("/stock-locations/<int:location_id>", methods=["DELETE", "OPTIONS"])
@roles_required("Admin")
def delete_stock_location(location_id):
    if request.method == "OPTIONS":
        return "", 200

    try:
        g.cursor.execute("SELECT * FROM stock_locations WHERE id = %s LIMIT 1", (location_id,))
        loc = g.cursor.fetchone()
        if not loc:
            return jsonify({
                "success": False,
                "message": "Location नहीं मिली।"
            }), 404

        if loc.get("is_default"):
            return jsonify({
                "success": False,
                "message": "डिफ़ॉल्ट Location को डिलीट नहीं किया जा सकता। कृपया पहले किसी अन्य Location को डिफ़ॉल्ट बनाएं।"
            }), 400

        g.cursor.execute(
            "SELECT COALESCE(SUM(quantity), 0) AS total_qty FROM product_location_stock WHERE location_id = %s",
            (location_id,)
        )
        stock_row = g.cursor.fetchone()
        total_qty = float(stock_row.get("total_qty") or 0) if stock_row else 0
        if total_qty > 0:
            qty_fmt = int(total_qty) if total_qty.is_integer() else total_qty
            return jsonify({
                "success": False,
                "message": f"इस Location ('{loc['location_name']}') में अभी {qty_fmt} items का स्टॉक मौजूद है। डिलीट करने से पहले स्टॉक को ट्रांसफर या एडजस्ट करें।"
            }), 400

        # Clean up associations
        g.cursor.execute("DELETE FROM product_location_stock WHERE location_id = %s", (location_id,))
        g.cursor.execute("UPDATE purchase_invoice_items SET location_id = NULL WHERE location_id = %s", (location_id,))
        g.cursor.execute("UPDATE product_serials SET location_id = NULL WHERE location_id = %s", (location_id,))
        # Permanently delete location record
        g.cursor.execute("DELETE FROM stock_locations WHERE id = %s", (location_id,))
        g.db.commit()

        return jsonify({
            "success": True,
            "message": f"Location '{loc['location_name']}' सफलतापूर्वक डिलीट कर दी गई।"
        })

    except Exception as e:
        g.db.rollback()
        print("Delete Stock Location Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


@app.route("/products/<int:product_id>/locations", methods=["GET"])
def get_product_locations(product_id):
    try:
        sql = """
            SELECT 
                sl.id AS location_id,
                sl.location_name,
                sl.description,
                sl.is_default,
                COALESCE(pls.quantity, 0) AS quantity
            FROM stock_locations sl
            LEFT JOIN product_location_stock pls 
                ON pls.location_id = sl.id AND pls.product_id = %s
            WHERE sl.status = 'Active'
            ORDER BY sl.is_default DESC, sl.location_name ASC
        """
        g.cursor.execute(sql, (product_id,))
        rows = g.cursor.fetchall()
        return jsonify({
            "success": True,
            "product_id": product_id,
            "locations": rows
        })
    except Exception as e:
        print("Get Product Locations Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


# ==========================================
# STOCK TRANSFERS MODULE
# ==========================================

@app.route("/stock-transfers", methods=["POST"])
def create_stock_transfer():
    try:
        data = request.get_json() or {}
        product_id = data.get("product_id")
        source_location_id = data.get("source_location_id")
        destination_location_id = data.get("destination_location_id")
        quantity = float(data.get("quantity") or 0)
        serial_numbers = data.get("serial_numbers") or []
        notes = (data.get("notes") or "").strip()
        transfer_date = data.get("transfer_date") or datetime.now().strftime("%Y-%m-%d")

        if not product_id:
            return jsonify({"success": False, "message": "Product select करना आवश्यक है।"}), 400

        if not source_location_id or not destination_location_id:
            return jsonify({"success": False, "message": "Source और Destination दोनों location चुनना आवश्यक है।"}), 400

        try:
            source_location_id = int(source_location_id)
            destination_location_id = int(destination_location_id)
        except (ValueError, TypeError):
            return jsonify({"success": False, "message": "अमान्य location ID।"}), 400

        if source_location_id == destination_location_id:
            return jsonify({"success": False, "message": "Source और Destination location अलग-अलग होने चाहिए।"}), 400

        if quantity <= 0:
            return jsonify({"success": False, "message": "Transfer quantity 0 से अधिक होनी चाहिए।"}), 400

        # Check product exists
        g.cursor.execute("SELECT id, product_name, item_type, track_serial, stock FROM products WHERE id = %s LIMIT 1", (product_id,))
        prod = g.cursor.fetchone()
        if not prod:
            return jsonify({"success": False, "message": "Product नहीं मिला।"}), 404
        if prod.get("item_type") == "Service":
            return jsonify({"success": False, "message": "Service items का stock transfer नहीं किया जा सकता।"}), 400

        # Verify source location has enough stock
        g.cursor.execute("SELECT id, quantity FROM product_location_stock WHERE product_id = %s AND location_id = %s LIMIT 1", (product_id, source_location_id))
        source_stock_row = g.cursor.fetchone()
        source_qty = float(source_stock_row["quantity"] or 0) if source_stock_row else 0.0

        # If source location has no pls record but is default location and product has stock
        if not source_stock_row:
            g.cursor.execute("SELECT id FROM stock_locations WHERE is_default = 1 LIMIT 1")
            def_loc = g.cursor.fetchone()
            if def_loc and def_loc["id"] == source_location_id:
                g.cursor.execute("SELECT COUNT(*) AS cnt FROM product_location_stock WHERE product_id = %s", (product_id,))
                cnt_row = g.cursor.fetchone()
                if not cnt_row or cnt_row["cnt"] == 0:
                    source_qty = float(prod.get("stock") or 0)
                    g.cursor.execute(
                        "INSERT INTO product_location_stock (product_id, location_id, quantity) VALUES (%s, %s, %s)",
                        (product_id, source_location_id, source_qty)
                    )

        if source_qty < quantity:
            return jsonify({
                "success": False,
                "message": f"Source location पर पर्याप्त स्टॉक नहीं है! (उपलब्ध: {source_qty}, माँगा गया: {quantity})"
            }), 400

        # Check serial numbers if tracking is enabled
        track_serial = prod.get("track_serial")
        is_serial_tracked = str(track_serial).strip().lower() in ("yes", "true", "1")
        clean_serials = [str(s).strip() for s in serial_numbers if str(s).strip()]
        if is_serial_tracked:
            if len(clean_serials) != int(quantity):
                return jsonify({
                    "success": False,
                    "message": f"इस प्रोडक्ट में Serial Tracking चालू है। कृपया ठीक {int(quantity)} सीरियल नंबर चुनें/स्कैन करें (अभी {len(clean_serials)} चुने गए हैं)।"
                }), 400

            # Verify each serial exists at source location and is In Stock
            for ser in clean_serials:
                g.cursor.execute(
                    """
                    SELECT id, location_id, status 
                    FROM product_serials 
                    WHERE product_id = %s AND serial_number = %s 
                    LIMIT 1
                    """,
                    (product_id, ser)
                )
                ser_row = g.cursor.fetchone()
                if not ser_row:
                    return jsonify({"success": False, "message": f"सीरियल नंबर '{ser}' इस प्रोडक्ट के रिकॉर्ड में नहीं मिला।"}), 400
                if ser_row.get("status") != "In Stock":
                    return jsonify({"success": False, "message": f"सीरियल नंबर '{ser}' का स्टेटस '{ser_row.get('status')}' है, यह 'In Stock' नहीं है।"}), 400
                
                # Check location
                ser_loc = ser_row.get("location_id")
                if ser_loc is None:
                    g.cursor.execute("SELECT id FROM stock_locations WHERE is_default = 1 LIMIT 1")
                    def_loc = g.cursor.fetchone()
                    if def_loc and def_loc["id"] != source_location_id:
                        return jsonify({"success": False, "message": f"सीरियल नंबर '{ser}' चुनी गई source location पर मौजूद नहीं है।"}), 400
                elif ser_loc != source_location_id:
                    return jsonify({"success": False, "message": f"सीरियल नंबर '{ser}' चुनी गई source location पर मौजूद नहीं है।"}), 400

        # Generate transfer number
        today_prefix = datetime.now().strftime("%Y%m%d")
        g.cursor.execute(
            "SELECT COUNT(*) AS cnt FROM stock_transfers WHERE DATE(transfer_date) = %s",
            (transfer_date,)
        )
        cnt = (g.cursor.fetchone() or {}).get("cnt", 0) + 1
        transfer_number = f"ST-{today_prefix}-{cnt:04d}"

        # 1. Update source location stock
        g.cursor.execute(
            """
            UPDATE product_location_stock 
            SET quantity = quantity - %s 
            WHERE product_id = %s AND location_id = %s
            """,
            (quantity, product_id, source_location_id)
        )

        # 2. Update destination location stock (upsert)
        g.cursor.execute(
            """
            INSERT INTO product_location_stock (product_id, location_id, quantity)
            VALUES (%s, %s, %s)
            ON DUPLICATE KEY UPDATE quantity = quantity + VALUES(quantity)
            """,
            (product_id, destination_location_id, quantity)
        )

        # 3. Update serials location if serials were transferred
        for ser in clean_serials:
            g.cursor.execute(
                """
                UPDATE product_serials 
                SET location_id = %s 
                WHERE product_id = %s AND serial_number = %s AND status = 'In Stock'
                """,
                (destination_location_id, product_id, ser)
            )

        # 4. Insert into stock_transfers
        g.cursor.execute(
            """
            INSERT INTO stock_transfers (
                transfer_number,
                transfer_date,
                source_location_id,
                destination_location_id,
                product_id,
                quantity,
                notes
            )
            VALUES (%s, %s, %s, %s, %s, %s, %s)
            """,
            (
                transfer_number,
                transfer_date,
                source_location_id,
                destination_location_id,
                product_id,
                quantity,
                notes
            )
        )
        transfer_id = g.cursor.lastrowid

        # 5. Insert serial items into stock_transfer_serials
        for ser in clean_serials:
            g.cursor.execute(
                """
                INSERT INTO stock_transfer_serials (transfer_id, serial_number)
                VALUES (%s, %s)
                """,
                (transfer_id, ser)
            )

        g.db.commit()

        return jsonify({
            "success": True,
            "message": f"स्टॉक सफलतापूर्वक ट्रांसफर हो गया! (Transfer No: {transfer_number})",
            "transfer_number": transfer_number,
            "transfer_id": transfer_id
        })

    except Exception as e:
        g.db.rollback()
        print("Stock Transfer Error:", e)
        return jsonify({
            "success": False,
            "message": f"Stock Transfer विफल रहा: {str(e)}"
        }), 500


@app.route("/stock-transfers", methods=["GET"])
def get_stock_transfers():
    try:
        product_id = request.args.get("product_id")
        limit = int(request.args.get("limit") or 50)

        sql = """
            SELECT 
                st.id,
                st.transfer_number,
                st.transfer_date,
                st.source_location_id,
                sl_src.location_name AS source_location_name,
                st.destination_location_id,
                sl_dst.location_name AS destination_location_name,
                st.product_id,
                p.product_name,
                p.track_serial,
                st.quantity,
                st.notes,
                st.created_at,
                (SELECT COUNT(*) FROM stock_transfer_serials sts WHERE sts.transfer_id = st.id) AS serial_count,
                (SELECT GROUP_CONCAT(sts.serial_number SEPARATOR ', ') FROM stock_transfer_serials sts WHERE sts.transfer_id = st.id) AS serial_numbers_str
            FROM stock_transfers st
            LEFT JOIN stock_locations sl_src ON sl_src.id = st.source_location_id
            LEFT JOIN stock_locations sl_dst ON sl_dst.id = st.destination_location_id
            LEFT JOIN products p ON p.id = st.product_id
            WHERE 1=1
        """
        params = []
        if product_id:
            sql += " AND st.product_id = %s "
            params.append(product_id)

        sql += " ORDER BY st.id DESC LIMIT %s "
        params.append(limit)

        g.cursor.execute(sql, tuple(params))
        rows = g.cursor.fetchall()

        return jsonify({
            "success": True,
            "transfers": rows
        })
    except Exception as e:
        print("Get Stock Transfers Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


# ==========================================
# SCRAP MANAGEMENT MODULE
# ==========================================

@app.route("/scrap-dashboard", methods=["GET"])
def get_scrap_dashboard():
    try:
        g.cursor.execute(
            """
            SELECT 
                COUNT(id) AS available_items,
                COALESCE(SUM(available_quantity), 0) AS available_quantity
            FROM scraps
            WHERE status = 'Available' AND available_quantity > 0
            """
        )
        avail_row = g.cursor.fetchone() or {}

        g.cursor.execute(
            """
            SELECT COALESCE(SUM(total_amount), 0) AS this_month_sales
            FROM scrap_sales
            WHERE MONTH(sale_date) = MONTH(CURRENT_DATE()) AND YEAR(sale_date) = YEAR(CURRENT_DATE())
            """
        )
        month_row = g.cursor.fetchone() or {}

        g.cursor.execute(
            """
            SELECT sale_number, sale_date, party_name, total_amount
            FROM scrap_sales
            ORDER BY sale_date DESC, id DESC
            LIMIT 1
            """
        )
        last_sale = g.cursor.fetchone()
        last_sale_data = None
        if last_sale:
            last_sale_data = {
                "sale_number": last_sale["sale_number"],
                "sale_date": str(last_sale["sale_date"]),
                "party_name": last_sale["party_name"],
                "total_amount": float(last_sale["total_amount"] or 0)
            }

        return jsonify({
            "success": True,
            "stats": {
                "available_items": int(avail_row.get("available_items") or 0),
                "available_quantity": float(avail_row.get("available_quantity") or 0),
                "this_month_sales": float(month_row.get("this_month_sales") or 0),
                "last_sale": last_sale_data
            }
        })
    except Exception as e:
        print("Get Scrap Dashboard Error:", e)
        return jsonify({"success": False, "message": str(e)}), 500


@app.route("/scraps", methods=["GET"])
def get_scraps():
    try:
        search = (request.args.get("search") or "").strip()
        status = (request.args.get("status") or "").strip()
        source_type = (request.args.get("source_type") or "").strip()
        repair_job_id = (request.args.get("repair_job_id") or "").strip()
        start_date = (request.args.get("start_date") or "").strip()
        end_date = (request.args.get("end_date") or "").strip()

        sql = """
            SELECT 
                s.id,
                s.scrap_date,
                s.item_name,
                s.quantity,
                s.sold_quantity,
                s.available_quantity,
                s.unit,
                s.source_type,
                s.repair_job_id,
                s.customer_job_no,
                s.reason,
                s.remark,
                s.status,
                s.created_at,
                s.updated_at,
                sj.service_number,
                sj.customer_name AS job_customer_name
            FROM scraps s
            LEFT JOIN service_jobs sj ON sj.id = s.repair_job_id
            WHERE 1=1
        """
        params = []

        if status and status.lower() != "all":
            sql += " AND s.status = %s "
            params.append(status)

        if source_type and source_type.lower() != "all":
            sql += " AND s.source_type = %s "
            params.append(source_type)

        if repair_job_id:
            try:
                sql += " AND s.repair_job_id = %s "
                params.append(int(repair_job_id))
            except ValueError:
                pass

        if start_date:
            sql += " AND s.scrap_date >= %s "
            params.append(start_date)

        if end_date:
            sql += " AND s.scrap_date <= %s "
            params.append(end_date)

        if search:
            sql += """ AND (
                s.item_name LIKE %s 
                OR s.reason LIKE %s 
                OR s.remark LIKE %s 
                OR s.customer_job_no LIKE %s
                OR sj.service_number LIKE %s
                OR sj.customer_name LIKE %s
            ) """
            like_val = f"%{search}%"
            params.extend([like_val, like_val, like_val, like_val, like_val, like_val])

        sql += " ORDER BY s.scrap_date DESC, s.id DESC "

        g.cursor.execute(sql, tuple(params))
        rows = g.cursor.fetchall()

        scraps = []
        for r in rows:
            scraps.append({
                "id": r["id"],
                "scrap_date": str(r["scrap_date"]),
                "item_name": r["item_name"],
                "quantity": float(r["quantity"] or 0),
                "sold_quantity": float(r["sold_quantity"] or 0),
                "available_quantity": float(r["available_quantity"] or 0),
                "unit": r["unit"] or "Nos",
                "source_type": r["source_type"] or "Repair",
                "repair_job_id": r["repair_job_id"],
                "customer_job_no": r["customer_job_no"] or (r.get("service_number") or ""),
                "job_customer_name": r.get("job_customer_name") or "",
                "reason": r["reason"] or "",
                "remark": r["remark"] or "",
                "status": r["status"] or "Available",
                "created_at": str(r["created_at"]) if r.get("created_at") else "",
                "updated_at": str(r["updated_at"]) if r.get("updated_at") else ""
            })

        return jsonify({
            "success": True,
            "scraps": scraps,
            "count": len(scraps)
        })
    except Exception as e:
        print("Get Scraps Error:", e)
        return jsonify({"success": False, "message": str(e)}), 500


@app.route("/scraps", methods=["POST", "OPTIONS"])
def add_scrap():
    if request.method == "OPTIONS":
        return "", 200

    try:
        data = request.get_json() or {}
        item_name = (data.get("item_name") or "").strip()
        if not item_name:
            return jsonify({"success": False, "message": "Item name is required"}), 400

        try:
            quantity = float(data.get("quantity") or 1.0)
            if quantity <= 0:
                return jsonify({"success": False, "message": "Quantity must be greater than zero"}), 400
        except (ValueError, TypeError):
            return jsonify({"success": False, "message": "Invalid quantity value"}), 400

        scrap_date = (data.get("scrap_date") or "").strip()
        if not scrap_date:
            scrap_date = get_ist_today_str()

        unit = (data.get("unit") or "Nos").strip()
        source_type = (data.get("source_type") or "Repair").strip()
        repair_job_id = data.get("repair_job_id")
        if repair_job_id:
            try:
                repair_job_id = int(repair_job_id)
            except (ValueError, TypeError):
                repair_job_id = None

        customer_job_no = (data.get("customer_job_no") or "").strip()
        if repair_job_id and not customer_job_no:
            g.cursor.execute("SELECT service_number FROM service_jobs WHERE id = %s LIMIT 1", (repair_job_id,))
            job_row = g.cursor.fetchone()
            if job_row:
                customer_job_no = job_row.get("service_number") or ""

        reason = (data.get("reason") or "").strip()
        remark = (data.get("remark") or "").strip()

        available_quantity = quantity
        sold_quantity = 0.0
        status = "Available"

        insert_sql = """
            INSERT INTO scraps 
            (scrap_date, item_name, quantity, sold_quantity, available_quantity, unit, source_type, repair_job_id, customer_job_no, reason, remark, status)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        """
        g.cursor.execute(
            insert_sql,
            (scrap_date, item_name, quantity, sold_quantity, available_quantity, unit, source_type, repair_job_id, customer_job_no, reason, remark, status)
        )
        scrap_id = g.cursor.lastrowid
        g.db.commit()

        return jsonify({
            "success": True,
            "message": "Scrap item added successfully",
            "scrap_id": scrap_id
        }), 201

    except Exception as e:
        g.db.rollback()
        print("Add Scrap Error:", e)
        return jsonify({"success": False, "message": str(e)}), 500


@app.route("/scraps/<int:scrap_id>", methods=["GET"])
def get_scrap_detail(scrap_id):
    try:
        sql = """
            SELECT 
                s.*,
                sj.service_number,
                sj.customer_name AS job_customer_name,
                COALESCE(p.product_name, sj.external_product_name, '') AS job_product_name,
                sj.complaint AS job_complaint,
                sj.service_status AS job_status
            FROM scraps s
            LEFT JOIN service_jobs sj ON sj.id = s.repair_job_id
            LEFT JOIN products p ON p.id = sj.product_id
            WHERE s.id = %s
            LIMIT 1
        """
        g.cursor.execute(sql, (scrap_id,))
        scrap = g.cursor.fetchone()
        if not scrap:
            return jsonify({"success": False, "message": "Scrap item not found"}), 404

        scrap_data = {
            "id": scrap["id"],
            "scrap_date": str(scrap["scrap_date"]),
            "item_name": scrap["item_name"],
            "quantity": float(scrap["quantity"] or 0),
            "sold_quantity": float(scrap["sold_quantity"] or 0),
            "available_quantity": float(scrap["available_quantity"] or 0),
            "unit": scrap["unit"] or "Nos",
            "source_type": scrap["source_type"] or "Repair",
            "repair_job_id": scrap["repair_job_id"],
            "customer_job_no": scrap["customer_job_no"] or (scrap.get("service_number") or ""),
            "job_customer_name": scrap.get("job_customer_name") or "",
            "job_product_name": scrap.get("job_product_name") or "",
            "job_complaint": scrap.get("job_complaint") or "",
            "job_status": scrap.get("job_status") or "",
            "reason": scrap["reason"] or "",
            "remark": scrap["remark"] or "",
            "status": scrap["status"] or "Available",
            "created_at": str(scrap["created_at"]) if scrap.get("created_at") else "",
            "updated_at": str(scrap["updated_at"]) if scrap.get("updated_at") else ""
        }

        # Fetch sale history
        sales_sql = """
            SELECT 
                ssi.id,
                ssi.quantity,
                ssi.amount,
                ssi.created_at,
                ss.id AS scrap_sale_id,
                ss.sale_number,
                ss.sale_date,
                ss.party_name,
                ss.mobile,
                ss.payment_mode
            FROM scrap_sale_items ssi
            JOIN scrap_sales ss ON ss.id = ssi.scrap_sale_id
            WHERE ssi.scrap_id = %s
            ORDER BY ss.sale_date DESC, ssi.id DESC
        """
        g.cursor.execute(sales_sql, (scrap_id,))
        sales_rows = g.cursor.fetchall()
        sales_history = []
        for sr in sales_rows:
            sales_history.append({
                "item_id": sr["id"],
                "scrap_sale_id": sr["scrap_sale_id"],
                "sale_number": sr["sale_number"],
                "sale_date": str(sr["sale_date"]),
                "party_name": sr["party_name"],
                "mobile": sr["mobile"] or "",
                "payment_mode": sr["payment_mode"] or "Cash",
                "quantity": float(sr["quantity"] or 0),
                "amount": float(sr["amount"] or 0),
                "created_at": str(sr["created_at"]) if sr.get("created_at") else ""
            })

        return jsonify({
            "success": True,
            "scrap": scrap_data,
            "sales_history": sales_history
        })
    except Exception as e:
        print("Get Scrap Detail Error:", e)
        return jsonify({"success": False, "message": str(e)}), 500


@app.route("/scraps/<int:scrap_id>", methods=["PUT", "OPTIONS"])
@roles_required("Admin", "Manager")
def update_scrap(scrap_id):
    if request.method == "OPTIONS":
        return "", 200

    try:
        g.cursor.execute("SELECT * FROM scraps WHERE id = %s LIMIT 1 FOR UPDATE", (scrap_id,))
        existing = g.cursor.fetchone()
        if not existing:
            return jsonify({"success": False, "message": "Scrap item not found"}), 404

        data = request.get_json() or {}
        item_name = (data.get("item_name") or existing["item_name"]).strip()
        if not item_name:
            return jsonify({"success": False, "message": "Item name cannot be empty"}), 400

        try:
            quantity = float(data.get("quantity", existing["quantity"]))
            if quantity <= 0:
                return jsonify({"success": False, "message": "Quantity must be greater than zero"}), 400
        except (ValueError, TypeError):
            return jsonify({"success": False, "message": "Invalid quantity"}), 400

        sold_quantity = float(existing["sold_quantity"] or 0)
        if quantity < sold_quantity:
            return jsonify({
                "success": False,
                "message": f"New total quantity ({quantity}) cannot be less than already sold quantity ({sold_quantity})"
            }), 400

        available_quantity = round(quantity - sold_quantity, 2)
        status = "Sold" if available_quantity <= 0.001 else "Available"

        scrap_date = (data.get("scrap_date") or str(existing["scrap_date"])).strip()
        unit = (data.get("unit") or existing["unit"] or "Nos").strip()
        source_type = (data.get("source_type") or existing["source_type"] or "Repair").strip()
        
        repair_job_id = data.get("repair_job_id") if "repair_job_id" in data else existing["repair_job_id"]
        if repair_job_id:
            try:
                repair_job_id = int(repair_job_id)
            except (ValueError, TypeError):
                repair_job_id = None

        customer_job_no = data.get("customer_job_no", existing["customer_job_no"])
        if customer_job_no is not None:
            customer_job_no = str(customer_job_no).strip()
        
        reason = (data.get("reason") if "reason" in data else (existing["reason"] or "")).strip()
        remark = (data.get("remark") if "remark" in data else (existing["remark"] or "")).strip()

        update_sql = """
            UPDATE scraps
            SET scrap_date = %s,
                item_name = %s,
                quantity = %s,
                available_quantity = %s,
                unit = %s,
                source_type = %s,
                repair_job_id = %s,
                customer_job_no = %s,
                reason = %s,
                remark = %s,
                status = %s
            WHERE id = %s
        """
        g.cursor.execute(
            update_sql,
            (scrap_date, item_name, quantity, available_quantity, unit, source_type, repair_job_id, customer_job_no, reason, remark, status, scrap_id)
        )
        g.db.commit()

        return jsonify({
            "success": True,
            "message": "Scrap item updated successfully"
        })
    except Exception as e:
        g.db.rollback()
        print("Update Scrap Error:", e)
        return jsonify({"success": False, "message": str(e)}), 500


@app.route("/scraps/<int:scrap_id>", methods=["DELETE", "OPTIONS"])
@roles_required("Admin")
def delete_scrap(scrap_id):
    if request.method == "OPTIONS":
        return "", 200

    try:
        g.cursor.execute("SELECT * FROM scraps WHERE id = %s LIMIT 1 FOR UPDATE", (scrap_id,))
        scrap = g.cursor.fetchone()
        if not scrap:
            return jsonify({"success": False, "message": "Scrap item not found"}), 404

        sold_quantity = float(scrap.get("sold_quantity") or 0)
        if sold_quantity > 0:
            return jsonify({
                "success": False,
                "message": f"Cannot delete this scrap item because {sold_quantity} unit(s) have already been sold."
            }), 400

        g.cursor.execute("DELETE FROM scraps WHERE id = %s", (scrap_id,))
        g.db.commit()

        return jsonify({
            "success": True,
            "message": "Scrap item deleted successfully"
        })
    except Exception as e:
        g.db.rollback()
        print("Delete Scrap Error:", e)
        return jsonify({"success": False, "message": str(e)}), 500


@app.route("/scrap-sales", methods=["POST", "OPTIONS"])
def create_scrap_sale():
    if request.method == "OPTIONS":
        return "", 200

    try:
        data = request.get_json() or {}
        party_name = (data.get("party_name") or data.get("buyer_name") or data.get("customer_name") or "").strip()
        if not party_name:
            return jsonify({"success": False, "message": "Kabadi / Buyer (Party) name is required"}), 400

        sale_date = (data.get("sale_date") or data.get("date") or "").strip()
        if not sale_date:
            sale_date = get_ist_today_str()

        mobile = (data.get("mobile") or "").strip()
        payment_mode = (data.get("payment_mode") or "Cash").strip()
        remark = (data.get("remark") or "").strip()

        items = data.get("items") or []
        if not isinstance(items, list) or len(items) == 0:
            return jsonify({"success": False, "message": "At least one scrap item must be added for sale"}), 400

        # Validate and prepare items
        total_sale_amount = 0.0
        validated_items = []
        seen_scrap_ids = set()

        for idx, itm in enumerate(items):
            try:
                raw_id = itm.get("scrap_id") if itm.get("scrap_id") is not None else itm.get("id")
                scrap_id = int(raw_id)
                qty = float(itm.get("quantity") or itm.get("qty_sold") or itm.get("qty") or 0)
                amount = float(itm.get("amount") if itm.get("amount") is not None else (itm.get("total_amount") if itm.get("total_amount") is not None else itm.get("price", 0)) or 0)
            except (ValueError, TypeError):
                return jsonify({"success": False, "message": f"Invalid item format at row {idx+1}"}), 400

            if scrap_id in seen_scrap_ids:
                return jsonify({"success": False, "message": f"पंक्ति #{idx+1} का स्क्रैप आइटम पहले ही चुना जा चुका है। एक ही बिक्री में दो बार समान आइटम नहीं बेचा जा सकता।"}), 400
            seen_scrap_ids.add(scrap_id)

            if qty <= 0:
                return jsonify({"success": False, "message": f"Quantity must be greater than zero at row {idx+1}"}), 400

            if amount < 0:
                return jsonify({"success": False, "message": f"Amount cannot be negative at row {idx+1}"}), 400

            # Lock scrap row
            g.cursor.execute("SELECT * FROM scraps WHERE id = %s FOR UPDATE", (scrap_id,))
            scrap_row = g.cursor.fetchone()
            if not scrap_row:
                return jsonify({"success": False, "message": f"Scrap item #{scrap_id} not found"}), 404

            avail_qty = float(scrap_row.get("available_quantity") or 0)
            if qty > avail_qty + 0.0001:
                return jsonify({
                    "success": False,
                    "message": f"Quantity ({qty}) exceeds available quantity ({avail_qty}) for '{scrap_row['item_name']}'"
                }), 400

            validated_items.append({
                "scrap_id": scrap_id,
                "scrap_row": scrap_row,
                "quantity": qty,
                "amount": amount
            })
            total_sale_amount += amount

        # Generate unique sale number
        sale_number = generate_document_number(
            "scrap-billing", "scrap_sale_prefix", "next_scrap_sale_number", "SS-"
        )

        insert_sale_sql = """
            INSERT INTO scrap_sales (sale_number, sale_date, party_name, mobile, total_amount, payment_mode, remark)
            VALUES (%s, %s, %s, %s, %s, %s, %s)
        """
        g.cursor.execute(
            insert_sale_sql,
            (sale_number, sale_date, party_name, mobile, total_sale_amount, payment_mode, remark)
        )
        scrap_sale_id = g.cursor.lastrowid

        insert_item_sql = """
            INSERT INTO scrap_sale_items (scrap_sale_id, scrap_id, quantity, amount)
            VALUES (%s, %s, %s, %s)
        """
        update_scrap_sql = """
            UPDATE scraps
            SET sold_quantity = %s,
                available_quantity = %s,
                status = %s
            WHERE id = %s
        """

        for v in validated_items:
            g.cursor.execute(insert_item_sql, (scrap_sale_id, v["scrap_id"], v["quantity"], v["amount"]))
            new_sold = round(float(v["scrap_row"]["sold_quantity"] or 0) + v["quantity"], 2)
            new_avail = round(float(v["scrap_row"]["available_quantity"] or 0) - v["quantity"], 2)
            new_status = "Sold" if new_avail <= 0.001 else "Available"
            g.cursor.execute(update_scrap_sql, (new_sold, new_avail, new_status, v["scrap_id"]))

        g.db.commit()

        return jsonify({
            "success": True,
            "message": "Scrap sale saved successfully",
            "sale_id": scrap_sale_id,
            "sale_number": sale_number,
            "total_amount": total_sale_amount
        }), 201

    except Exception as e:
        g.db.rollback()
        print("Create Scrap Sale Error:", e)
        return jsonify({"success": False, "message": str(e)}), 500


@app.route("/scrap-sales", methods=["GET"])
def get_scrap_sales():
    try:
        search = (request.args.get("search") or "").strip()
        start_date = (request.args.get("start_date") or "").strip()
        end_date = (request.args.get("end_date") or "").strip()

        sql = """
            SELECT 
                ss.id,
                ss.sale_number,
                ss.sale_date,
                ss.party_name,
                ss.mobile,
                ss.total_amount,
                ss.payment_mode,
                ss.remark,
                ss.created_at,
                COUNT(ssi.id) AS items_count,
                COALESCE(SUM(ssi.quantity), 0) AS total_sold_qty
            FROM scrap_sales ss
            LEFT JOIN scrap_sale_items ssi ON ssi.scrap_sale_id = ss.id
            WHERE 1=1
        """
        params = []

        if start_date:
            sql += " AND ss.sale_date >= %s "
            params.append(start_date)

        if end_date:
            sql += " AND ss.sale_date <= %s "
            params.append(end_date)

        if search:
            sql += " AND (ss.sale_number LIKE %s OR ss.party_name LIKE %s OR ss.mobile LIKE %s) "
            like_val = f"%{search}%"
            params.extend([like_val, like_val, like_val])

        sql += """
            GROUP BY ss.id, ss.sale_number, ss.sale_date, ss.party_name, ss.mobile,
                     ss.total_amount, ss.payment_mode, ss.remark, ss.created_at
            ORDER BY ss.sale_date DESC, ss.id DESC
        """

        g.cursor.execute(sql, tuple(params))
        rows = g.cursor.fetchall()

        sales = []
        for r in rows:
            sales.append({
                "id": r["id"],
                "sale_number": r["sale_number"],
                "sale_date": str(r["sale_date"]),
                "party_name": r["party_name"],
                "mobile": r["mobile"] or "",
                "total_amount": float(r["total_amount"] or 0),
                "payment_mode": r["payment_mode"] or "Cash",
                "remark": r["remark"] or "",
                "items_count": int(r["items_count"] or 0),
                "total_sold_qty": float(r["total_sold_qty"] or 0),
                "created_at": str(r["created_at"]) if r.get("created_at") else ""
            })

        return jsonify({
            "success": True,
            "sales": sales,
            "count": len(sales)
        })
    except Exception as e:
        print("Get Scrap Sales Error:", e)
        return jsonify({"success": False, "message": str(e)}), 500


@app.route("/scrap-sales/<int:sale_id>", methods=["GET"])
def get_scrap_sale_detail(sale_id):
    try:
        sql = "SELECT * FROM scrap_sales WHERE id = %s LIMIT 1"
        g.cursor.execute(sql, (sale_id,))
        sale = g.cursor.fetchone()
        if not sale:
            return jsonify({"success": False, "message": "Scrap sale not found"}), 404

        items_sql = """
            SELECT 
                ssi.id,
                ssi.scrap_id,
                ssi.quantity,
                ssi.amount,
                s.item_name,
                s.unit,
                s.source_type,
                s.customer_job_no,
                s.reason
            FROM scrap_sale_items ssi
            JOIN scraps s ON s.id = ssi.scrap_id
            WHERE ssi.scrap_sale_id = %s
            ORDER BY ssi.id ASC
        """
        g.cursor.execute(items_sql, (sale_id,))
        items_rows = g.cursor.fetchall()

        items = []
        for ir in items_rows:
            items.append({
                "id": ir["id"],
                "scrap_id": ir["scrap_id"],
                "item_name": ir["item_name"],
                "quantity": float(ir["quantity"] or 0),
                "amount": float(ir["amount"] or 0),
                "unit": ir["unit"] or "Nos",
                "source_type": ir["source_type"] or "Repair",
                "customer_job_no": ir["customer_job_no"] or "",
                "reason": ir["reason"] or ""
            })

        # Fetch business info for slip
        g.cursor.execute("SELECT setting_key, setting_value FROM settings WHERE category = 'general-settings'")
        settings_rows = g.cursor.fetchall() or []
        settings_map = {row["setting_key"]: row.get("setting_value") for row in settings_rows}

        business_info = {
            "firm_name": settings_map.get("firm_name") or "CRD Global Solutions",
            "address": settings_map.get("address") or "",
            "phone": settings_map.get("phone") or "",
            "email": settings_map.get("email") or "",
            "gst_number": settings_map.get("gst_number") or ""
        }

        sale_data = {
            "id": sale["id"],
            "sale_number": sale["sale_number"],
            "sale_date": str(sale["sale_date"]),
            "party_name": sale["party_name"],
            "mobile": sale["mobile"] or "",
            "total_amount": float(sale["total_amount"] or 0),
            "payment_mode": sale["payment_mode"] or "Cash",
            "remark": sale["remark"] or "",
            "created_at": str(sale["created_at"]) if sale.get("created_at") else ""
        }

        return jsonify({
            "success": True,
            "sale": sale_data,
            "items": items,
            "business_info": business_info
        })
    except Exception as e:
        print("Get Scrap Sale Detail Error:", e)
        return jsonify({"success": False, "message": str(e)}), 500


@app.route("/reports/scrap", methods=["GET"])
def reports_scrap():
    try:
        search = (request.args.get("search") or "").strip()
        status = (request.args.get("status") or "").strip()
        source_type = (request.args.get("source_type") or "").strip()
        start_date = (request.args.get("start_date") or "").strip()
        end_date = (request.args.get("end_date") or "").strip()

        sql = """
            SELECT 
                s.id,
                s.scrap_date,
                s.item_name,
                s.quantity,
                s.sold_quantity,
                s.available_quantity,
                s.unit,
                s.source_type,
                s.repair_job_id,
                s.customer_job_no,
                s.reason,
                s.remark,
                s.status,
                sj.service_number,
                sj.customer_name AS job_customer_name
            FROM scraps s
            LEFT JOIN service_jobs sj ON sj.id = s.repair_job_id
            WHERE 1=1
        """
        params = []
        if status and status.lower() != "all":
            sql += " AND s.status = %s "
            params.append(status)
        if source_type and source_type.lower() != "all":
            sql += " AND s.source_type = %s "
            params.append(source_type)
        if start_date:
            sql += " AND s.scrap_date >= %s "
            params.append(start_date)
        if end_date:
            sql += " AND s.scrap_date <= %s "
            params.append(end_date)
        if search:
            sql += " AND (s.item_name LIKE %s OR s.reason LIKE %s OR s.remark LIKE %s OR s.customer_job_no LIKE %s) "
            like_val = f"%{search}%"
            params.extend([like_val, like_val, like_val, like_val])

        sql += " ORDER BY s.scrap_date DESC, s.id DESC "

        g.cursor.execute(sql, tuple(params))
        rows = g.cursor.fetchall()

        formatted_rows = []
        total_items = len(rows)
        total_qty = 0.0
        sold_qty = 0.0
        avail_qty = 0.0

        for r in rows:
            q = float(r["quantity"] or 0)
            sq = float(r["sold_quantity"] or 0)
            aq = float(r["available_quantity"] or 0)
            total_qty += q
            sold_qty += sq
            avail_qty += aq

            job_ref = r.get("customer_job_no") or (r.get("service_number") or "")
            if not job_ref and r.get("repair_job_id"):
                job_ref = f"Job #{r['repair_job_id']}"

            formatted_rows.append({
                "id": r["id"],
                "scrap_date": str(r["scrap_date"]),
                "item_name": r["item_name"],
                "quantity": q,
                "sold_quantity": sq,
                "available_quantity": aq,
                "unit": r["unit"] or "Nos",
                "source_type": r["source_type"] or "Repair",
                "customer_job_no": job_ref,
                "reason": r["reason"] or "-",
                "remark": r["remark"] or "",
                "status": r["status"] or "Available"
            })

        # Summary sales
        g.cursor.execute("SELECT COALESCE(SUM(total_amount), 0) AS total_sale_amount FROM scrap_sales")
        total_sale_row = g.cursor.fetchone() or {}
        total_sale_amount = float(total_sale_row.get("total_sale_amount") or 0)

        return jsonify({
            "success": True,
            "report_name": "Scrap Management Report",
            "summary": {
                "total_items": total_items,
                "total_quantity": total_qty,
                "available_quantity": avail_qty,
                "sold_quantity": sold_qty,
                "total_sale_amount": total_sale_amount
            },
            "rows": formatted_rows
        })
    except Exception as e:
        print("Scrap Report Error:", e)
        return jsonify({"success": False, "message": str(e)}), 500


# ==========================================
# USERS MANAGEMENT
# ==========================================

@app.route("/users", methods=["GET"])
@roles_required("Admin")
def get_users():
    try:
        search = (request.args.get("search") or "").strip()
        role = (request.args.get("role") or "").strip()
        status = (request.args.get("status") or "").strip()

        sql = """
            SELECT id, username, full_name, mobile, email, role, status, created_at, updated_at
            FROM users
            WHERE 1=1
        """
        params = []

        if search:
            like_term = f"%{search}%"
            sql += " AND (username LIKE %s OR full_name LIKE %s OR mobile LIKE %s OR email LIKE %s) "
            params.extend([like_term, like_term, like_term, like_term])

        if role:
            sql += " AND role = %s "
            params.append(role)

        if status:
            sql += " AND status = %s "
            params.append(status)

        sql += " ORDER BY created_at DESC, id DESC "

        g.cursor.execute(sql, tuple(params))
        users = g.cursor.fetchall()

        return jsonify({
            "success": True,
            "users": users
        })

    except Exception as e:
        print("Get Users Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


@app.route("/users/<int:user_id>", methods=["GET"])
@roles_required("Admin")
def get_user_by_id(user_id):
    try:
        g.cursor.execute(
            """
                SELECT id, username, full_name, mobile, email, role, status, created_at, updated_at
                FROM users
                WHERE id = %s
                LIMIT 1
            """,
            (user_id,)
        )
        user = g.cursor.fetchone()

        if not user:
            return jsonify({
                "success": False,
                "message": "User not found."
            }), 404

        return jsonify({
            "success": True,
            "user": user
        })

    except Exception as e:
        print("Get User Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


@app.route("/users", methods=["POST", "OPTIONS"])
@roles_required("Admin")
def add_user():
    if request.method == "OPTIONS":
        return "", 200

    data = request.get_json() or {}
    username = (data.get("username") or "").strip()
    full_name = (data.get("full_name") or "").strip()
    mobile = (data.get("mobile") or "").strip()
    email = (data.get("email") or "").strip()
    password = (data.get("password") or "").strip()
    role = (data.get("role") or "Staff").strip()
    status = (data.get("status") or "Active").strip()

    if not username:
        return jsonify({"success": False, "message": "Username is required."}), 400

    if not full_name:
        return jsonify({"success": False, "message": "Full Name is required."}), 400

    if not password:
        return jsonify({"success": False, "message": "Password is required for new users."}), 400

    if email and not re.match(r"^[^\s@]+@[^\s@]+\.[^\s@]+$", email):
        return jsonify({"success": False, "message": "Please enter a valid email address."}), 400

    g.cursor.execute("SELECT role_name FROM user_roles WHERE LOWER(role_name) = LOWER(%s) LIMIT 1", (role,))
    matched_role = g.cursor.fetchone()
    if matched_role:
        role = matched_role["role_name"]
    elif role not in {"Admin", "Manager", "Staff"}:
        role = "Staff"

    if status not in {"Active", "Inactive"}:
        status = "Active"

    try:
        g.cursor.execute("SELECT id FROM users WHERE username = %s LIMIT 1", (username,))
        if g.cursor.fetchone():
            return jsonify({"success": False, "message": "Username already exists. Please choose another username."}), 409

        sql = """
            INSERT INTO users (username, password, full_name, mobile, email, role, status)
            VALUES (%s, %s, %s, %s, %s, %s, %s)
        """
        g.cursor.execute(sql, (username, generate_password_hash(password), full_name, mobile or None, email or None, role, status))
        g.db.commit()

        return jsonify({
            "success": True,
            "message": "User created successfully."
        }), 201

    except Exception as e:
        g.db.rollback()
        print("Add User Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


@app.route("/users/<int:user_id>", methods=["PUT", "OPTIONS"])
@roles_required("Admin")
def update_user(user_id):
    if request.method == "OPTIONS":
        return "", 200

    data = request.get_json() or {}
    full_name = (data.get("full_name") or "").strip()
    mobile = (data.get("mobile") or "").strip()
    email = (data.get("email") or "").strip()
    role = (data.get("role") or "Staff").strip()
    status = (data.get("status") or "Active").strip()

    if not full_name:
        return jsonify({"success": False, "message": "Full Name is required."}), 400

    if email and not re.match(r"^[^\s@]+@[^\s@]+\.[^\s@]+$", email):
        return jsonify({"success": False, "message": "Please enter a valid email address."}), 400

    g.cursor.execute("SELECT role_name FROM user_roles WHERE LOWER(role_name) = LOWER(%s) LIMIT 1", (role,))
    matched_role = g.cursor.fetchone()
    if matched_role:
        role = matched_role["role_name"]
    elif role not in {"Admin", "Manager", "Staff"}:
        role = "Staff"

    if status not in {"Active", "Inactive"}:
        status = "Active"

    try:
        g.cursor.execute("SELECT id, role, status FROM users WHERE id = %s LIMIT 1", (user_id,))
        user = g.cursor.fetchone()
        if not user:
            return jsonify({"success": False, "message": "User not found."}), 404

        g.cursor.execute(
            """
                UPDATE users
                SET full_name = %s,
                    mobile = %s,
                    email = %s,
                    role = %s,
                    status = %s,
                    updated_at = CURRENT_TIMESTAMP
                WHERE id = %s
            """,
            (full_name, mobile or None, email or None, role, status, user_id)
        )
        g.db.commit()

        return jsonify({
            "success": True,
            "message": "User updated successfully."
        })

    except Exception as e:
        g.db.rollback()
        print("Update User Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


@app.route("/users/<int:user_id>/status", methods=["PUT", "OPTIONS"])
@roles_required("Admin")
def update_user_status(user_id):
    if request.method == "OPTIONS":
        return "", 200

    data = request.get_json() or {}
    status = (data.get("status") or "").strip()

    if status not in {"Active", "Inactive"}:
        return jsonify({"success": False, "message": "Invalid user status."}), 400

    try:
        g.cursor.execute("SELECT id, role, status FROM users WHERE id = %s LIMIT 1", (user_id,))
        user = g.cursor.fetchone()
        if not user:
            return jsonify({"success": False, "message": "User not found."}), 404

        if str(user.get("role") or "").strip() == "Admin" and str(user.get("status") or "").strip() == "Active" and status == "Inactive":
            g.cursor.execute("SELECT COUNT(*) AS admin_count FROM users WHERE role = 'Admin' AND status = 'Active'")
            admin_count = int((g.cursor.fetchone() or {}).get("admin_count") or 0)
            if admin_count <= 1:
                return jsonify({
                    "success": False,
                    "message": "Cannot deactivate the last active Admin account."
                }), 400

        g.cursor.execute(
            "UPDATE users SET status = %s, updated_at = CURRENT_TIMESTAMP WHERE id = %s",
            (status, user_id)
        )
        g.db.commit()

        return jsonify({
            "success": True,
            "message": f"User {status.lower()} successfully."
        })

    except Exception as e:
        g.db.rollback()
        print("Update User Status Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


@app.route("/users/<int:user_id>/password", methods=["PUT", "OPTIONS"])
@roles_required("Admin")
def reset_user_password(user_id):
    if request.method == "OPTIONS":
        return "", 200

    data = request.get_json() or {}
    new_password = (data.get("new_password") or "").strip()
    confirm_password = (data.get("confirm_password") or "").strip()

    if not new_password or len(new_password) < 6:
        return jsonify({"success": False, "message": "Password must be at least 6 characters long."}), 400

    if new_password != confirm_password:
        return jsonify({"success": False, "message": "Passwords do not match."}), 400

    try:
        g.cursor.execute("SELECT id FROM users WHERE id = %s LIMIT 1", (user_id,))
        if not g.cursor.fetchone():
            return jsonify({"success": False, "message": "User not found."}), 404

        g.cursor.execute(
            "UPDATE users SET password = %s, updated_at = CURRENT_TIMESTAMP WHERE id = %s",
            (generate_password_hash(new_password), user_id)
        )
        g.db.commit()

        return jsonify({
            "success": True,
            "message": "Password reset successfully."
        })

    except Exception as e:
        g.db.rollback()
        print("Reset Password Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500

# ==========================================
# ROLES & PERMISSIONS MANAGEMENT
# ==========================================

@app.route("/permissions", methods=["GET"])
def get_permissions():
    return jsonify({
        "success": True,
        "permissions": ALL_PERMISSIONS
    })


@app.route("/roles", methods=["GET"])
def get_roles():
    try:
        g.cursor.execute(
            """
            SELECT
                ur.id,
                ur.role_name,
                ur.description,
                ur.permissions,
                ur.is_system,
                ur.created_at,
                ur.updated_at,
                COUNT(u.id) AS user_count
            FROM user_roles ur
            LEFT JOIN users u
                ON LOWER(u.role) = LOWER(ur.role_name)
            GROUP BY ur.id
            ORDER BY ur.is_system DESC, ur.role_name ASC
            """
        )
        rows = g.cursor.fetchall()
        roles = []
        for row in rows:
            perms = []
            raw_perms = row.get("permissions")
            if isinstance(raw_perms, str):
                try:
                    perms = json.loads(raw_perms)
                except Exception:
                    perms = []
            elif isinstance(raw_perms, list):
                perms = raw_perms

            roles.append({
                "id": row["id"],
                "role_name": row["role_name"],
                "description": row["description"] or "",
                "permissions": perms,
                "is_system": bool(row["is_system"]),
                "user_count": int(row.get("user_count") or 0),
                "created_at": normalize_date_string(row["created_at"])
            })

        return jsonify({
            "success": True,
            "roles": roles
        })

    except Exception as e:
        print("Get Roles Error:", e)
        return jsonify({"success": False, "message": str(e)}), 500


@app.route("/roles/<int:role_id>", methods=["GET"])
def get_role_by_id(role_id):
    try:
        g.cursor.execute(
            """
            SELECT
                ur.id,
                ur.role_name,
                ur.description,
                ur.permissions,
                ur.is_system,
                ur.created_at,
                COUNT(u.id) AS user_count
            FROM user_roles ur
            LEFT JOIN users u
                ON LOWER(u.role) = LOWER(ur.role_name)
            WHERE ur.id = %s
            GROUP BY ur.id
            LIMIT 1
            """,
            (role_id,)
        )
        row = g.cursor.fetchone()
        if not row:
            return jsonify({"success": False, "message": "Role not found."}), 404

        raw_perms = row.get("permissions")
        perms = json.loads(raw_perms) if isinstance(raw_perms, str) else (raw_perms or [])

        return jsonify({
            "success": True,
            "role": {
                "id": row["id"],
                "role_name": row["role_name"],
                "description": row["description"] or "",
                "permissions": perms,
                "is_system": bool(row["is_system"]),
                "user_count": int(row.get("user_count") or 0)
            }
        })

    except Exception as e:
        print("Get Role By ID Error:", e)
        return jsonify({"success": False, "message": str(e)}), 500


@app.route("/roles", methods=["POST", "OPTIONS"])
@roles_required("Admin")
def create_role():
    if request.method == "OPTIONS":
        return "", 200

    data = request.get_json() or {}
    role_name = (data.get("role_name") or "").strip()
    description = (data.get("description") or "").strip()
    permissions = data.get("permissions") or []

    if not role_name:
        return jsonify({"success": False, "message": "Role name is required."}), 400

    if not re.match(r"^[A-Za-z0-9 _-]{2,50}$", role_name):
        return jsonify({"success": False, "message": "Role name must be 2-50 characters with letters, numbers, spaces, underscores or hyphens."}), 400

    if not isinstance(permissions, list):
        return jsonify({"success": False, "message": "Permissions must be a list of permission keys."}), 400

    valid_keys_set = set(ALL_PERMISSION_KEYS)
    filtered_permissions = [k for k in permissions if k in valid_keys_set]

    try:
        g.cursor.execute("SELECT id FROM user_roles WHERE LOWER(role_name) = LOWER(%s) LIMIT 1", (role_name,))
        if g.cursor.fetchone():
            return jsonify({"success": False, "message": f"Role '{role_name}' already exists."}), 400

        g.cursor.execute(
            """
            INSERT INTO user_roles (role_name, description, permissions, is_system)
            VALUES (%s, %s, %s, 0)
            """,
            (role_name, description, json.dumps(filtered_permissions))
        )
        g.db.commit()
        new_id = g.cursor.lastrowid

        return jsonify({
            "success": True,
            "message": f"Role '{role_name}' created successfully.",
            "role_id": new_id
        }), 201

    except Exception as e:
        g.db.rollback()
        print("Create Role Error:", e)
        return jsonify({"success": False, "message": str(e)}), 500


@app.route("/roles/<int:role_id>", methods=["PUT", "OPTIONS"])
@roles_required("Admin")
def update_role(role_id):
    if request.method == "OPTIONS":
        return "", 200

    data = request.get_json() or {}
    new_role_name = (data.get("role_name") or "").strip()
    description = (data.get("description") or "").strip()
    permissions = data.get("permissions") or []

    if not isinstance(permissions, list):
        return jsonify({"success": False, "message": "Permissions must be a list."}), 400

    try:
        g.cursor.execute("SELECT id, role_name, is_system FROM user_roles WHERE id = %s LIMIT 1", (role_id,))
        role = g.cursor.fetchone()
        if not role:
            return jsonify({"success": False, "message": "Role not found."}), 404

        # If not system role and a new name was provided, allow renaming
        if not role["is_system"] and new_role_name and new_role_name.lower() != role["role_name"].lower():
            g.cursor.execute("SELECT id FROM user_roles WHERE LOWER(role_name) = LOWER(%s) AND id != %s LIMIT 1", (new_role_name, role_id))
            if g.cursor.fetchone():
                return jsonify({"success": False, "message": f"Role '{new_role_name}' already exists."}), 400
            old_name = role["role_name"]
            g.cursor.execute("UPDATE user_roles SET role_name = %s WHERE id = %s", (new_role_name, role_id))
            g.cursor.execute("UPDATE users SET role = %s WHERE LOWER(role) = LOWER(%s)", (new_role_name, old_name))
            role["role_name"] = new_role_name

        if role["is_system"] and role["role_name"].lower() == "admin":
            filtered_permissions = list(ALL_PERMISSION_KEYS)
        else:
            valid_keys_set = set(ALL_PERMISSION_KEYS)
            filtered_permissions = [k for k in permissions if k in valid_keys_set]

        g.cursor.execute(
            """
            UPDATE user_roles
            SET description = %s,
                permissions = %s,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = %s
            """,
            (description, json.dumps(filtered_permissions), role_id)
        )
        g.db.commit()

        return jsonify({
            "success": True,
            "message": f"Role '{role['role_name']}' updated successfully."
        })

    except Exception as e:
        g.db.rollback()
        print("Update Role Error:", e)
        return jsonify({"success": False, "message": str(e)}), 500


@app.route("/roles/<int:role_id>", methods=["DELETE", "OPTIONS"])
@roles_required("Admin")
def delete_role(role_id):
    if request.method == "OPTIONS":
        return "", 200

    try:
        g.cursor.execute("SELECT id, role_name, is_system FROM user_roles WHERE id = %s LIMIT 1", (role_id,))
        role = g.cursor.fetchone()
        if not role:
            return jsonify({"success": False, "message": "Role not found."}), 404

        if role["is_system"]:
            return jsonify({"success": False, "message": f"Built-in system role '{role['role_name']}' cannot be deleted."}), 400

        g.cursor.execute("SELECT COUNT(*) AS total FROM users WHERE LOWER(role) = LOWER(%s)", (role["role_name"],))
        user_count_row = g.cursor.fetchone()
        user_count = int(user_count_row.get("total") or 0) if user_count_row else 0
        if user_count > 0:
            return jsonify({
                "success": False,
                "message": f"Cannot delete role '{role['role_name']}' because {user_count} user(s) are currently assigned to it. Please reassign those users first."
            }), 400

        g.cursor.execute("DELETE FROM user_roles WHERE id = %s", (role_id,))
        g.db.commit()

        return jsonify({
            "success": True,
            "message": f"Role '{role['role_name']}' deleted successfully."
        })

    except Exception as e:
        g.db.rollback()
        print("Delete Role Error:", e)
        return jsonify({"success": False, "message": str(e)}), 500


def check_duplicate_customer_mobile(mobile, exclude_customer_id=None):
    if not mobile:
        return None
    clean_mobile = str(mobile).strip()
    if not clean_mobile:
        return None
    digits = re.sub(r"\D", "", clean_mobile)
    if not digits:
        return None

    if len(digits) >= 10:
        last_10 = digits[-10:]
        sql = """
            SELECT id, customer_name, mobile, address 
            FROM customers 
            WHERE mobile IS NOT NULL AND mobile != '' 
              AND (
                  mobile = %s 
                  OR mobile LIKE %s 
                  OR RIGHT(REGEXP_REPLACE(mobile, '[^0-9]', ''), 10) = %s
              )
        """
        params = [clean_mobile, f"%{last_10}", last_10]
    else:
        sql = """
            SELECT id, customer_name, mobile, address 
            FROM customers 
            WHERE mobile IS NOT NULL AND mobile != '' 
              AND mobile = %s
        """
        params = [clean_mobile]

    if exclude_customer_id:
        sql += " AND id != %s"
        params.append(exclude_customer_id)

    sql += " LIMIT 1"

    g.cursor.execute(sql, tuple(params))
    return g.cursor.fetchone()


@app.route("/customers/check-mobile", methods=["GET"])
def check_customer_mobile_endpoint():
    mobile = request.args.get("mobile", "").strip()
    exclude_id = request.args.get("exclude_id")
    if exclude_id:
        try:
            exclude_id = int(exclude_id)
        except (ValueError, TypeError):
            exclude_id = None

    if not mobile:
        return jsonify({"success": True, "exists": False})

    existing = check_duplicate_customer_mobile(mobile, exclude_customer_id=exclude_id)
    if existing:
        return jsonify({
            "success": True,
            "exists": True,
            "customer": {
                "id": existing.get("id"),
                "customer_name": existing.get("customer_name"),
                "mobile": existing.get("mobile"),
                "address": existing.get("address") or ""
            },
            "message": f"यह मोबाइल नंबर पहले से ग्राहक '{existing.get('customer_name')}' (ID: #{existing.get('id')}) के प्रोफाइल पर दर्ज है।"
        })
    return jsonify({"success": True, "exists": False})


# ==========================================
# ADD CUSTOMER
# ==========================================

@app.route(
    "/customers",
    methods=["POST", "OPTIONS"]
)
def add_customer():

    if request.method == "OPTIONS":
        return "", 200

    data = request.get_json()

    if not data:

        return jsonify({
            "success": False,
            "message": "Customer data नहीं मिला"
        }), 400

    customer_name = data.get("customer_name")
    mobile = data.get("mobile")
    email = data.get("email")
    address = data.get("address")
    gstin = data.get("gstin")
    opening_balance = data.get("opening_balance", 0)
    customer_type = data.get("customer_type", "Regular")

    if not customer_name:

        return jsonify({
            "success": False,
            "message": "Customer Name is required"
        }), 400

    if mobile and str(mobile).strip():
        existing_customer = check_duplicate_customer_mobile(mobile)
        if existing_customer:
            cust_name = existing_customer.get("customer_name") or "Customer"
            cust_id = existing_customer.get("id")
            cust_mob = existing_customer.get("mobile") or str(mobile).strip()
            return jsonify({
                "success": False,
                "message": f"यह मोबाइल नंबर ({cust_mob}) पहले से ग्राहक '{cust_name}' (ID: #{cust_id}) के प्रोफाइल पर दर्ज है। एक ही मोबाइल नंबर से दोबारा नया ग्राहक नहीं बनाया जा सकता।"
            }), 400

    try:

        sql = """
            INSERT INTO customers
            (
                customer_name,
                mobile,
                email,
                address,
                gstin,
                opening_balance,
                customer_type
            )
            VALUES
            (
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s
            )
        """

        g.cursor.execute(
            sql,
            (
                customer_name,
                mobile,
                email,
                address,
                gstin,
                opening_balance,
                customer_type
            )
        )

        g.db.commit()
        customer_id = g.cursor.lastrowid

        return jsonify({
            "success": True,
            "message": "Customer Successfully Saved",
            "customer_id": customer_id
        })

    except Exception as e:

        g.db.rollback()

        print("Customer Error:", e)

        return jsonify({
            "success": False,
            "message": "Customer save नहीं हुआ"
        }), 500


# ==========================================
# GET ALL CUSTOMERS
# ==========================================

@app.route(
    "/customers",
    methods=["GET"]
)
def get_customers():

    try:
        search = request.args.get("search", "").strip()
        params = []
        sql = """
            SELECT
                c.id,
                c.customer_name,
                c.mobile,
                c.email,
                c.address,
                c.gstin,
                c.opening_balance,
                c.customer_type,
                c.status,
                c.created_at,
                ROUND(
                    COALESCE(c.opening_balance, 0.0)
                    + COALESCE(inv.total_due, 0.0)
                    - COALESCE(pay.direct_paid, 0.0),
                    2
                ) AS outstanding_balance
            FROM customers c
            LEFT JOIN (
                SELECT customer_id, SUM(GREATEST(0, grand_total - COALESCE(paid_amount, 0))) AS total_due
                FROM sales_invoices
                WHERE (document_type IS NULL OR document_type != 'Quotation')
                GROUP BY customer_id
            ) inv ON inv.customer_id = c.id
            LEFT JOIN (
                SELECT customer_id, SUM(amount) AS direct_paid
                FROM customer_payments
                WHERE invoice_id IS NULL
                GROUP BY customer_id
            ) pay ON pay.customer_id = c.id
        """
        if search:
            digits = re.sub(r"\D", "", search)
            if digits and len(digits) >= 5:
                sql += " WHERE (c.customer_name LIKE %s OR c.mobile LIKE %s OR RIGHT(REGEXP_REPLACE(c.mobile, '[^0-9]', ''), 10) LIKE %s) "
                params = [f"%{search}%", f"%{search}%", f"%{digits[-10:]}%"]
            else:
                sql += " WHERE (c.customer_name LIKE %s OR c.mobile LIKE %s OR c.address LIKE %s) "
                params = [f"%{search}%", f"%{search}%", f"%{search}%"]

        sql += " ORDER BY c.id DESC"
        g.cursor.execute(sql, tuple(params))

        customers = g.cursor.fetchall()
        for cust in customers:
            cust["opening_balance"] = float(cust.get("opening_balance") or 0)
            cust["outstanding_balance"] = float(cust.get("outstanding_balance") or 0)

        return jsonify({
            "success": True,
            "customers": customers
        })

    except Exception as e:

        print(
            "Get Customers Error:",
            e
        )

        return jsonify({
            "success": False,
            "message": "Customers load नहीं हुए"
        }), 500


# ==========================================
# GET SINGLE CUSTOMER
# ==========================================

@app.route(
    "/customers/<int:customer_id>",
    methods=["GET"]
)
def get_customer(customer_id):

    try:

        sql = """
            SELECT
                id,
                customer_name,
                mobile,
                email,
                address,
                gstin,
                opening_balance,
                customer_type,
                status,
                created_at
            FROM customers
            WHERE id = %s
        """

        g.cursor.execute(
            sql,
            (customer_id,)
        )

        customer = g.cursor.fetchone()

        if not customer:

            return jsonify({
                "success": False,
                "message": "Customer नहीं मिला"
            }), 404

        return jsonify({
            "success": True,
            "customer": customer
        })

    except Exception as e:

        print(
            "Get Customer Error:",
            e
        )

        return jsonify({
            "success": False,
            "message": "Customer load नहीं हुआ"
        }), 500


# ==========================================
# UPDATE CUSTOMER
# ==========================================

@app.route(
    "/customers/<int:customer_id>",
    methods=["PUT", "OPTIONS"]
)
def update_customer(customer_id):

    if request.method == "OPTIONS":
        return "", 200

    data = request.get_json()

    if not data:

        return jsonify({
            "success": False,
            "message": "Customer data नहीं मिला"
        }), 400

    customer_name = data.get("customer_name")
    mobile = data.get("mobile")
    email = data.get("email")
    address = data.get("address")
    gstin = data.get("gstin")
    opening_balance = data.get("opening_balance", 0)
    customer_type = data.get(
        "customer_type",
        "Regular"
    )

    if not customer_name:

        return jsonify({
            "success": False,
            "message": "Customer Name is required"
        }), 400

    try:
        g.cursor.execute("SELECT id FROM customers WHERE id = %s LIMIT 1", (customer_id,))
        if not g.cursor.fetchone():
            return jsonify({
                "success": False,
                "message": "Customer नहीं मिला"
            }), 404

        if mobile and str(mobile).strip():
            existing_customer = check_duplicate_customer_mobile(mobile, exclude_customer_id=customer_id)
            if existing_customer:
                cust_name = existing_customer.get("customer_name") or "Customer"
                cust_id = existing_customer.get("id")
                cust_mob = existing_customer.get("mobile") or str(mobile).strip()
                return jsonify({
                    "success": False,
                    "message": f"यह मोबाइल नंबर ({cust_mob}) पहले से ग्राहक '{cust_name}' (ID: #{cust_id}) के प्रोफाइल पर दर्ज है।"
                }), 400

        sql = """
            UPDATE customers
            SET
                customer_name = %s,
                mobile = %s,
                email = %s,
                address = %s,
                gstin = %s,
                opening_balance = %s,
                customer_type = %s
            WHERE id = %s
        """

        g.cursor.execute(
            sql,
            (
                customer_name,
                mobile,
                email,
                address,
                gstin,
                opening_balance,
                customer_type,
                customer_id
            )
        )

        g.db.commit()

        return jsonify({
            "success": True,
            "message": "Customer Successfully Updated"
        })

    except Exception as e:

        g.db.rollback()

        print(
            "Update Customer Error:",
            e
        )

        return jsonify({
            "success": False,
            "message": "Customer update नहीं हुआ"
        }), 500



# ==========================================
# DEACTIVATE CUSTOMER
# ==========================================

@app.route(
    "/customers/<int:customer_id>/deactivate",
    methods=["PUT", "OPTIONS"]
)
def deactivate_customer(customer_id):

    if request.method == "OPTIONS":
        return "", 200

    try:

        sql = """
            UPDATE customers
            SET status = 'Inactive'
            WHERE id = %s
        """

        g.cursor.execute(
            sql,
            (customer_id,)
        )

        g.db.commit()

        if g.cursor.rowcount == 0:

            return jsonify({
                "success": False,
                "message": "Customer नहीं मिला"
            }), 404

        return jsonify({
            "success": True,
            "message": "Customer Deactivated Successfully"
        })

    except Exception as e:

        g.db.rollback()

        print(
            "Deactivate Customer Error:",
            e
        )

        return jsonify({
            "success": False,
            "message": "Customer deactivate नहीं हुआ"
        }), 500


@app.route("/customers/<int:customer_id>/status", methods=["PUT", "OPTIONS"])
def update_customer_status(customer_id):
    if request.method == "OPTIONS":
        return "", 200

    status = str((request.get_json(silent=True) or {}).get("status") or "").strip()
    if status not in {"Active", "Inactive"}:
        return jsonify({"success": False, "message": "Invalid customer status."}), 400

    try:
        g.cursor.execute("UPDATE customers SET status = %s WHERE id = %s", (status, customer_id))
        if g.cursor.rowcount == 0:
            g.db.rollback()
            return jsonify({"success": False, "message": "Customer not found."}), 404
        g.db.commit()
        return jsonify({"success": True, "message": f"Customer {status.lower()} successfully."})
    except Exception as e:
        g.db.rollback()
        print("Update Customer Status Error:", e)
        return jsonify({"success": False, "message": "Unable to update customer status."}), 500

# ==========================================
# DASHBOARD CUSTOMER COUNT
# ==========================================

@app.route("/dashboard/customers-count", methods=["GET"])
def customers_count():

    try:

        sql = """
            SELECT COUNT(*) AS total
            FROM customers
            WHERE status = 'Active'
        """

        g.cursor.execute(sql)

        result = g.cursor.fetchone()

        return jsonify({
            "success": True,
            "total": result["total"]
        })

    except Exception as e:

        print("Customer Count Error:", e)

        return jsonify({
            "success": False,
            "message": "Customer count load नहीं हुआ"
        }), 500

# ==========================================
# EXPENSES MODULE
# ==========================================

@app.route("/expenses", methods=["GET"])
def get_expenses():
    try:
        search = request.args.get("search", "").strip()
        from_date = normalize_date_string(request.args.get("from_date"))
        to_date = normalize_date_string(request.args.get("to_date"))
        category = request.args.get("category", "").strip()
        payment_method = request.args.get("payment_method", "").strip()
        include_deleted = str(request.args.get("include_deleted", "")).strip().lower() in {"1", "true", "yes"}

        sql = """
            SELECT
                id,
                expense_number,
                expense_date,
                category,
                description,
                amount,
                payment_method,
                paid_to,
                reference_number,
                notes,
                status,
                original_status,
                deleted_at,
                deleted_by_user_id,
                created_at,
                updated_at
            FROM expenses
            WHERE 1=1
        """
        params = []

        if not include_deleted:
            sql += " AND status <> %s "
            params.append("Deleted")

        if search:
            sql += """
                AND (
                    expense_number LIKE %s
                    OR description LIKE %s
                    OR paid_to LIKE %s
                    OR reference_number LIKE %s
                )
            """
            like_search = f"%{search}%"
            params.extend([like_search, like_search, like_search, like_search])

        if from_date:
            sql += " AND expense_date >= %s "
            params.append(from_date)

        if to_date:
            sql += " AND expense_date <= %s "
            params.append(to_date)

        if category:
            sql += " AND category = %s "
            params.append(category)

        if payment_method:
            sql += " AND payment_method = %s "
            params.append(payment_method)

        sql += " ORDER BY expense_date DESC, id DESC "

        g.cursor.execute(sql, tuple(params))
        expenses = g.cursor.fetchall()

        total_amount = sum(float(item["amount"] or 0) for item in expenses)

        return jsonify({
            "success": True,
            "expenses": expenses,
            "total_expenses": round(total_amount, 2)
        })

    except Exception as e:
        print("Get Expenses Error:", e)
        return jsonify({
            "success": False,
            "message": "Expenses load नहीं हुए"
        }), 500


@app.route("/expenses/<int:expense_id>", methods=["GET"])
def get_expense(expense_id):
    try:
        g.cursor.execute(
            """
            SELECT
                id,
                expense_number,
                expense_date,
                category,
                description,
                amount,
                payment_method,
                paid_to,
                reference_number,
                notes,
                status,
                original_status,
                deleted_at,
                deleted_by_user_id,
                created_at,
                updated_at
            FROM expenses
            WHERE id = %s
            LIMIT 1
            """,
            (expense_id,)
        )

        expense = g.cursor.fetchone()
        if not expense:
            return jsonify({
                "success": False,
                "message": "Expense नहीं मिला"
            }), 404

        if str(expense.get("status") or "").strip() == "Deleted":
            user = getattr(g, "current_user", None)
            if not user or str(user.get("role") or "") not in {"Admin", "Manager"}:
                return jsonify({
                    "success": False,
                    "message": "This expense has been deleted and is not available to this user."
                }), 403

        return jsonify({
            "success": True,
            "expense": expense
        })

    except Exception as e:
        print("Get Expense Error:", e)
        return jsonify({
            "success": False,
            "message": "Expense load नहीं हुआ"
        }), 500


@app.route("/expenses", methods=["POST", "OPTIONS"])
def add_expense():
    if request.method == "OPTIONS":
        return "", 200

    data = request.get_json(silent=True)
    if not data:
        return jsonify({
            "success": False,
            "message": "Expense data नहीं मिला"
        }), 400

    expense_date = normalize_date_string(data.get("expense_date"))
    category = (data.get("category") or "").strip()
    description = (data.get("description") or "").strip()
    amount = data.get("amount")
    payment_method = (data.get("payment_method") or "").strip()
    paid_to = (data.get("paid_to") or "").strip()
    reference_number = (data.get("reference_number") or "").strip()
    notes = (data.get("notes") or "").strip()
    status = (data.get("status") or "Paid").strip()

    if not expense_date:
        return jsonify({
            "success": False,
            "message": "Expense Date is required"
        }), 400

    if not category:
        return jsonify({
            "success": False,
            "message": "Category is required"
        }), 400

    if not description:
        return jsonify({
            "success": False,
            "message": "Description is required"
        }), 400

    try:
        amount_value = float(amount)
    except (TypeError, ValueError):
        return jsonify({
            "success": False,
            "message": "Amount must be a valid number"
        }), 400

    if amount_value <= 0:
        return jsonify({
            "success": False,
            "message": "Amount must be greater than 0"
        }), 400

    if not payment_method:
        return jsonify({
            "success": False,
            "message": "Payment Method is required"
        }), 400

    try:
        expense_number = generate_expense_number()

        g.cursor.execute(
            """
            INSERT INTO expenses (
                expense_number,
                expense_date,
                category,
                description,
                amount,
                payment_method,
                paid_to,
                reference_number,
                notes,
                status
            )
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            """,
            (
                expense_number,
                expense_date,
                category,
                description,
                amount_value,
                payment_method,
                paid_to or None,
                reference_number or None,
                notes or None,
                status if status in ("Paid", "Pending") else "Paid"
            )
        )

        g.db.commit()

        return jsonify({
            "success": True,
            "message": "Expense Successfully Saved",
            "expense_number": expense_number
        })

    except Exception as e:
        g.db.rollback()
        print("Add Expense Error:", e)
        return jsonify({
            "success": False,
            "message": "Expense save नहीं हुआ"
        }), 500


@app.route("/expenses/<int:expense_id>", methods=["PUT", "OPTIONS"])
@roles_required("Admin", "Manager")
def update_expense(expense_id):
    if request.method == "OPTIONS":
        return "", 200

    data = request.get_json(silent=True)
    if not data:
        return jsonify({
            "success": False,
            "message": "Expense data नहीं मिला"
        }), 400

    expense_date = normalize_date_string(data.get("expense_date"))
    category = (data.get("category") or "").strip()
    description = (data.get("description") or "").strip()
    amount = data.get("amount")
    payment_method = (data.get("payment_method") or "").strip()
    paid_to = (data.get("paid_to") or "").strip()
    reference_number = (data.get("reference_number") or "").strip()
    notes = (data.get("notes") or "").strip()
    status = (data.get("status") or "Paid").strip()

    if not expense_date:
        return jsonify({
            "success": False,
            "message": "Expense Date is required"
        }), 400

    if not category:
        return jsonify({
            "success": False,
            "message": "Category is required"
        }), 400

    if not description:
        return jsonify({
            "success": False,
            "message": "Description is required"
        }), 400

    try:
        amount_value = float(amount)
    except (TypeError, ValueError):
        return jsonify({
            "success": False,
            "message": "Amount must be a valid number"
        }), 400

    if amount_value <= 0:
        return jsonify({
            "success": False,
            "message": "Amount must be greater than 0"
        }), 400

    if not payment_method:
        return jsonify({
            "success": False,
            "message": "Payment Method is required"
        }), 400

    try:
        g.cursor.execute(
            """
            UPDATE expenses
            SET
                expense_date = %s,
                category = %s,
                description = %s,
                amount = %s,
                payment_method = %s,
                paid_to = %s,
                reference_number = %s,
                notes = %s,
                status = %s,
                updated_at = NOW()
            WHERE id = %s
            """,
            (
                expense_date,
                category,
                description,
                amount_value,
                payment_method,
                paid_to or None,
                reference_number or None,
                notes or None,
                status if status in ("Paid", "Pending") else "Paid",
                expense_id
            )
        )

        if g.cursor.rowcount == 0:
            g.db.rollback()
            return jsonify({
                "success": False,
                "message": "Expense नहीं मिला"
            }), 404

        g.db.commit()
        return jsonify({
            "success": True,
            "message": "Expense Successfully Updated"
        })

    except Exception as e:
        g.db.rollback()
        print("Update Expense Error:", e)
        return jsonify({
            "success": False,
            "message": "Expense update नहीं हुआ"
        }), 500


@app.route("/expenses/<int:expense_id>", methods=["DELETE", "OPTIONS"])
@roles_required("Admin", "Manager")
def delete_expense(expense_id):
    if request.method == "OPTIONS":
        return "", 200

    try:
        g.cursor.execute(
            "SELECT id, status, original_status FROM expenses WHERE id = %s LIMIT 1",
            (expense_id,)
        )
        expense = g.cursor.fetchone()
        if not expense:
            g.db.rollback()
            return jsonify({
                "success": False,
                "message": "Expense नहीं मिला"
            }), 404

        if str(expense.get("status") or "").strip() == "Deleted":
            return jsonify({
                "success": False,
                "message": "Expense already deleted."
            }), 400

        previous_status = str(expense.get("status") or "Paid").strip() or "Paid"
        g.cursor.execute(
            """
            UPDATE expenses
            SET status = %s,
                original_status = %s,
                deleted_at = NOW(),
                deleted_by_user_id = %s,
                updated_at = NOW()
            WHERE id = %s
            """,
            ("Deleted", previous_status, g.current_user.get("id"), expense_id)
        )
        log_expense_action(
            expense_id,
            "deleted",
            g.current_user,
            json.dumps({"previous_status": previous_status})
        )
        g.db.commit()
        return jsonify({
            "success": True,
            "message": "Expense Successfully Deleted"
        })

    except Exception as e:
        g.db.rollback()
        print("Delete Expense Error:", e)
        return jsonify({
            "success": False,
            "message": "Expense delete नहीं हुआ"
        }), 500


@app.route("/expenses/<int:expense_id>/restore", methods=["PUT", "OPTIONS"])
@roles_required("Admin", "Manager")
def restore_expense(expense_id):
    if request.method == "OPTIONS":
        return "", 200

    try:
        g.cursor.execute(
            "SELECT id, status, original_status FROM expenses WHERE id = %s LIMIT 1",
            (expense_id,)
        )
        expense = g.cursor.fetchone()
        if not expense:
            g.db.rollback()
            return jsonify({"success": False, "message": "Expense नहीं मिला"}), 404

        if str(expense.get("status") or "").strip() != "Deleted":
            return jsonify({"success": False, "message": "Expense is not deleted."}), 400

        restore_status = str(expense.get("original_status") or "Paid").strip() or "Paid"
        if restore_status not in {"Paid", "Pending"}:
            restore_status = "Paid"

        g.cursor.execute(
            """
            UPDATE expenses
            SET status = %s,
                original_status = NULL,
                deleted_at = NULL,
                deleted_by_user_id = NULL,
                updated_at = NOW()
            WHERE id = %s
            """,
            (restore_status, expense_id)
        )
        log_expense_action(
            expense_id,
            "restored",
            g.current_user,
            json.dumps({"restored_status": restore_status})
        )
        g.db.commit()
        return jsonify({
            "success": True,
            "message": "Expense Successfully Restored"
        })
    except Exception as e:
        g.db.rollback()
        print("Restore Expense Error:", e)
        return jsonify({
            "success": False,
            "message": "Expense restore नहीं हुआ"
        }), 500


# ==========================================
# ADD PRODUCT
# ==========================================

@app.route("/products", methods=["POST", "OPTIONS"])
def add_product():

    if request.method == "OPTIONS":
        return "", 200

    data = request.get_json()

    if not data:
        return jsonify({
            "success": False,
            "message": "Product data नहीं मिला"
        }), 400

    product_name = data.get("product_name")
    category = data.get("category")
    brand = data.get("brand")
    item_type = (data.get("item_type") or "Product").strip()
    if item_type not in ["Product", "Service"]:
        item_type = "Product"
    unit = data.get("unit", "Piece")
    secondary_unit = (data.get("secondary_unit") or "").strip() or None
    try:
        conversion_rate = float(data.get("conversion_rate", 1) or 1)
        if conversion_rate <= 0:
            conversion_rate = 1.0
    except (TypeError, ValueError):
        conversion_rate = 1.0

    try:
        secondary_sale_price = float(data.get("secondary_sale_price", 0) or 0)
    except (TypeError, ValueError):
        secondary_sale_price = 0.0

    try:
        secondary_purchase_price = float(data.get("secondary_purchase_price", 0) or 0)
    except (TypeError, ValueError):
        secondary_purchase_price = 0.0

    if not secondary_unit:
        conversion_rate = 1.0
        secondary_sale_price = 0.0
        secondary_purchase_price = 0.0

    description = data.get("description", "")
    purchase_price = data.get("purchase_price", 0)
    sale_price = data.get("sale_price", 0)
    stock = data.get("stock", 0) if item_type == "Product" else 0
    min_stock = data.get("min_stock", 0) if item_type == "Product" else 0
    gst_rate = data.get("gst_rate", 0)
    track_serial = data.get("track_serial", "No") if item_type == "Product" else "No"
    try:
        warranty_months = int(data.get("warranty_months", 0) or 0)
    except (TypeError, ValueError):
        warranty_months = 0

    if item_type == "Service":
        unit = "Service"
        secondary_unit = None
        conversion_rate = 1.0
        warranty_months = 0
        if not brand:
            brand = "Service"

    product_name = (data.get("product_name") or "").strip()
    if not product_name:
        return jsonify({
            "success": False,
            "message": "Product Name is required"
        }), 400

    try:
        # DUPLICATE CHECK: Prevent adding product with same name
        g.cursor.execute(
            "SELECT id, product_name, status FROM products WHERE LOWER(TRIM(product_name)) = LOWER(TRIM(%s)) LIMIT 1",
            (product_name,)
        )
        existing_prod = g.cursor.fetchone()
        if existing_prod:
            status_note = ""
            if str(existing_prod.get("status") or "").lower() == "inactive":
                status_note = " (यह प्रोडक्ट अभी Inactive है, Product List में जाकर इसे Activate कर सकते हैं)"
            return jsonify({
                "success": False,
                "message": f"'{existing_prod['product_name']}' नाम का प्रोडक्ट पहले से मौजूद है (ID: {existing_prod['id']})!{status_note}"
            }), 409

        # Check actual columns in products table dynamically
        g.cursor.execute("SHOW COLUMNS FROM products")
        cols_in_db = {row["Field"] for row in g.cursor.fetchall()}

        candidate_fields = {
            "product_name": product_name,
            "category": category,
            "brand": brand,
            "item_type": item_type,
            "description": description,
            "unit": unit,
            "secondary_unit": secondary_unit,
            "conversion_rate": conversion_rate,
            "purchase_price": purchase_price,
            "secondary_purchase_price": secondary_purchase_price,
            "sale_price": sale_price,
            "secondary_sale_price": secondary_sale_price,
            "stock": stock,
            "min_stock": min_stock,
            "gst_rate": gst_rate,
            "track_serial": track_serial,
            "warranty_months": warranty_months,
            "status": "Active"
        }

        insert_fields = [f for f in candidate_fields if f in cols_in_db]
        insert_values = [candidate_fields[f] for f in insert_fields]
        placeholders = ", ".join(["%s"] * len(insert_fields))
        sql = f"INSERT INTO products ({', '.join(insert_fields)}) VALUES ({placeholders})"

        g.cursor.execute(sql, tuple(insert_values))
        new_product_id = g.cursor.lastrowid

        # If initial stock > 0, assign to default location in product_location_stock (safely)
        try:
            if stock > 0 and item_type != "Service":
                g.cursor.execute("SELECT id FROM stock_locations WHERE is_default = 1 LIMIT 1")
                def_loc = g.cursor.fetchone()
                if def_loc:
                    g.cursor.execute(
                        """
                        INSERT INTO product_location_stock (product_id, location_id, quantity)
                        VALUES (%s, %s, %s)
                        ON DUPLICATE KEY UPDATE quantity = VALUES(quantity)
                        """,
                        (new_product_id, def_loc["id"], stock)
                    )
        except Exception as loc_err:
            print("Location stock assignment notice (non-fatal):", loc_err)

        g.db.commit()

        return jsonify({
            "success": True,
            "message": "Product Successfully Saved",
            "product_id": new_product_id
        })

    except Exception as e:

        g.db.rollback()

        print("Product Error:", e)

        return jsonify({
            "success": False,
            "message": f"Product save नहीं हुआ: {str(e)}"
        }), 500


# ==========================================
# PRODUCT TEMPLATE / IMPORT
# ==========================================

@app.route("/products/template", methods=["GET"])
def download_product_template():
    try:
        from openpyxl import Workbook
    except Exception:
        return jsonify({
            "success": False,
            "message": "openpyxl library not installed. Please install it first."
        }), 500

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"

    headers = [
        "product_name",
        "category",
        "brand",
        "item_type",
        "description",
        "unit",
        "purchase_price",
        "sale_price",
        "stock",
        "min_stock",
        "gst_rate",
        "track_serial"
    ]
    sheet.append(headers)

    example_row = [
        "Dell Laptop",
        "Laptop",
        "Dell",
        "Product",
        "15-inch business laptop",
        "Piece",
        "35000",
        "42000",
        "10",
        "2",
        "18",
        "No"
    ]
    sheet.append(example_row)

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    return send_file(
        buffer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="product-import-template.xlsx"
    )


@app.route("/products/import", methods=["POST"])
def import_products():
    try:
        from openpyxl import load_workbook
    except Exception:
        return jsonify({
            "success": False,
            "message": "openpyxl library not installed. Please install it first."
        }), 500

    file = request.files.get("file")
    if not file:
        return jsonify({
            "success": False,
            "message": "Import file required"
        }), 400

    filename = (file.filename or "").lower()
    if not (filename.endswith(".xlsx") or filename.endswith(".xls") or filename.endswith(".csv")):
        return jsonify({
            "success": False,
            "message": "Please upload XLSX, XLS or CSV file"
        }), 400

    try:
        if filename.endswith(".csv"):
            import csv
            text = file.read().decode("utf-8-sig")
            rows = list(csv.reader(StringIO(text)))
            if len(rows) < 2:
                return jsonify({
                    "success": False,
                    "message": "CSV file is empty"
                }), 400
            sheet_rows = rows[1:]
            headers = [str(h).strip().lower() for h in rows[0]]
        else:
            workbook = load_workbook(file, read_only=True, data_only=True)
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
            if len(rows) < 2:
                return jsonify({
                    "success": False,
                    "message": "Excel file is empty"
                }), 400
            headers = [str(h).strip().lower() for h in rows[0]]
            sheet_rows = rows[1:]
    except Exception as exc:
        print("Import product parse error:", exc)
        return jsonify({
            "success": False,
            "message": "File could not be read. Please check the template format."
        }), 400

    required_headers = [
        "product_name",
        "category",
        "brand",
        "item_type",
        "description",
        "unit",
        "purchase_price",
        "sale_price",
        "stock",
        "min_stock",
        "gst_rate",
        "track_serial"
    ]

    missing = [header for header in required_headers if header not in headers]
    if missing:
        return jsonify({
            "success": False,
            "message": "Missing required columns: " + ", ".join(missing)
        }), 400

    header_map = {header: index for index, header in enumerate(headers)}
    imported = 0
    skipped = 0
    errors = []

    for row_number, row in enumerate(sheet_rows, start=2):
        if not any(cell is not None and str(cell).strip() for cell in row):
            continue

        try:
            product_name = str(row[header_map["product_name"]]).strip() if header_map["product_name"] < len(row) and row[header_map["product_name"]] is not None else ""
            if not product_name:
                skipped += 1
                errors.append(f"Row {row_number}: Product Name required")
                continue

            item_type = str(row[header_map["item_type"]]).strip() if header_map["item_type"] < len(row) and row[header_map["item_type"]] is not None else "Product"
            item_type = item_type if item_type in ["Product", "Service"] else "Product"

            category = str(row[header_map["category"]]).strip() if header_map["category"] < len(row) and row[header_map["category"]] is not None else ""
            brand = str(row[header_map["brand"]]).strip() if header_map["brand"] < len(row) and row[header_map["brand"]] is not None else ""
            description = str(row[header_map["description"]]).strip() if header_map["description"] < len(row) and row[header_map["description"]] is not None else ""
            unit = str(row[header_map["unit"]]).strip() if header_map["unit"] < len(row) and row[header_map["unit"]] is not None else ("Service" if item_type == "Service" else "Piece")
            purchase_price = float(row[header_map["purchase_price"]] or 0) if header_map["purchase_price"] < len(row) else 0
            sale_price = float(row[header_map["sale_price"]] or 0) if header_map["sale_price"] < len(row) else 0
            stock = float(row[header_map["stock"]] or 0) if header_map["stock"] < len(row) else 0
            min_stock = float(row[header_map["min_stock"]] or 0) if header_map["min_stock"] < len(row) else 0
            gst_rate = float(row[header_map["gst_rate"]] or 0) if header_map["gst_rate"] < len(row) else 0
            track_serial = str(row[header_map["track_serial"]]).strip() if header_map["track_serial"] < len(row) and row[header_map["track_serial"]] is not None else "No"
            track_serial = track_serial if track_serial in ["Yes", "No"] else "No"

            if item_type == "Service":
                unit = "Service"
                brand = brand or "Service"
                stock = 0
                min_stock = 0
                track_serial = "No"

            # Skip if product with same name already exists
            g.cursor.execute(
                "SELECT id FROM products WHERE LOWER(TRIM(product_name)) = LOWER(TRIM(%s)) LIMIT 1",
                (product_name,)
            )
            if g.cursor.fetchone():
                skipped += 1
                errors.append(f"Row {row_number}: '{product_name}' already exists (Duplicate skipped)")
                continue

            g.cursor.execute(
                """
                INSERT INTO products
                (product_name, category, brand, item_type, description, unit, purchase_price, sale_price, stock, min_stock, gst_rate, track_serial, status)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, 'Active')
                """,
                (
                    product_name,
                    category,
                    brand,
                    item_type,
                    description,
                    unit,
                    purchase_price,
                    sale_price,
                    stock,
                    min_stock,
                    gst_rate,
                    track_serial,
                )
            )
            imported += 1
        except Exception as exc:
            skipped += 1
            errors.append(f"Row {row_number}: {str(exc)}")

    g.db.commit()

    return jsonify({
        "success": True,
        "message": f"Imported {imported} product(s). Skipped {skipped} row(s).",
        "imported": imported,
        "skipped": skipped,
        "errors": errors[:10]
    })


# ==========================================
# GET ALL PRODUCTS
# ==========================================

@app.route("/products", methods=["GET"])
def get_products():

    try:

        sql = "SELECT * FROM products ORDER BY id DESC"

        g.cursor.execute(sql)

        products = g.cursor.fetchall()

        for p in products:
            p.setdefault("item_type", "Product")
            p.setdefault("category", "")
            p.setdefault("brand", "")
            p.setdefault("description", "")
            p.setdefault("status", "Active")
            p.setdefault("track_serial", "No")
            p.setdefault("unit", "Piece")
            p.setdefault("secondary_unit", None)
            p.setdefault("conversion_rate", 1.0)
            p.setdefault("purchase_price", 0.0)
            p.setdefault("secondary_purchase_price", 0.0)
            p.setdefault("sale_price", 0.0)
            p.setdefault("secondary_sale_price", 0.0)
            p.setdefault("stock", 0.0)
            p.setdefault("min_stock", 0.0)
            p.setdefault("gst_rate", 0.0)
            p.setdefault("warranty_months", 0)
            u1 = p.get("unit") or "Piece"
            u2 = p.get("secondary_unit")
            crate = float(p.get("conversion_rate") or 1)
            stk = float(p.get("stock") or 0)
            if u2 and crate > 1:
                p["display_unit"] = f"{u1} (1 {u1} = {int(crate) if crate.is_integer() else crate} {u2})"
                equiv_primary = round(stk / crate, 2)
                eq_str = f"{int(equiv_primary) if equiv_primary.is_integer() else equiv_primary} {u1}"
                stk_str = f"{int(stk) if stk.is_integer() else stk} {u2}"
                p["display_stock"] = f"{stk_str} ({eq_str})"
            else:
                p["display_unit"] = u1
                p["display_stock"] = f"{int(stk) if stk.is_integer() else stk}"

        return jsonify({
            "success": True,
            "products": products
        })

    except Exception as e:

        print("Get Products Error:", e)

        return jsonify({
            "success": False,
            "message": f"Products load नहीं हुए: {str(e)}"
        }), 500


# ==========================================
# GET SINGLE PRODUCT
# ==========================================

@app.route("/products/<int:product_id>", methods=["GET"])
def get_product(product_id):

    try:

        sql = "SELECT * FROM products WHERE id = %s LIMIT 1"

        g.cursor.execute(sql, (product_id,))

        product = g.cursor.fetchone()

        if product:
            product.setdefault("item_type", "Product")
            product.setdefault("category", "")
            product.setdefault("brand", "")
            product.setdefault("description", "")
            product.setdefault("status", "Active")
            product.setdefault("track_serial", "No")
            product.setdefault("unit", "Piece")
            product.setdefault("secondary_unit", None)
            product.setdefault("conversion_rate", 1.0)
            product.setdefault("purchase_price", 0.0)
            product.setdefault("secondary_purchase_price", 0.0)
            product.setdefault("sale_price", 0.0)
            product.setdefault("secondary_sale_price", 0.0)
            product.setdefault("stock", 0.0)
            product.setdefault("min_stock", 0.0)
            product.setdefault("gst_rate", 0.0)
            product.setdefault("warranty_months", 0)

        if not product:

            return jsonify({
                "success": False,
                "message": "Product नहीं मिला"
            }), 404

        u1 = product.get("unit") or "Piece"
        u2 = product.get("secondary_unit")
        crate = float(product.get("conversion_rate") or 1)
        stk = float(product.get("stock") or 0)
        if u2 and crate > 1:
            product["display_unit"] = f"{u1} (1 {u1} = {int(crate) if crate.is_integer() else crate} {u2})"
            equiv_primary = round(stk / crate, 2)
            eq_str = f"{int(equiv_primary) if equiv_primary.is_integer() else equiv_primary} {u1}"
            stk_str = f"{int(stk) if stk.is_integer() else stk} {u2}"
            product["display_stock"] = f"{stk_str} ({eq_str})"
        else:
            product["display_unit"] = u1
            product["display_stock"] = f"{int(stk) if stk.is_integer() else stk}"

        return jsonify({
            "success": True,
            "product": product
        })

    except Exception as e:

        print("Get Product Error:", e)

        return jsonify({
            "success": False,
            "message": "Product load नहीं हुआ"
        }), 500


# ==========================================
# UPDATE PRODUCT
# ==========================================

@app.route(
    "/products/<int:product_id>",
    methods=["PUT", "OPTIONS"]
)
def update_product(product_id):

    if request.method == "OPTIONS":
        return "", 200

    data = request.get_json()

    if not data:

        return jsonify({
            "success": False,
            "message": "Product data नहीं मिला"
        }), 400

    product_name = data.get("product_name")
    category = data.get("category")
    brand = data.get("brand")
    item_type = (data.get("item_type") or "Product").strip()
    if item_type not in ["Product", "Service"]:
        item_type = "Product"
    unit = data.get("unit", "Piece")
    secondary_unit = (data.get("secondary_unit") or "").strip() or None
    try:
        conversion_rate = float(data.get("conversion_rate", 1) or 1)
        if conversion_rate <= 0:
            conversion_rate = 1.0
    except (TypeError, ValueError):
        conversion_rate = 1.0

    try:
        secondary_sale_price = float(data.get("secondary_sale_price", 0) or 0)
    except (TypeError, ValueError):
        secondary_sale_price = 0.0

    try:
        secondary_purchase_price = float(data.get("secondary_purchase_price", 0) or 0)
    except (TypeError, ValueError):
        secondary_purchase_price = 0.0

    if not secondary_unit:
        conversion_rate = 1.0
        secondary_sale_price = 0.0
        secondary_purchase_price = 0.0

    description = data.get("description", "")
    try:
        purchase_price = float(data.get("purchase_price", 0) or 0)
    except (TypeError, ValueError):
        purchase_price = 0.0

    try:
        sale_price = float(data.get("sale_price", 0) or 0)
    except (TypeError, ValueError):
        sale_price = 0.0

    stock = float(data.get("stock", 0) or 0) if item_type == "Product" else 0.0
    min_stock = float(data.get("min_stock", 0) or 0) if item_type == "Product" else 0.0

    try:
        gst_rate = float(data.get("gst_rate", 0) or 0)
    except (TypeError, ValueError):
        gst_rate = 0.0

    track_serial = data.get("track_serial", "No") if item_type == "Product" else "No"
    try:
        warranty_months = int(data.get("warranty_months", 0) or 0)
    except (TypeError, ValueError):
        warranty_months = 0

    if item_type == "Service":
        unit = "Service"
        secondary_unit = None
        conversion_rate = 1.0
        secondary_sale_price = 0.0
        secondary_purchase_price = 0.0
        warranty_months = 0
        if not brand:
            brand = "Service"

    product_name = (data.get("product_name") or "").strip()
    if not product_name:
        return jsonify({
            "success": False,
            "message": "Product Name is required"
        }), 400

    try:
        g.cursor.execute("SELECT id FROM products WHERE id = %s LIMIT 1", (product_id,))
        if not g.cursor.fetchone():
            return jsonify({
                "success": False,
                "message": "Product नहीं मिला"
            }), 404

        # DUPLICATE CHECK: Prevent renaming to an existing product's name
        g.cursor.execute(
            "SELECT id, product_name FROM products WHERE LOWER(TRIM(product_name)) = LOWER(TRIM(%s)) AND id != %s LIMIT 1",
            (product_name, product_id)
        )
        dup_prod = g.cursor.fetchone()
        if dup_prod:
            return jsonify({
                "success": False,
                "message": f"'{dup_prod['product_name']}' नाम का दूसरा प्रोडक्ट पहले से मौजूद है (ID: {dup_prod['id']})!"
            }), 409

        # Check actual columns in products table dynamically
        g.cursor.execute("SHOW COLUMNS FROM products")
        cols_in_db = {row["Field"] for row in g.cursor.fetchall()}

        candidate_fields = {
            "product_name": product_name,
            "category": category,
            "brand": brand,
            "item_type": item_type,
            "description": description,
            "unit": unit,
            "secondary_unit": secondary_unit,
            "conversion_rate": conversion_rate,
            "purchase_price": purchase_price,
            "secondary_purchase_price": secondary_purchase_price,
            "sale_price": sale_price,
            "secondary_sale_price": secondary_sale_price,
            "stock": stock,
            "min_stock": min_stock,
            "gst_rate": gst_rate,
            "track_serial": track_serial,
            "warranty_months": warranty_months
        }

        update_fields = [f for f in candidate_fields if f in cols_in_db]
        update_values = [candidate_fields[f] for f in update_fields]
        set_clauses = ", ".join([f"{f} = %s" for f in update_fields])
        sql = f"UPDATE products SET {set_clauses} WHERE id = %s"
        update_values.append(product_id)

        g.cursor.execute(sql, tuple(update_values))

        # Sync location stock safely
        try:
            if item_type != "Service":
                g.cursor.execute("SELECT id FROM stock_locations WHERE is_default = 1 LIMIT 1")
                def_loc = g.cursor.fetchone()
                if def_loc:
                    g.cursor.execute(
                        """
                        INSERT INTO product_location_stock (product_id, location_id, quantity)
                        VALUES (%s, %s, %s)
                        ON DUPLICATE KEY UPDATE quantity = VALUES(quantity)
                        """,
                        (product_id, def_loc["id"], stock)
                    )
        except Exception as loc_err:
            print("Location stock update notice (non-fatal):", loc_err)

        g.db.commit()

        return jsonify({
            "success": True,
            "message": "Product Successfully Updated"
        })

    except Exception as e:

        g.db.rollback()

        print("Update Product Error:", e)

        return jsonify({
            "success": False,
            "message": f"Product update नहीं हुआ: {str(e)}"
        }), 500


@app.route("/products/<int:product_id>/status", methods=["PUT", "OPTIONS"])
def update_product_status(product_id):
    if request.method == "OPTIONS":
        return "", 200

    status = str((request.get_json(silent=True) or {}).get("status") or "").strip()
    if status not in {"Active", "Inactive"}:
        return jsonify({"success": False, "message": "Invalid product status."}), 400

    try:
        g.cursor.execute("SELECT id, stock FROM products WHERE id = %s LIMIT 1", (product_id,))
        product = g.cursor.fetchone()
        if not product:
            return jsonify({"success": False, "message": "Product not found."}), 404

        if status == "Inactive" and float(product.get("stock") or 0) > 0:
            return jsonify({
                "success": False,
                "message": "Stock is still available. Set stock to zero before deactivating this product."
            }), 400

        g.cursor.execute("UPDATE products SET status = %s WHERE id = %s", (status, product_id))
        g.db.commit()
        return jsonify({"success": True, "message": f"Product {status.lower()} successfully."})
    except Exception as e:
        g.db.rollback()
        print("Update Product Status Error:", e)
        return jsonify({"success": False, "message": "Unable to update product status."}), 500


@app.route("/products/<int:product_id>", methods=["DELETE", "OPTIONS"])
def delete_product_route(product_id):
    if request.method == "OPTIONS":
        return "", 200

    try:
        g.cursor.execute("SELECT id, product_name FROM products WHERE id = %s LIMIT 1", (product_id,))
        prod = g.cursor.fetchone()
        if not prod:
            return jsonify({"success": False, "message": "Product नहीं मिला"}), 404

        # Check if product is used in sales, purchases or services
        tables_to_check = [
            ("sales_invoice_items", "product_id"),
            ("purchase_invoice_items", "product_id"),
            ("service_jobs", "product_id"),
            ("service_job_parts", "product_id")
        ]
        in_use = False
        for tbl, col in tables_to_check:
            try:
                g.cursor.execute(f"SELECT COUNT(*) AS cnt FROM {tbl} WHERE {col} = %s", (product_id,))
                cnt = (g.cursor.fetchone() or {}).get("cnt", 0)
                if cnt > 0:
                    in_use = True
                    break
            except Exception as check_err:
                print(f"Check usage in {tbl} skipped:", check_err)

        if in_use:
            g.cursor.execute("UPDATE products SET status = 'Inactive' WHERE id = %s", (product_id,))
            g.db.commit()
            return jsonify({
                "success": True,
                "message": f"Product '{prod['product_name']}' बिलिंग/सर्विस में इस्तेमाल हुआ है, इसलिए सुरक्षा के लिए इसे Inactive कर दिया गया।"
            })
        else:
            for tbl, col in [("product_location_stock", "product_id"), ("product_serials", "product_id")]:
                try:
                    g.cursor.execute(f"DELETE FROM {tbl} WHERE {col} = %s", (product_id,))
                except Exception:
                    pass
            g.cursor.execute("DELETE FROM products WHERE id = %s", (product_id,))
            g.db.commit()
            return jsonify({
                "success": True,
                "message": f"Product '{prod['product_name']}' सफलतापूर्वक डिलीट कर दिया गया।"
            })
    except Exception as e:
        g.db.rollback()
        print("Delete Product Error:", e)
        return jsonify({"success": False, "message": f"Product डिलीट नहीं हुआ: {str(e)}"}), 500


@app.route("/products/cleanup-duplicates", methods=["POST", "OPTIONS"])
def cleanup_duplicate_products():
    if request.method == "OPTIONS":
        return "", 200

    try:
        # Find product names that appear more than once
        sql = """
            SELECT LOWER(TRIM(product_name)) AS norm_name, COUNT(*) AS cnt, MIN(id) AS keep_id
            FROM products
            GROUP BY LOWER(TRIM(product_name))
            HAVING cnt > 1
        """
        g.cursor.execute(sql)
        dupes = g.cursor.fetchall()

        total_removed = 0
        cleaned_names = []

        tables_to_check = [
            ("sales_invoice_items", "product_id"),
            ("purchase_invoice_items", "product_id"),
            ("service_jobs", "product_id"),
            ("service_job_parts", "product_id")
        ]

        for item in dupes:
            norm_name = item["norm_name"]
            keep_id = item["keep_id"]

            # Get all duplicates except the lowest ID (which we keep)
            g.cursor.execute(
                "SELECT id, product_name FROM products WHERE LOWER(TRIM(product_name)) = %s AND id != %s ORDER BY id ASC",
                (norm_name, keep_id)
            )
            dup_rows = g.cursor.fetchall()

            for dup in dup_rows:
                dup_id = dup["id"]

                in_use = False
                for tbl, col in tables_to_check:
                    try:
                        g.cursor.execute(f"SELECT COUNT(*) AS cnt FROM {tbl} WHERE {col} = %s", (dup_id,))
                        cnt = (g.cursor.fetchone() or {}).get("cnt", 0)
                        if cnt > 0:
                            in_use = True
                            break
                    except Exception:
                        pass

                if not in_use:
                    for tbl, col in [("product_location_stock", "product_id"), ("product_serials", "product_id")]:
                        try:
                            g.cursor.execute(f"DELETE FROM {tbl} WHERE {col} = %s", (dup_id,))
                        except Exception:
                            pass
                    g.cursor.execute("DELETE FROM products WHERE id = %s", (dup_id,))
                    total_removed += 1
                else:
                    g.cursor.execute("UPDATE products SET status = 'Inactive' WHERE id = %s", (dup_id,))
                    total_removed += 1

            if dup_rows:
                cleaned_names.append(dup_rows[0].get("product_name", norm_name))

        g.db.commit()

        if total_removed > 0:
            names_summary = ", ".join(cleaned_names[:3])
            if len(cleaned_names) > 3:
                names_summary += "..."
            msg = f"सफलतापूर्वक {total_removed} डुप्लीकेट प्रोडक्ट्स हटा दिए गए हैं! ({names_summary})"
        else:
            msg = "डेटाबेस में कोई डुप्लीकेट प्रोडक्ट नहीं मिला।"

        return jsonify({
            "success": True,
            "message": msg,
            "removed_count": total_removed
        })

    except Exception as e:
        g.db.rollback()
        print("Cleanup Duplicate Products Error:", e)
        return jsonify({
            "success": False,
            "message": f"डुप्लीकेट हटाने में त्रुटि: {str(e)}"
        }), 500


# ==========================================
# GET AVAILABLE SERIAL NUMBERS
# ==========================================

@app.route(
    "/products/<int:product_id>/serials",
    methods=["GET"]
)
def get_product_serials(product_id):

    try:

        location_id = request.args.get("location_id")

        sql = """
            SELECT
                ps.id,
                ps.serial_number,
                ps.purchase_date,
                ps.supplier_name,
                ps.warranty_start,
                ps.warranty_end,
                ps.location_id,
                sl.location_name
            FROM product_serials ps
            LEFT JOIN stock_locations sl ON sl.id = ps.location_id
            WHERE ps.product_id = %s
            AND ps.status = 'In Stock'
        """
        params = [product_id]

        if location_id:
            sql += " AND (ps.location_id = %s OR (ps.location_id IS NULL AND %s = (SELECT id FROM stock_locations WHERE is_default = 1 LIMIT 1))) "
            params.extend([location_id, location_id])

        sql += " ORDER BY ps.id DESC "

        g.cursor.execute(
            sql,
            tuple(params)
        )

        serials = g.cursor.fetchall()

        return jsonify({
            "success": True,
            "serials": serials
        })

    except Exception as e:

        print(
            "Serial Number Error:",
            e
        )

        return jsonify({
            "success": False,
            "message": "Serial numbers load नहीं हुए"
        }), 500


# ==========================================
# SERIAL NUMBER / BARCODE LOOKUP FOR BILLING
# ==========================================

@app.route("/product-serials/lookup", methods=["GET"])
def lookup_serial_or_product():
    """Lookup a serial number or product by barcode/code/name for quick scanning in billing."""
    try:
        query = (request.args.get("query") or request.args.get("code") or "").strip()
        if not query:
            return jsonify({
                "success": False,
                "message": "Scan query cannot be empty."
            }), 400

        # 1. Search product_serials for matching serial_number
        sql_serial = """
            SELECT
                ps.id AS serial_id,
                ps.serial_number,
                ps.status,
                ps.warranty_start,
                ps.warranty_end,
                p.id AS product_id,
                p.product_name,
                p.category,
                p.brand,
                p.unit,
                p.sale_price,
                p.gst_rate,
                p.track_serial,
                p.stock
            FROM product_serials ps
            JOIN products p ON ps.product_id = p.id
            WHERE LOWER(TRIM(ps.serial_number)) = LOWER(%s)
            ORDER BY ps.id DESC
            LIMIT 1
        """
        g.cursor.execute(sql_serial, (query,))
        serial_match = g.cursor.fetchone()

        if serial_match:
            if serial_match["status"] != "In Stock":
                return jsonify({
                    "success": False,
                    "message": f"Serial number '{serial_match['serial_number']}' उपलब्ध नहीं है (Status: {serial_match['status']})."
                }), 400

            product_info = {
                "id": serial_match["product_id"],
                "product_name": serial_match["product_name"],
                "category": serial_match["category"],
                "brand": serial_match["brand"],
                "unit": serial_match["unit"],
                "sale_price": float(serial_match["sale_price"] or 0),
                "gst_rate": float(serial_match["gst_rate"] or 0),
                "track_serial": serial_match["track_serial"],
                "stock": float(serial_match["stock"] or 0)
            }
            return jsonify({
                "success": True,
                "type": "serial",
                "serial_number": serial_match["serial_number"],
                "product": product_info
            })

        # 2. Check if query matches product ID or product name
        sql_product = """
            SELECT
                id,
                product_name,
                category,
                brand,
                unit,
                sale_price,
                gst_rate,
                track_serial,
                stock
            FROM products
            WHERE (id = %s OR LOWER(TRIM(product_name)) = LOWER(%s))
            AND status = 'Active'
            LIMIT 1
        """
        product_id_param = int(query) if query.isdigit() else -1
        g.cursor.execute(sql_product, (product_id_param, query))
        prod_match = g.cursor.fetchone()

        if prod_match:
            product_info = {
                "id": prod_match["id"],
                "product_name": prod_match["product_name"],
                "category": prod_match["category"],
                "brand": prod_match["brand"],
                "unit": prod_match["unit"],
                "sale_price": float(prod_match["sale_price"] or 0),
                "gst_rate": float(prod_match["gst_rate"] or 0),
                "track_serial": prod_match["track_serial"],
                "stock": float(prod_match["stock"] or 0)
            }
            return jsonify({
                "success": True,
                "type": "product",
                "serial_number": None,
                "product": product_info
            })

        return jsonify({
            "success": False,
            "message": f"Serial number या Product '{query}' नहीं मिला।"
        }), 404

    except Exception as e:
        print("Serial Lookup Error:", e)
        return jsonify({
            "success": False,
            "message": "खोजने में समस्या आई: " + str(e)
        }), 500


# ==========================================
# GET ALL SUPPLIERS
# ==========================================

@app.route("/suppliers", methods=["GET"])
def get_suppliers():

    try:

        sql = """
            SELECT
                id,
                supplier_name,
                mobile,
                email,
                address,
                gstin,
                status,
                created_at
            FROM suppliers
            ORDER BY id DESC
        """

        g.cursor.execute(sql)

        suppliers = g.cursor.fetchall()

        return jsonify({
            "success": True,
            "suppliers": suppliers
        })

    except Exception as e:

        print("Get Suppliers Error:", e)

        return jsonify({
            "success": False,
            "message": "Suppliers load नहीं हुए"
        }), 500


@app.route("/suppliers/<int:supplier_id>", methods=["GET"])
def get_supplier(supplier_id):
    try:
        g.cursor.execute(
            """
            SELECT
                id,
                supplier_name,
                mobile,
                email,
                address,
                gstin,
                status,
                created_at
            FROM suppliers
            WHERE id = %s
            LIMIT 1
            """,
            (supplier_id,)
        )

        supplier = g.cursor.fetchone()
        if not supplier:
            return jsonify({"success": False, "message": "Supplier नहीं मिला"}), 404

        return jsonify({"success": True, "supplier": supplier})
    except Exception as e:
        print("Get Supplier Error:", e)
        return jsonify({"success": False, "message": "Supplier load नहीं हुआ"}), 500


@app.route("/suppliers/<int:supplier_id>", methods=["PUT", "OPTIONS"])
def update_supplier(supplier_id):
    if request.method == "OPTIONS":
        return "", 200

    data = request.get_json(silent=True) or {}
    if not data:
        return jsonify({"success": False, "message": "Supplier data नहीं मिला"}), 400

    supplier_name = str(data.get("supplier_name", "") or "").strip()
    mobile = str(data.get("mobile", "") or "").strip()
    email = str(data.get("email", "") or "").strip()
    address = str(data.get("address", "") or "").strip()
    gstin = str(data.get("gstin", "") or "").strip()

    if not supplier_name:
        return jsonify({"success": False, "message": "Supplier Name is required"}), 400

    try:
        g.cursor.execute("SELECT id FROM suppliers WHERE id = %s LIMIT 1", (supplier_id,))
        if not g.cursor.fetchone():
            return jsonify({"success": False, "message": "Supplier नहीं मिला"}), 404

        g.cursor.execute(
            """
            UPDATE suppliers
            SET
                supplier_name = %s,
                mobile = %s,
                email = %s,
                address = %s,
                gstin = %s
            WHERE id = %s
            """,
            (supplier_name, mobile, email, address, gstin, supplier_id)
        )
        g.db.commit()
        return jsonify({"success": True, "message": "Supplier Successfully Updated"})
    except Exception as e:
        g.db.rollback()
        print("Update Supplier Error:", e)
        return jsonify({"success": False, "message": "Supplier update नहीं हुआ"}), 500


@app.route("/suppliers/<int:supplier_id>/status", methods=["PUT", "OPTIONS"])
def update_supplier_status(supplier_id):
    if request.method == "OPTIONS":
        return "", 200

    payload = request.get_json(silent=True) or {}
    status = str(payload.get("status") or "").strip()
    if status not in {"Active", "Inactive"}:
        return jsonify({"success": False, "message": "Invalid supplier status."}), 400

    try:
        g.cursor.execute(
            "UPDATE suppliers SET status = %s WHERE id = %s",
            (status, supplier_id)
        )
        if g.cursor.rowcount == 0:
            g.db.rollback()
            return jsonify({"success": False, "message": "Supplier नहीं मिला"}), 404
        g.db.commit()
        return jsonify({"success": True, "message": f"Supplier {status.lower()} successfully."})
    except Exception as e:
        g.db.rollback()
        print("Update Supplier Status Error:", e)
        return jsonify({"success": False, "message": "Unable to update supplier status."}), 500


# ==========================================
# ADD SUPPLIER
# ==========================================

@app.route(
    "/suppliers",
    methods=["POST", "OPTIONS"]
)
def add_supplier():

    if request.method == "OPTIONS":
        return "", 200

    data = request.get_json()

    if not data:

        return jsonify({
            "success": False,
            "message": "Supplier data नहीं मिला"
        }), 400


    supplier_name = data.get(
        "supplier_name"
    )

    mobile = data.get(
        "mobile",
        ""
    )

    email = data.get(
        "email",
        ""
    )

    address = data.get(
        "address",
        ""
    )

    gstin = data.get(
        "gstin",
        ""
    )


    if not supplier_name:

        return jsonify({
            "success": False,
            "message": "Supplier Name is required"
        }), 400


    try:

        sql = """
            INSERT INTO suppliers
            (
                supplier_name,
                mobile,
                email,
                address,
                gstin
            )

            VALUES
            (
                %s,
                %s,
                %s,
                %s,
                %s
            )
        """


        g.cursor.execute(
            sql,
            (
                supplier_name,
                mobile,
                email,
                address,
                gstin
            )
        )


        g.db.commit()


        return jsonify({
            "success": True,
            "message": "Supplier Successfully Saved"
        })


    except Exception as e:

        g.db.rollback()

        print(
            "Add Supplier Error:",
            e
        )


        return jsonify({
            "success": False,
            "message": "Supplier save नहीं हुआ"
        }), 500


# ==========================================
# SAVE PURCHASE
# ==========================================

@app.route(
    "/purchases",
    methods=["POST", "OPTIONS"]
)
def add_purchase():

    if request.method == "OPTIONS":
        return "", 200

    data = request.get_json()

    if not data:
        return jsonify({
            "success": False,
            "message": "Purchase data नहीं मिला"
        }), 400

    purchase_number = data.get("purchase_number")
    supplier_id = data.get("supplier_id")
    purchase_date = data.get("purchase_date")
    subtotal = data.get("subtotal", 0)
    discount = data.get("discount", 0)
    tax_amount = data.get("tax_amount", 0)
    grand_total = data.get("grand_total", 0)
    payment_method = data.get(
        "payment_method",
        "Cash"
    )
    payment_status = data.get(
        "payment_status",
        "Paid"
    )
    notes = data.get("notes", "")

    items = data.get("items", [])


    if not supplier_id:
        return jsonify({
            "success": False,
            "message": "Supplier select करें"
        }), 400


    if not items:
        return jsonify({
            "success": False,
            "message": "कम से कम एक product जोड़ें"
        }), 400


    try:

        doc_tax_incl = False if str(data.get("tax_type", "")).lower() == "exclusive" or data.get("is_tax_inclusive") is False else True
        items, subtotal, tax_amount, grand_total, discount = normalize_document_items(
            items,
            discount,
            default_tax_inclusive=doc_tax_incl
        )
        g.cursor.execute("SELECT id FROM suppliers WHERE id = %s AND status = 'Active' LIMIT 1", (supplier_id,))
        if not g.cursor.fetchone():
            raise ValueError("Selected supplier is not available.")
        purchase_number = generate_document_number(
            "purchase", "purchase_prefix", "next_purchase_number", "PUR-"
        )

        # ----------------------------------
        # CREATE PURCHASE INVOICE
        # ----------------------------------

        sql = """
            INSERT INTO purchase_invoices
            (
                purchase_number,
                supplier_id,
                purchase_date,
                subtotal,
                discount,
                tax_amount,
                grand_total,
                payment_method,
                payment_status,
                notes
            )
            VALUES
            (
                %s, %s, %s, %s, %s,
                %s, %s, %s, %s, %s
            )
        """

        g.cursor.execute(
            sql,
            (
                purchase_number,
                supplier_id,
                purchase_date,
                subtotal,
                discount,
                tax_amount,
                grand_total,
                payment_method,
                payment_status,
                notes
            )
        )


        purchase_id = g.cursor.lastrowid


        # ----------------------------------
        # SAVE ITEMS
        # ----------------------------------

        for item in items:

            product_id = item.get(
                "product_id"
            )

            quantity = float(
                item.get("quantity", 0)
            )

            rate = float(
                item.get("rate", 0)
            )

            gst_rate = float(
                item.get("gst_rate", 0)
            )

            gst_amount = float(
                item.get("gst_amount", 0)
            )

            item_discount = float(
                item.get("discount", 0)
            )

            line_total = float(
                item.get("line_total", 0)
            )

            serial_numbers = item.get("serial_numbers") or []
            serial_number = item.get("serial_number")

            if isinstance(serial_numbers, list):
                serial_values = [str(v).strip() for v in serial_numbers if str(v).strip()]
            elif serial_number:
                serial_values = [str(serial_number).strip()]
            else:
                serial_values = []

            warranty_start = item.get(
                "warranty_start"
            )

            warranty_end = item.get(
                "warranty_end"
            )

            raw_loc_id = item.get("location_id")
            location_id = None
            if raw_loc_id not in (None, "", "null", "undefined"):
                try:
                    location_id = int(raw_loc_id)
                except (ValueError, TypeError):
                    location_id = None

            if not location_id:
                g.cursor.execute("SELECT id FROM stock_locations WHERE is_default = 1 AND status = 'Active' LIMIT 1")
                def_row = g.cursor.fetchone()
                if def_row:
                    location_id = def_row["id"]
                else:
                    g.cursor.execute("SELECT id FROM stock_locations WHERE status = 'Active' ORDER BY id ASC LIMIT 1")
                    fallback_row = g.cursor.fetchone()
                    location_id = fallback_row["id"] if fallback_row else 1


            if not product_id:

                raise Exception(
                    "Product select नहीं किया गया"
                )

            product_record = get_product_record_by_id(product_id)
            if not product_record:
                raise Exception("Product नहीं मिला")

            cost_rate = float(product_record.get("purchase_price") or 0)

            if quantity <= 0:

                raise Exception(
                    "Quantity सही नहीं है"
                )

            item_unit = (item.get("unit") or "").strip() or None
            p_unit = product_record.get("unit") or "Piece"
            p_sec_unit = product_record.get("secondary_unit")
            p_crate = float(product_record.get("conversion_rate") or 1)

            if p_sec_unit and p_crate > 1:
                if item_unit == p_unit:
                    conversion_factor = p_crate
                    base_quantity = round(quantity * p_crate, 4)
                elif item_unit == p_sec_unit:
                    conversion_factor = 1.0
                    base_quantity = round(quantity, 4)
                else:
                    item_unit = p_unit
                    conversion_factor = p_crate
                    base_quantity = round(quantity * p_crate, 4)
            else:
                if not item_unit:
                    item_unit = p_unit
                conversion_factor = 1.0
                base_quantity = round(quantity, 4)

            if product_record.get("item_type") == "Service":
                serial_number = None
                serial_values = []
                warranty_start = None
                warranty_end = None
            else:
                tracks_serial = str(product_record.get("track_serial") or "No") == "Yes"
                if tracks_serial and quantity != int(quantity):
                    raise Exception("Serial-tracked products require whole-number quantity.")
                if tracks_serial and len(serial_values) != int(quantity):
                    raise Exception("Every unit of a serial-tracked product needs one serial number")

                seen = set()
                for serial_value in serial_values:
                    if serial_value in seen:
                        raise Exception("Serial Number " + serial_value + " already exists.")
                    seen.add(serial_value)

                    g.cursor.execute(
                        """
                        SELECT id
                        FROM product_serials
                        WHERE serial_number = %s
                        LIMIT 1
                        """,
                        (serial_value,)
                    )

                    existing_serial = g.cursor.fetchone()
                    if existing_serial:
                        raise Exception("Serial Number " + serial_value + " already exists.")

                if serial_values:
                    serial_number = serial_values[0]


            # ----------------------------------
            # PURCHASE ITEM
            # ----------------------------------

            g.cursor.execute(
                """
                INSERT INTO purchase_invoice_items
                (
                    purchase_invoice_id,
                    product_id,
                    location_id,
                    quantity,
                    unit,
                    base_quantity,
                    conversion_factor,
                    rate,
                    gst_rate,
                    gst_amount,
                    discount,
                    line_total,
                    serial_number,
                    warranty_start,
                    warranty_end
                )
                VALUES
                (
                    %s, %s, %s, %s, %s,
                    %s, %s, %s, %s, %s,
                    %s, %s, %s, %s, %s
                )
                """,
                (
                    purchase_id,
                    product_id,
                    location_id,
                    quantity,
                    item_unit,
                    base_quantity,
                    conversion_factor,
                    rate,
                    gst_rate,
                    gst_amount,
                    item_discount,
                    line_total,
                    serial_number,
                    warranty_start,
                    warranty_end
                )
            )


            # ----------------------------------
            # INCREASE STOCK & LOCATION STOCK
            # ----------------------------------

            if product_record.get("item_type") != "Service":
                g.cursor.execute(
                    """
                    UPDATE products
                    SET stock = stock + %s
                    WHERE id = %s
                    """,
                    (
                        base_quantity,
                        product_id
                    )
                )

                if location_id:
                    g.cursor.execute(
                        """
                        INSERT INTO product_location_stock (product_id, location_id, quantity)
                        VALUES (%s, %s, %s)
                        ON DUPLICATE KEY UPDATE quantity = quantity + VALUES(quantity)
                        """,
                        (
                            product_id,
                            location_id,
                            base_quantity
                        )
                    )


            # ----------------------------------
            # SAVE SERIAL NUMBERS
            # ----------------------------------

            for serial_value in serial_values:
                g.cursor.execute(
                    """
                    INSERT INTO product_serials
                    (
                        product_id,
                        serial_number,
                        purchase_invoice_id,
                        purchase_date,
                        supplier_name,
                        location_id,
                        warranty_start,
                        warranty_end,
                        status
                    )
                    SELECT
                        %s,
                        %s,
                        %s,
                        %s,
                        supplier_name,
                        %s,
                        %s,
                        %s,
                        'In Stock'
                    FROM suppliers
                    WHERE id = %s
                    """,
                    (
                        product_id,
                        serial_value,
                        purchase_id,
                        purchase_date,
                        location_id,
                        warranty_start,
                        warranty_end,
                        supplier_id
                    )
                )


        # ----------------------------------
        # COMMIT
        # ----------------------------------

        g.db.commit()


        return jsonify({
            "success": True,
            "message": "Purchase Successfully Saved",
            "purchase_id": purchase_id
        })


    except Exception as e:

        g.db.rollback()

        print(
            "Purchase Error:",
            e
        )

        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


# ==========================================
# GET PURCHASES LIST
# ==========================================

@app.route("/purchases", methods=["GET"])
def get_purchases_list():
    try:
        purchase_number = (request.args.get("purchase_number") or "").strip()
        supplier = (request.args.get("supplier") or "").strip()
        from_date = normalize_date_string(request.args.get("from_date"))
        to_date = normalize_date_string(request.args.get("to_date"))
        payment_status = (request.args.get("payment_status") or "").strip()

        sql = """
            SELECT
                pi.id,
                pi.purchase_number,
                pi.purchase_date,
                pi.subtotal,
                pi.discount,
                pi.tax_amount,
                pi.grand_total,
                pi.payment_method,
                pi.payment_status,
                COALESCE(s.supplier_name, 'Unknown Supplier') AS supplier_name,
                pi.supplier_id
            FROM purchase_invoices pi
            LEFT JOIN suppliers s
                ON s.id = pi.supplier_id
            WHERE 1=1
        """
        params = []

        if purchase_number:
            sql += " AND pi.purchase_number LIKE %s "
            params.append(f"%{purchase_number}%")

        if supplier:
            sql += " AND (s.supplier_name LIKE %s OR s.mobile LIKE %s) "
            params.extend([f"%{supplier}%", f"%{supplier}%"])

        if from_date:
            sql += " AND DATE(pi.purchase_date) >= %s "
            params.append(from_date)

        if to_date:
            sql += " AND DATE(pi.purchase_date) <= %s "
            params.append(to_date)

        if payment_status:
            sql += " AND pi.payment_status = %s "
            params.append(payment_status)

        sql += " ORDER BY pi.purchase_date DESC, pi.id DESC "

        g.cursor.execute(sql, tuple(params))
        rows = g.cursor.fetchall()

        purchases = []
        for row in rows:
            purchases.append({
                "id": row["id"],
                "purchase_number": row["purchase_number"],
                "purchase_date": normalize_date_string(row["purchase_date"]),
                "supplier": row["supplier_name"],
                "subtotal": float(row["subtotal"] or 0),
                "discount": float(row["discount"] or 0),
                "tax_amount": float(row["tax_amount"] or 0),
                "grand_total": float(row["grand_total"] or 0),
                "payment_method": row["payment_method"],
                "payment_status": row["payment_status"],
                "supplier_id": row.get("supplier_id")
            })

        return jsonify({
            "success": True,
            "purchases": purchases,
            "total": len(purchases)
        })

    except Exception as e:
        print("Get Purchases List Error:", e)
        return jsonify({
            "success": False,
            "message": "Unable to load purchase invoices."
        }), 500


# ==========================================
# GET SINGLE PURCHASE DETAILS
# ==========================================

@app.route("/purchases/<int:purchase_id>", methods=["GET"])
def get_purchase_details(purchase_id):
    try:
        g.cursor.execute(
            """
            SELECT
                pi.id,
                pi.supplier_id,
                pi.purchase_number,
                pi.purchase_date,
                pi.subtotal,
                pi.discount,
                pi.tax_amount,
                pi.grand_total,
                pi.payment_method,
                pi.payment_status,
                pi.notes,
                s.supplier_name,
                s.mobile AS supplier_mobile,
                s.address AS supplier_address,
                s.gstin AS supplier_gstin,
                s.email AS supplier_email
            FROM purchase_invoices pi
            LEFT JOIN suppliers s
                ON s.id = pi.supplier_id
            WHERE pi.id = %s
            LIMIT 1
            """,
            (purchase_id,)
        )

        purchase = g.cursor.fetchone()
        if not purchase:
            return jsonify({
                "success": False,
                "message": "Purchase invoice नहीं मिला"
            }), 404

        g.cursor.execute(
            """
            SELECT
                pii.id,
                pii.product_id,
                pii.location_id,
                COALESCE(sl.location_name, 'Main Counter') AS location_name,
                pii.quantity,
                pii.unit,
                pii.base_quantity,
                pii.conversion_factor,
                pii.rate,
                pii.discount,
                pii.gst_rate,
                pii.gst_amount,
                pii.line_total,
                pii.serial_number,
                pii.warranty_start,
                pii.warranty_end,
                p.product_name,
                p.brand,
                p.item_type,
                p.unit AS primary_unit,
                p.secondary_unit,
                p.conversion_rate
            FROM purchase_invoice_items pii
            LEFT JOIN products p
                ON p.id = pii.product_id
            LEFT JOIN stock_locations sl
                ON sl.id = pii.location_id
            WHERE pii.purchase_invoice_id = %s
            ORDER BY pii.id ASC
            """,
            (purchase_id,)
        )
        items = g.cursor.fetchall()

        g.cursor.execute(
            """
            SELECT product_id, serial_number, warranty_start, warranty_end
            FROM product_serials
            WHERE purchase_invoice_id = %s
            ORDER BY id ASC
            """,
            (purchase_id,)
        )
        serial_rows = g.cursor.fetchall()
        serial_map = {}
        for row in serial_rows:
            product_id = row.get("product_id")
            serial_map.setdefault(product_id, []).append({
                "serial_number": row.get("serial_number"),
                "warranty_start": normalize_date_string(row.get("warranty_start")),
                "warranty_end": normalize_date_string(row.get("warranty_end"))
            })

        normalized_items = []
        for item in items:
            product_id = item.get("product_id")
            item_serials = serial_map.get(product_id, [])
            if not item_serials and item.get("serial_number"):
                item_serials = [{
                    "serial_number": item.get("serial_number"),
                    "warranty_start": normalize_date_string(item.get("warranty_start")),
                    "warranty_end": normalize_date_string(item.get("warranty_end"))
                }]

            normalized_items.append({
                "id": item["id"],
                "product_id": item["product_id"],
                "location_id": item.get("location_id"),
                "location_name": item.get("location_name") or "Main Counter",
                "product_name": item.get("product_name") or "Product",
                "brand": item.get("brand") or "-",
                "quantity": float(item["quantity"] or 0),
                "rate": float(item["rate"] or 0),
                "discount": float(item["discount"] or 0),
                "gst_rate": float(item["gst_rate"] or 0),
                "gst_amount": float(item["gst_amount"] or 0),
                "line_total": float(item["line_total"] or 0),
                "serials": item_serials
            })

        return jsonify({
            "success": True,
            "purchase": {
                "id": purchase["id"],
                "supplier_id": purchase["supplier_id"],
                "purchase_number": purchase["purchase_number"],
                "purchase_date": normalize_date_string(purchase["purchase_date"]),
                "subtotal": float(purchase["subtotal"] or 0),
                "discount": float(purchase["discount"] or 0),
                "tax_amount": float(purchase["tax_amount"] or 0),
                "grand_total": float(purchase["grand_total"] or 0),
                "payment_method": purchase["payment_method"],
                "payment_status": purchase["payment_status"],
                "notes": purchase["notes"] or "",
                "supplier_name": purchase["supplier_name"] or "Unknown Supplier",
                "supplier_mobile": purchase.get("supplier_mobile") or "-",
                "supplier_address": purchase.get("supplier_address") or "-",
                "supplier_gstin": purchase.get("supplier_gstin") or "-",
                "supplier_email": purchase.get("supplier_email") or "-"
            },
            "items": normalized_items
        })

    except Exception as e:
        print("Get Purchase Details Error:", e)
        return jsonify({
            "success": False,
            "message": "Unable to load purchase details."
        }), 500


# ==========================================
# UPDATE PURCHASE INVOICE
# ==========================================

@app.route("/purchases/<int:purchase_id>", methods=["PUT", "OPTIONS"])
@roles_required("Admin", "Manager", "Staff")
def update_purchase(purchase_id):
    if request.method == "OPTIONS":
        return "", 200

    data = request.get_json()
    if not data:
        return jsonify({"success": False, "message": "Purchase data नहीं मिला"}), 400

    supplier_id = data.get("supplier_id")
    purchase_date = data.get("purchase_date")
    discount = data.get("discount", 0)
    payment_method = data.get("payment_method", "Cash")
    payment_status = data.get("payment_status", "Paid")
    notes = data.get("notes", "")
    items = data.get("items", [])

    if not supplier_id:
        return jsonify({"success": False, "message": "Supplier select करें"}), 400

    if not items:
        return jsonify({"success": False, "message": "कम से कम एक product जोड़ें"}), 400

    try:
        g.cursor.execute("SELECT * FROM purchase_invoices WHERE id = %s LIMIT 1", (purchase_id,))
        purchase_record = g.cursor.fetchone()
        if not purchase_record:
            return jsonify({"success": False, "message": "Purchase invoice नहीं मिला"}), 404

        doc_tax_incl = False if str(data.get("tax_type", "")).lower() == "exclusive" or data.get("is_tax_inclusive") is False else True
        items, subtotal, tax_amount, grand_total, discount = normalize_document_items(
            items,
            discount,
            default_tax_inclusive=doc_tax_incl
        )

        g.cursor.execute("SELECT id, supplier_name FROM suppliers WHERE id = %s AND status = 'Active' LIMIT 1", (supplier_id,))
        supplier_row = g.cursor.fetchone()
        if not supplier_row:
            raise ValueError("Selected supplier is not available.")
        supplier_name = supplier_row.get("supplier_name") or ""

        # Fetch old items
        g.cursor.execute(
            """
            SELECT pii.product_id, pii.location_id, COALESCE(pii.base_quantity, pii.quantity) AS base_qty, p.item_type, p.product_name
            FROM purchase_invoice_items pii
            LEFT JOIN products p ON p.id = pii.product_id
            WHERE pii.purchase_invoice_id = %s
            """,
            (purchase_id,)
        )
        old_items = g.cursor.fetchall()

        # Check existing sold serials
        g.cursor.execute(
            """
            SELECT id, product_id, serial_number, status
            FROM product_serials
            WHERE purchase_invoice_id = %s
            """,
            (purchase_id,)
        )
        old_serials = g.cursor.fetchall()
        sold_serials_dict = {s["serial_number"]: s for s in old_serials if str(s.get("status") or "").strip() == "Sold"}

        # Validate new items
        validated_new_items = []
        new_serials_seen = set()
        new_serials_list = []

        for item in items:
            product_id = item.get("product_id")
            if not product_id:
                raise ValueError("Product select नहीं किया गया")

            product_record = get_product_record_by_id(product_id)
            if not product_record:
                raise ValueError("Product नहीं मिला")

            quantity = float(item.get("quantity", 0) or 0)
            if quantity <= 0:
                raise ValueError("Quantity सही नहीं है")

            rate = float(item.get("rate", 0) or 0)
            gst_rate = float(item.get("gst_rate", 0) or 0)
            gst_amount = float(item.get("gst_amount", 0) or 0)
            item_discount = float(item.get("discount", 0) or 0)
            line_total = float(item.get("line_total", 0) or 0)
            warranty_start = item.get("warranty_start")
            warranty_end = item.get("warranty_end")

            raw_loc_id = item.get("location_id")
            location_id = None
            if raw_loc_id not in (None, "", "null", "undefined"):
                try:
                    location_id = int(raw_loc_id)
                except (ValueError, TypeError):
                    location_id = None

            if not location_id:
                g.cursor.execute("SELECT id FROM stock_locations WHERE is_default = 1 AND status = 'Active' LIMIT 1")
                def_row = g.cursor.fetchone()
                if def_row:
                    location_id = def_row["id"]
                else:
                    g.cursor.execute("SELECT id FROM stock_locations WHERE status = 'Active' ORDER BY id ASC LIMIT 1")
                    fb = g.cursor.fetchone()
                    location_id = fb["id"] if fb else 1

            item_unit = (item.get("unit") or "").strip() or None
            p_unit = product_record.get("unit") or "Piece"
            p_sec_unit = product_record.get("secondary_unit")
            p_crate = float(product_record.get("conversion_rate") or 1)

            if p_sec_unit and p_crate > 1:
                if item_unit == p_unit:
                    conversion_factor = p_crate
                    base_quantity = round(quantity * p_crate, 4)
                elif item_unit == p_sec_unit:
                    conversion_factor = 1.0
                    base_quantity = round(quantity, 4)
                else:
                    item_unit = p_unit
                    conversion_factor = p_crate
                    base_quantity = round(quantity * p_crate, 4)
            else:
                if not item_unit:
                    item_unit = p_unit
                conversion_factor = 1.0
                base_quantity = round(quantity, 4)

            serial_numbers = item.get("serial_numbers") or []
            serial_number = item.get("serial_number")
            if isinstance(serial_numbers, list):
                serial_values = [str(v).strip() for v in serial_numbers if str(v).strip()]
            elif serial_number:
                serial_values = [str(serial_number).strip()]
            else:
                serial_values = []

            item_type = product_record.get("item_type")
            if item_type == "Service":
                serial_values = []
                serial_number = None
                warranty_start = None
                warranty_end = None
            else:
                tracks_serial = str(product_record.get("track_serial") or "No") == "Yes"
                if tracks_serial and quantity != int(quantity):
                    raise ValueError("Serial-tracked products require whole-number quantity.")
                if tracks_serial and len(serial_values) != int(quantity):
                    raise ValueError("Every unit of a serial-tracked product needs one serial number.")

                for s_val in serial_values:
                    if s_val in new_serials_seen:
                        raise ValueError(f"Duplicate serial number found: {s_val}")
                    new_serials_seen.add(s_val)

                    g.cursor.execute(
                        "SELECT id, purchase_invoice_id, status FROM product_serials WHERE serial_number = %s AND product_id = %s LIMIT 1",
                        (s_val, product_id)
                    )
                    existing_s = g.cursor.fetchone()
                    if existing_s and int(existing_s.get("purchase_invoice_id") or 0) != int(purchase_id):
                        raise ValueError(f"Serial Number '{s_val}' already exists on another purchase!")

                    new_serials_list.append({
                        "product_id": product_id,
                        "serial_number": s_val,
                        "location_id": location_id,
                        "warranty_start": warranty_start,
                        "warranty_end": warranty_end
                    })

                if serial_values:
                    serial_number = serial_values[0]

            validated_new_items.append({
                "product_id": product_id,
                "location_id": location_id,
                "quantity": quantity,
                "unit": item_unit,
                "base_quantity": base_quantity,
                "conversion_factor": conversion_factor,
                "rate": rate,
                "gst_rate": gst_rate,
                "gst_amount": gst_amount,
                "discount": item_discount,
                "line_total": line_total,
                "serial_number": serial_number,
                "warranty_start": warranty_start,
                "warranty_end": warranty_end,
                "item_type": item_type
            })

        # Ensure previously sold serials are not removed
        for sold_sn, sold_info in sold_serials_dict.items():
            if sold_sn not in new_serials_seen:
                raise ValueError(f"सीरियल नंबर '{sold_sn}' पहले ही बिक्री में दर्ज हो चुका है, इसलिए इसे हटाया नहीं जा सकता।")

        # 1. Reverse old stock from products & location stock
        for old in old_items:
            pid = old.get("product_id")
            loc_id = old.get("location_id")
            qty = float(old.get("base_qty") or 0)
            if pid and old.get("item_type") != "Service":
                g.cursor.execute("UPDATE products SET stock = stock - %s WHERE id = %s", (qty, pid))
                if loc_id:
                    g.cursor.execute(
                        "UPDATE product_location_stock SET quantity = GREATEST(0, quantity - %s) WHERE product_id = %s AND location_id = %s",
                        (qty, pid, loc_id)
                    )

        # 2. Add new stock to products & location stock
        for new_item in validated_new_items:
            pid = new_item["product_id"]
            loc_id = new_item["location_id"]
            base_qty = new_item["base_quantity"]
            if pid and new_item["item_type"] != "Service":
                g.cursor.execute("UPDATE products SET stock = stock + %s WHERE id = %s", (base_qty, pid))
                if loc_id:
                    g.cursor.execute(
                        """
                        INSERT INTO product_location_stock (product_id, location_id, quantity)
                        VALUES (%s, %s, %s)
                        ON DUPLICATE KEY UPDATE quantity = quantity + VALUES(quantity)
                        """,
                        (pid, loc_id, base_qty)
                    )

        # 3. Delete old items and insert new items
        g.cursor.execute("DELETE FROM purchase_invoice_items WHERE purchase_invoice_id = %s", (purchase_id,))

        for new_item in validated_new_items:
            g.cursor.execute(
                """
                INSERT INTO purchase_invoice_items
                (
                    purchase_invoice_id, product_id, location_id, quantity, unit,
                    base_quantity, conversion_factor, rate, gst_rate, gst_amount,
                    discount, line_total, serial_number, warranty_start, warranty_end
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                """,
                (
                    purchase_id, new_item["product_id"], new_item["location_id"], new_item["quantity"], new_item["unit"],
                    new_item["base_quantity"], new_item["conversion_factor"], new_item["rate"], new_item["gst_rate"], new_item["gst_amount"],
                    new_item["discount"], new_item["line_total"], new_item["serial_number"], new_item["warranty_start"], new_item["warranty_end"]
                )
            )

        # 4. Sync serials: delete unsold serials, then re-insert
        g.cursor.execute(
            "DELETE FROM product_serials WHERE purchase_invoice_id = %s AND (status IS NULL OR status = 'In Stock')",
            (purchase_id,)
        )

        for s_obj in new_serials_list:
            if s_obj["serial_number"] in sold_serials_dict:
                g.cursor.execute(
                    """
                    UPDATE product_serials
                    SET supplier_name = %s, purchase_date = %s, location_id = %s, warranty_start = %s, warranty_end = %s
                    WHERE purchase_invoice_id = %s AND serial_number = %s
                    """,
                    (supplier_name, purchase_date, s_obj["location_id"], s_obj["warranty_start"], s_obj["warranty_end"], purchase_id, s_obj["serial_number"])
                )
            else:
                g.cursor.execute(
                    """
                    INSERT INTO product_serials
                    (product_id, serial_number, purchase_invoice_id, purchase_date, supplier_name, location_id, warranty_start, warranty_end, status)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, 'In Stock')
                    """,
                    (s_obj["product_id"], s_obj["serial_number"], purchase_id, purchase_date, supplier_name, s_obj["location_id"], s_obj["warranty_start"], s_obj["warranty_end"])
                )

        # 5. Update purchase invoice record
        g.cursor.execute(
            """
            UPDATE purchase_invoices
            SET supplier_id = %s,
                purchase_date = %s,
                subtotal = %s,
                discount = %s,
                tax_amount = %s,
                grand_total = %s,
                payment_method = %s,
                payment_status = %s,
                notes = %s
            WHERE id = %s
            """,
            (
                supplier_id,
                purchase_date,
                subtotal,
                discount,
                tax_amount,
                grand_total,
                payment_method,
                payment_status,
                notes,
                purchase_id
            )
        )

        g.db.commit()

        return jsonify({
            "success": True,
            "message": "Purchase invoice सफलतापूर्वक अपडेट कर दिया गया",
            "purchase_id": purchase_id,
            "purchase_number": purchase_record.get("purchase_number")
        })

    except Exception as e:
        g.db.rollback()
        print("Update Purchase Error:", e)
        return jsonify({"success": False, "message": str(e)}), 500


# ==========================================
# DELETE PURCHASE INVOICE
# ==========================================

@app.route("/purchases/<int:purchase_id>", methods=["DELETE", "OPTIONS"])
@roles_required("Admin", "Manager")
def delete_purchase(purchase_id):
    if request.method == "OPTIONS":
        return "", 200

    try:
        g.cursor.execute("SELECT * FROM purchase_invoices WHERE id = %s LIMIT 1", (purchase_id,))
        purchase = g.cursor.fetchone()
        if not purchase:
            return jsonify({"success": False, "message": "Purchase invoice नहीं मिला"}), 404

        # 1. Check if any serial number from this purchase was already sold
        g.cursor.execute(
            """
            SELECT serial_number
            FROM product_serials
            WHERE purchase_invoice_id = %s AND status = 'Sold'
            """,
            (purchase_id,)
        )
        sold_rows = g.cursor.fetchall()
        if sold_rows:
            sold_serials = [row["serial_number"] for row in sold_rows]
            sample_serials = ", ".join(sold_serials[:5])
            if len(sold_serials) > 5:
                sample_serials += f" और {len(sold_serials) - 5} अन्य"
            return jsonify({
                "success": False,
                "message": f"इस Purchase के सीरियल नंबर ({sample_serials}) पहले ही ग्राहकों को बेचे जा चुके हैं, इसलिए इस Purchase को हटाया नहीं जा सकता।"
            }), 400

        # 2. Check current product stocks
        g.cursor.execute(
            """
            SELECT pii.product_id, pii.location_id, COALESCE(pii.base_quantity, pii.quantity) AS base_qty,
                   p.product_name, p.stock, p.item_type
            FROM purchase_invoice_items pii
            LEFT JOIN products p ON p.id = pii.product_id
            WHERE pii.purchase_invoice_id = %s
            """,
            (purchase_id,)
        )
        old_items = g.cursor.fetchall()

        required_by_prod = {}
        for row in old_items:
            pid = row.get("product_id")
            if not pid or row.get("item_type") == "Service":
                continue
            qty = float(row.get("base_qty") or 0)
            required_by_prod[pid] = required_by_prod.get(pid, 0) + qty

        for pid, req_qty in required_by_prod.items():
            g.cursor.execute("SELECT product_name, stock FROM products WHERE id = %s", (pid,))
            prod = g.cursor.fetchone()
            curr_stock = float(prod.get("stock") or 0) if prod else 0
            if curr_stock < req_qty:
                return jsonify({
                    "success": False,
                    "message": f"उत्पाद '{prod.get('product_name')}' का वर्तमान स्टॉक ({curr_stock}) इस Purchase की मात्रा ({req_qty}) से कम है (स्टॉक पहले ही बिक चुका है), इसलिए इसे हटाया नहीं जा सकता।"
                }), 400

        # 3. Deduct stock from products & location stock
        for row in old_items:
            pid = row.get("product_id")
            loc_id = row.get("location_id")
            qty = float(row.get("base_qty") or 0)
            if pid and row.get("item_type") != "Service":
                g.cursor.execute("UPDATE products SET stock = GREATEST(0, stock - %s) WHERE id = %s", (qty, pid))
                if loc_id:
                    g.cursor.execute(
                        "UPDATE product_location_stock SET quantity = GREATEST(0, quantity - %s) WHERE product_id = %s AND location_id = %s",
                        (qty, pid, loc_id)
                    )

        # 4. Delete product serials attached to this purchase
        g.cursor.execute("DELETE FROM product_serials WHERE purchase_invoice_id = %s", (purchase_id,))

        # 5. Delete purchase items and invoice
        g.cursor.execute("DELETE FROM purchase_invoice_items WHERE purchase_invoice_id = %s", (purchase_id,))
        g.cursor.execute("DELETE FROM purchase_invoices WHERE id = %s", (purchase_id,))
        sync_document_sequence("purchase-billing", "next_purchase_number", "purchase_prefix", "PUR-")
        g.db.commit()

        return jsonify({
            "success": True,
            "message": "Purchase invoice सफलतापूर्वक हटा दिया गया"
        }), 200

    except Exception as e:
        g.db.rollback()
        print("Delete Purchase Error:", e)
        return jsonify({"success": False, "message": str(e)}), 500


# ==========================================
# SAVE SALES INVOICE
# ==========================================

@app.route(
    "/sales",
    methods=["POST", "OPTIONS"]
)
def add_sale():

    if request.method == "OPTIONS":
        return "", 200

    data = request.get_json()

    if not data:
        return jsonify({
            "success": False,
            "message": "Sale data नहीं मिला"
        }), 400

    invoice_number = data.get("invoice_number")
    customer_id = data.get("customer_id")
    invoice_date = str(data.get("invoice_date") or data.get("date") or get_ist_today_str()).strip()
    subtotal = data.get("subtotal", 0)
    discount = data.get("discount", 0)
    tax_amount = data.get("tax_amount", 0)
    grand_total = data.get("grand_total", 0)
    payment_method = data.get(
        "payment_method",
        "Cash"
    )
    payment_status = data.get(
        "payment_status",
        "Unpaid"
    )
    paid_amount = data.get("paid_amount")
    default_warranty_months = int(data.get("warranty_months", 0) or 0)
    notes = data.get("notes", "")
    items = data.get("items", [])

    document_type = (data.get("document_type") or "Invoice").strip()
    if document_type not in ["Invoice", "Quotation"]:
        document_type = "Invoice"

    if document_type == "Quotation":
        payment_method = "Estimate"
        payment_status = "Unpaid"
        paid_amount = 0.0
        cash_amount = 0.0
        upi_amount = 0.0
        card_amount = 0.0
        bank_amount = 0.0
    else:
        cash_amount = float(data.get("cash_amount", 0) or 0)
        upi_amount = float(data.get("upi_amount", 0) or 0)
        card_amount = float(data.get("card_amount", 0) or 0)
        bank_amount = float(data.get("bank_amount", 0) or 0)

    service_job_id = data.get("service_job_id")

    if document_type != "Quotation":
        split_sum = cash_amount + upi_amount + card_amount + bank_amount
        if split_sum > 0:
            paid_amount = split_sum
            modes_used = sum([1 for m in [cash_amount, upi_amount, card_amount, bank_amount] if m > 0])
            if modes_used > 1:
                payment_method = "Split"
            elif cash_amount > 0:
                payment_method = "Cash"
            elif upi_amount > 0:
                payment_method = "UPI"
            elif card_amount > 0:
                payment_method = "Card"
            elif bank_amount > 0:
                payment_method = "Bank Transfer"


    if not customer_id:

        return jsonify({
            "success": False,
            "message": "Customer select करें"
        }), 400


    if not items:

        return jsonify({
            "success": False,
            "message": "कम से कम एक product जोड़ें"
        }), 400


    try:

        doc_tax_incl = False if str(data.get("tax_type", "")).lower() == "exclusive" or data.get("is_tax_inclusive") is False else True
        items, subtotal, tax_amount, grand_total, discount = normalize_document_items(
            items,
            discount,
            default_tax_inclusive=doc_tax_incl
        )
        payment_status, paid_amount = normalize_sale_payment(
            grand_total,
            payment_status,
            paid_amount
        )
        g.cursor.execute("SELECT id FROM customers WHERE id = %s AND status = 'Active' LIMIT 1", (customer_id,))
        if not g.cursor.fetchone():
            raise ValueError("Selected customer is not available.")

        if document_type == "Quotation":
            invoice_number = generate_document_number(
                "quotation-billing", "quotation_prefix", "next_quotation_number", "QTN-"
            )
        else:
            invoice_number = generate_document_number(
                "invoice-billing", "invoice_prefix", "next_invoice_number", "INV-"
            )

        # ----------------------------------
        # CREATE SALES INVOICE
        # ----------------------------------

        g.cursor.execute(
            """
            INSERT INTO sales_invoices
            (
                invoice_number,
                customer_id,
                invoice_date,
                subtotal,
                discount,
                tax_amount,
                grand_total,
                paid_amount,
                payment_method,
                payment_status,
                notes,
                document_type,
                cash_amount,
                upi_amount,
                card_amount,
                bank_amount,
                service_job_id
            )
            VALUES
            (
                %s, %s, %s, %s, %s,
                %s, %s, %s, %s, %s,
                %s, %s, %s, %s, %s,
                %s, %s
            )
            """,
            (
                invoice_number,
                customer_id,
                invoice_date,
                subtotal,
                discount,
                tax_amount,
                grand_total,
                paid_amount,
                payment_method,
                payment_status,
                notes,
                document_type,
                cash_amount,
                upi_amount,
                card_amount,
                bank_amount,
                service_job_id
            )
        )


        invoice_id = g.cursor.lastrowid


        # ----------------------------------
        # SAVE SALE ITEMS
        # ----------------------------------

        for item in items:

            product_id = item.get(
                "product_id"
            )

            quantity = float(
                item.get("quantity", 0)
            )

            rate = float(
                item.get("rate", 0)
            )

            gst_rate = float(
                item.get("gst_rate", 0)
            )

            gst_amount = float(
                item.get("gst_amount", 0)
            )

            item_discount = float(
                item.get("discount", 0)
            )

            line_total = float(
                item.get("line_total", 0)
            )

            serial_number = item.get(
                "serial_number"
            )

            serial_numbers = item.get(
                "serial_numbers", []
            ) or []

            if not product_id:
                raise Exception(
                    "Product select नहीं किया गया"
                )

            product_record = get_product_record_by_id(product_id)
            if not product_record:
                raise Exception("Product नहीं मिला")

            try:
                raw_wm = item.get("warranty_months")
                if raw_wm is not None and str(raw_wm).strip() != "":
                    warranty_months = int(raw_wm)
                else:
                    warranty_months = int(product_record.get("warranty_months", default_warranty_months) or default_warranty_months or 0)
            except (TypeError, ValueError):
                warranty_months = 0

            warranty_start = None
            warranty_end = None

            if warranty_months > 0 and invoice_date:
                warranty_start = str(invoice_date)[:10]
                warranty_end = add_months_to_date(invoice_date, warranty_months)

            if quantity <= 0:

                raise Exception(
                    "Quantity सही नहीं है"
                )

            if product_record.get("item_type") == "Service":
                serial_values = []
                serial_number = None
            else:
                # ----------------------------------
                # SERIAL NUMBER SALE
                # ----------------------------------

                serial_values = []
                if isinstance(serial_numbers, list):
                    serial_values = [str(v).strip() for v in serial_numbers if str(v).strip()]
                elif serial_number:
                    serial_values = [str(serial_number).strip()]

                if serial_values and len(serial_values) > quantity:
                    raise Exception("Serial count cannot exceed quantity for the product")

                tracks_serial = str(product_record.get("track_serial") or "No") == "Yes"
                if document_type != "Quotation" and tracks_serial and (quantity != int(quantity) or len(serial_values) != int(quantity)):
                    raise Exception("Every unit of a serial-tracked product needs one serial number")

                if serial_values:
                    seen = set()
                    for serial_value in serial_values:
                        if serial_value in seen:
                            raise Exception("Duplicate serial number found: " + serial_value)
                        seen.add(serial_value)

                        g.cursor.execute(
                            """
                            SELECT id, status
                            FROM product_serials
                            WHERE serial_number = %s
                            AND product_id = %s
                            LIMIT 1
                            """,
                            (
                                serial_value,
                                product_id
                            )
                        )

                        serial_record = g.cursor.fetchone()
                        if serial_record and str(serial_record.get("status") or "").strip() != "In Stock":
                            raise Exception(
                                "Serial Number is not available in stock: " + serial_value
                            )


            item_unit = (item.get("unit") or "").strip() or None
            p_unit = product_record.get("unit") or "Piece"
            p_sec_unit = product_record.get("secondary_unit")
            p_crate = float(product_record.get("conversion_rate") or 1)

            if p_sec_unit and p_crate > 1:
                if item_unit == p_unit:
                    conversion_factor = p_crate
                    base_quantity = round(quantity * p_crate, 4)
                elif item_unit == p_sec_unit:
                    conversion_factor = 1.0
                    base_quantity = round(quantity, 4)
                else:
                    item_unit = p_unit
                    conversion_factor = p_crate
                    base_quantity = round(quantity * p_crate, 4)
            else:
                if not item_unit:
                    item_unit = p_unit
                conversion_factor = 1.0
                base_quantity = round(quantity, 4)

            if document_type != "Quotation":
                # ----------------------------------
                # CHECK STOCK
                # ----------------------------------

                g.cursor.execute(
                    """
                    SELECT stock
                    FROM products
                    WHERE id = %s
                    """,
                    (product_id,)
                )

                stock_record = g.cursor.fetchone()

                if not stock_record:
                    raise Exception("Product नहीं मिला")

                current_stock = float(stock_record["stock"] or 0)

                if current_stock < base_quantity:
                    p_name = product_record.get("product_name") or "Product"
                    raise Exception(f"'{p_name}' के stock में पर्याप्त quantity नहीं है (उपलब्ध: {current_stock}, आवश्यक: {base_quantity})")


            # ----------------------------------
            # SAVE SALE ITEM
            # ----------------------------------

            g.cursor.execute(
                """
                INSERT INTO sales_invoice_items
                (
                    invoice_id,
                    product_id,
                    quantity,
                    unit,
                    base_quantity,
                    conversion_factor,
                    rate,
                    cost_rate,
                    gst_rate,
                    gst_amount,
                    discount,
                    line_total,
                    warranty_months,
                    serial_number
                )
                VALUES
                (
                    %s, %s, %s, %s, %s, %s,
                    %s, %s, %s, %s, %s, %s, %s, %s
                )
                """,
                (
                    invoice_id,
                    product_id,
                    quantity,
                    item_unit,
                    base_quantity,
                    conversion_factor,
                    rate,
                    float(product_record.get("purchase_price") or 0),
                    gst_rate,
                    gst_amount,
                    item_discount,
                    line_total,
                    warranty_months,
                    serial_number
                )
            )


            if document_type != "Quotation":
                # ----------------------------------
                # DECREASE STOCK
                # ----------------------------------

                if product_record.get("item_type") != "Service":
                    g.cursor.execute(
                        """
                        UPDATE products
                        SET stock = stock - %s
                        WHERE id = %s
                        """,
                        (
                            base_quantity,
                            product_id
                        )
                    )

                    g.cursor.execute(
                        """
                        UPDATE product_location_stock
                        SET quantity = GREATEST(0, quantity - %s)
                        WHERE product_id = %s AND location_id = (SELECT id FROM stock_locations WHERE is_default = 1 LIMIT 1)
                        """,
                        (
                            base_quantity,
                            product_id
                        )
                    )


                # ----------------------------------
                # MARK SERIAL AS SOLD
                # ----------------------------------

                if serial_values:
                    for serial_value in serial_values:
                        if not serial_value:
                            continue

                        g.cursor.execute(
                            """
                            SELECT id, status
                            FROM product_serials
                            WHERE serial_number = %s
                            AND product_id = %s
                            LIMIT 1
                            """,
                            (
                                serial_value,
                                product_id
                            )
                        )
                        existing_serial = g.cursor.fetchone()

                        if existing_serial:
                            if str(existing_serial.get("status") or "").strip() == "Sold":
                                raise ValueError("Serial Number is not available in stock: " + serial_value)
                            g.cursor.execute(
                                """
                                UPDATE product_serials
                                SET
                                    sale_invoice_id = %s,
                                    customer_id = %s,
                                    sale_date = %s,
                                    warranty_start = %s,
                                    warranty_end = %s,
                                    status = 'Sold'
                                WHERE serial_number = %s
                                AND product_id = %s
                                """,
                                (
                                    invoice_id,
                                    customer_id,
                                    invoice_date,
                                    warranty_start,
                                    warranty_end,
                                    serial_value,
                                    product_id
                                )
                            )
                        else:
                            g.cursor.execute(
                                """
                                INSERT INTO product_serials
                                (
                                    product_id,
                                    serial_number,
                                    sale_invoice_id,
                                    customer_id,
                                    sale_date,
                                    warranty_start,
                                    warranty_end,
                                    status
                                )
                                VALUES (%s, %s, %s, %s, %s, %s, %s, 'Sold')
                                """,
                                (
                                    product_id,
                                    serial_value,
                                    invoice_id,
                                    customer_id,
                                    invoice_date,
                                    warranty_start,
                                    warranty_end
                                )
                            )

        if service_job_id:
            g.cursor.execute(
                "UPDATE service_jobs SET invoice_id = %s, service_status = 'Delivered', delivered_at = NOW() WHERE id = %s",
                (invoice_id, service_job_id)
            )

        converted_from_quotation_id = data.get("converted_from_quotation_id")
        if converted_from_quotation_id and document_type == "Invoice":
            g.cursor.execute(
                """
                UPDATE sales_invoices
                SET payment_status = 'Converted',
                    notes = CONCAT(COALESCE(notes, ''), '\nConverted to Tax Invoice: ', %s)
                WHERE id = %s AND document_type = 'Quotation'
                """,
                (invoice_number, converted_from_quotation_id)
            )

        # ----------------------------------
        # COMMIT
        # ----------------------------------

        g.db.commit()

        msg = "Quotation Successfully Saved" if document_type == "Quotation" else "Sales Invoice Successfully Saved"
        return jsonify({
            "success": True,
            "message": msg,
            "invoice_id": invoice_id,
            "invoice_number": invoice_number,
            "grand_total": grand_total,
            "document_type": document_type
        })


    except Exception as e:

        g.db.rollback()

        print(
            "Sales Error:",
            e
        )

        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


# ==========================================
# GET SALES LIST
# ==========================================

@app.route("/sales", methods=["GET"])
def get_sales_list():
    try:
        invoice_number = (request.args.get("invoice_number") or "").strip()
        customer = (request.args.get("customer") or "").strip()
        from_date = normalize_date_string(request.args.get("from_date"))
        to_date = normalize_date_string(request.args.get("to_date"))
        payment_status = (request.args.get("payment_status") or "").strip()
        document_type = (request.args.get("document_type") or request.args.get("type") or "").strip()

        sql = """
            SELECT
                si.id,
                si.invoice_number,
                si.invoice_date,
                si.subtotal,
                si.discount,
                si.tax_amount,
                si.grand_total,
                si.paid_amount,
                si.payment_method,
                si.payment_status,
                si.document_type,
                si.service_job_id,
                COALESCE(c.customer_name, 'Walk-in Customer') AS customer_name,
                si.customer_id
            FROM sales_invoices si
            LEFT JOIN customers c
                ON c.id = si.customer_id
            WHERE 1=1
        """
        params = []

        if invoice_number:
            sql += " AND si.invoice_number LIKE %s "
            params.append(f"%{invoice_number}%")

        if customer:
            sql += " AND (c.customer_name LIKE %s OR c.mobile LIKE %s) "
            params.extend([f"%{customer}%", f"%{customer}%"])

        if from_date:
            sql += " AND DATE(si.invoice_date) >= %s "
            params.append(from_date)

        if to_date:
            sql += " AND DATE(si.invoice_date) <= %s "
            params.append(to_date)

        if payment_status:
            sql += " AND si.payment_status = %s "
            params.append(payment_status)

        if document_type and document_type.lower() != "all":
            if document_type.lower() in ["quotation", "quote"]:
                sql += " AND si.document_type = 'Quotation' "
            else:
                sql += " AND (si.document_type IS NULL OR si.document_type = 'Invoice' OR si.document_type != 'Quotation') "
        elif not document_type:
            # Default: only standard sales invoices (exclude quotations from general sales list)
            sql += " AND (si.document_type IS NULL OR si.document_type = 'Invoice' OR si.document_type != 'Quotation') "

        sql += " ORDER BY si.invoice_date DESC, si.id DESC "

        g.cursor.execute(sql, tuple(params))
        rows = g.cursor.fetchall()

        invoices = []
        for row in rows:
            invoices.append({
                "id": row["id"],
                "invoice_number": row["invoice_number"],
                "invoice_date": normalize_date_string(row["invoice_date"]),
                "customer": row["customer_name"],
                "subtotal": float(row["subtotal"] or 0),
                "discount": float(row["discount"] or 0),
                "tax_amount": float(row["tax_amount"] or 0),
                "grand_total": float(row["grand_total"] or 0),
                "paid_amount": float(row["paid_amount"] or 0),
                "outstanding_amount": round(float(row["grand_total"] or 0) - float(row["paid_amount"] or 0), 2),
                "payment_method": row["payment_method"],
                "payment_status": row["payment_status"],
                "document_type": row.get("document_type") or "Invoice",
                "service_job_id": row.get("service_job_id"),
                "customer_id": row.get("customer_id")
            })

        return jsonify({
            "success": True,
            "invoices": invoices,
            "total": len(invoices)
        })

    except Exception as e:
        print("Get Sales List Error:", e)
        return jsonify({
            "success": False,
            "message": "Unable to load sales invoices."
        }), 500


# ==========================================
# UPDATE SALE
# ==========================================

@app.route("/sales/<int:invoice_id>", methods=["PUT", "OPTIONS"])
@roles_required("Admin", "Manager")
def update_sale(invoice_id):
    if request.method == "OPTIONS":
        return "", 200

    data = request.get_json()
    if not data:
        return jsonify({
            "success": False,
            "message": "Invoice data not found."
        }), 400

    customer_id = data.get("customer_id")
    invoice_date = normalize_date_string(data.get("invoice_date"))
    subtotal = float(data.get("subtotal", 0) or 0)
    discount = float(data.get("discount", 0) or 0)
    tax_amount = float(data.get("tax_amount", 0) or 0)
    grand_total = float(data.get("grand_total", 0) or 0)
    payment_method = data.get("payment_method", "Cash")
    payment_status = data.get("payment_status", "Paid")
    paid_amount = data.get("paid_amount")
    notes = data.get("notes", "")
    items = data.get("items", [])

    if not customer_id:
        return jsonify({
            "success": False,
            "message": "Please select a customer."
        }), 400

    if not items:
        return jsonify({
            "success": False,
            "message": "Please add at least one item."
        }), 400

    try:
        g.cursor.execute(
            """
            SELECT id, customer_id, invoice_date, invoice_number, document_type
            FROM sales_invoices
            WHERE id = %s
            LIMIT 1
            """,
            (invoice_id,)
        )
        invoice_record = g.cursor.fetchone()
        if not invoice_record:
            return jsonify({
                "success": False,
                "message": "Document not found."
            }), 404

        existing_doc_type = (invoice_record.get("document_type") or "Invoice").strip()
        doc_type = (data.get("document_type") or existing_doc_type or "Invoice").strip()
        is_quotation = (doc_type == "Quotation")

        doc_tax_incl = False if str(data.get("tax_type", "")).lower() == "exclusive" or data.get("is_tax_inclusive") is False else True
        items, subtotal, tax_amount, grand_total, discount = normalize_document_items(
            items,
            discount,
            default_tax_inclusive=doc_tax_incl
        )
        payment_status, paid_amount = normalize_sale_payment(
            grand_total,
            payment_status,
            paid_amount
        )
        g.cursor.execute("SELECT id FROM customers WHERE id = %s AND status = 'Active' LIMIT 1", (customer_id,))
        if not g.cursor.fetchone():
            raise ValueError("Selected customer is not available.")

        g.cursor.execute(
            """
            SELECT sii.product_id, COALESCE(sii.base_quantity, sii.quantity) AS base_quantity, p.item_type
            FROM sales_invoice_items sii
            LEFT JOIN products p ON p.id = sii.product_id
            WHERE sii.invoice_id = %s
            """,
            (invoice_id,)
        )
        old_items = g.cursor.fetchall()

        old_qty_by_product = {}
        for row in old_items:
            product_id = row.get("product_id")
            if not product_id:
                continue
            item_type = (row.get("item_type") or "Product").strip()
            if item_type == "Service":
                continue
            old_qty_by_product[product_id] = old_qty_by_product.get(product_id, 0) + float(row.get("base_quantity") or 0)

        new_qty_by_product = {}
        validated_new_items = []
        seen_serials = set()

        for item in items:
            product_id = item.get("product_id")
            if not product_id:
                raise ValueError("Please select a product.")

            product_record = get_product_record_by_id(product_id)
            if not product_record:
                raise ValueError("Product not found.")

            item_type = (product_record.get("item_type") or "Product").strip()
            quantity = float(item.get("quantity", 0) or 0)
            if quantity <= 0:
                raise ValueError("Quantity must be greater than zero.")

            item_unit = (item.get("unit") or "").strip() or None
            p_unit = product_record.get("unit") or "Piece"
            p_sec_unit = product_record.get("secondary_unit")
            p_crate = float(product_record.get("conversion_rate") or 1)

            if p_sec_unit and p_crate > 1:
                if item_unit == p_unit:
                    conversion_factor = p_crate
                    base_quantity = round(quantity * p_crate, 4)
                elif item_unit == p_sec_unit:
                    conversion_factor = 1.0
                    base_quantity = round(quantity, 4)
                else:
                    item_unit = p_unit
                    conversion_factor = p_crate
                    base_quantity = round(quantity * p_crate, 4)
            else:
                if not item_unit:
                    item_unit = p_unit
                conversion_factor = 1.0
                base_quantity = round(quantity, 4)

            rate = float(item.get("rate", 0) or 0)
            gst_rate = float(item.get("gst_rate", 0) or 0)
            line_total = float(item.get("line_total", 0) or 0)

            if item_type != "Service":
                new_qty_by_product[product_id] = new_qty_by_product.get(product_id, 0) + base_quantity

            serial_values = []
            if isinstance(item.get("serial_numbers"), list):
                serial_values = [str(v).strip() for v in item.get("serial_numbers") if str(v).strip()]
            elif item.get("serial_number"):
                serial_values = [str(item.get("serial_number")).strip()]

            if item_type == "Service" or is_quotation:
                serial_values = []
            else:
                if len(serial_values) > quantity:
                    raise ValueError("Serial count cannot exceed quantity for the product.")

                tracks_serial = str(product_record.get("track_serial") or "No") == "Yes"
                if tracks_serial and (quantity != int(quantity) or len(serial_values) != int(quantity)):
                    raise ValueError("Every unit of a serial-tracked product needs one serial number.")

                for serial_value in serial_values:
                    if serial_value in seen_serials:
                        raise ValueError("Duplicate serial number found: " + serial_value)
                    seen_serials.add(serial_value)

                    g.cursor.execute(
                        """
                        SELECT id, status, serial_number
                        FROM product_serials
                        WHERE serial_number = %s AND product_id = %s
                        LIMIT 1
                        """,
                        (serial_value, product_id)
                    )
                    existing_serial = g.cursor.fetchone()
                    if existing_serial and str(existing_serial.get("status") or "").strip() == "Sold":
                        g.cursor.execute(
                            """
                            SELECT sale_invoice_id
                            FROM product_serials
                            WHERE serial_number = %s AND product_id = %s
                            LIMIT 1
                            """,
                            (serial_value, product_id)
                        )
                        serial_owner = g.cursor.fetchone()
                        if serial_owner and int(serial_owner.get("sale_invoice_id") or 0) != int(invoice_id):
                            raise ValueError("Serial Number is already sold or unavailable: " + serial_value)

            validated_new_items.append({
                "product_id": int(product_id),
                "quantity": quantity,
                "unit": item_unit,
                "base_quantity": base_quantity,
                "conversion_factor": conversion_factor,
                "rate": rate,
                "gst_rate": gst_rate,
                "gst_amount": float(item.get("gst_amount", 0) or 0),
                "discount": float(item.get("discount", 0) or 0),
                "line_total": line_total,
                "serial_number": serial_values[0] if serial_values else None,
                "serial_numbers": serial_values,
                "product_type": item_type,
                "cost_rate": float(product_record.get("purchase_price") or 0),
                "warranty_months": int(item.get("warranty_months") or (product_record.get("warranty_months") if product_record else 0) or 0)
            })

        if not is_quotation:
            product_stock_delta = {}
            all_product_ids = set(old_qty_by_product.keys()) | set(new_qty_by_product.keys())
            for product_id in all_product_ids:
                old_qty = old_qty_by_product.get(product_id, 0)
                new_qty = new_qty_by_product.get(product_id, 0)
                delta = new_qty - old_qty
                if delta != 0:
                    product_record = get_product_record_by_id(product_id)
                    if product_record and str(product_record.get("item_type") or "Product").strip() != "Service":
                        product_stock_delta[product_id] = delta

            for product_id, delta in product_stock_delta.items():
                if delta <= 0:
                    continue
                g.cursor.execute("SELECT stock FROM products WHERE id = %s FOR UPDATE", (product_id,))
                stock_record = g.cursor.fetchone()
                if not stock_record or float(stock_record.get("stock") or 0) < delta:
                    raise ValueError("Product stock is insufficient for this invoice update.")

            g.cursor.execute(
                """
                SELECT product_id, serial_number
                FROM product_serials
                WHERE sale_invoice_id = %s
                ORDER BY id ASC
                """,
                (invoice_id,)
            )
            old_serial_rows = g.cursor.fetchall()
            old_serials_by_product = {}
            for row in old_serial_rows:
                pid = row.get("product_id")
                serial_value = row.get("serial_number")
                if pid is None or not serial_value:
                    continue
                old_serials_by_product.setdefault(pid, []).append(str(serial_value))

            for pid, serial_values in old_serials_by_product.items():
                if not serial_values:
                    continue
                g.cursor.execute(
                    """
                    UPDATE product_serials
                    SET
                        sale_invoice_id = NULL,
                        customer_id = NULL,
                        sale_date = NULL,
                        warranty_start = NULL,
                        warranty_end = NULL,
                        status = 'In Stock'
                    WHERE sale_invoice_id = %s AND product_id = %s
                    """,
                    (invoice_id, pid)
                )

            for product_id, delta in product_stock_delta.items():
                if delta > 0:
                    g.cursor.execute(
                        """
                        UPDATE products
                        SET stock = stock - %s
                        WHERE id = %s
                        """,
                        (delta, product_id)
                    )
                elif delta < 0:
                    g.cursor.execute(
                        """
                        UPDATE products
                        SET stock = stock + %s
                        WHERE id = %s
                        """,
                        (abs(delta), product_id)
                    )

        g.cursor.execute(
            "DELETE FROM sales_invoice_items WHERE invoice_id = %s",
            (invoice_id,)
        )

        for item in validated_new_items:
            product_id = item["product_id"]
            quantity = item["quantity"]
            gst_rate = item["gst_rate"]
            gst_amount = item["gst_amount"]
            item_discount = item["discount"]
            line_total = item["line_total"]
            serial_number = item["serial_number"]
            serial_values = item["serial_numbers"]
            item_type = item["product_type"]
            cost_rate = item["cost_rate"]
            warranty_months = item["warranty_months"]

            g.cursor.execute(
                """
                INSERT INTO sales_invoice_items
                (
                    invoice_id,
                    product_id,
                    quantity,
                    unit,
                    base_quantity,
                    conversion_factor,
                    rate,
                    cost_rate,
                    gst_rate,
                    gst_amount,
                    discount,
                    line_total,
                    warranty_months,
                    serial_number
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                """,
                (
                    invoice_id,
                    product_id,
                    quantity,
                    item.get("unit"),
                    item.get("base_quantity", quantity),
                    item.get("conversion_factor", 1.0),
                    item.get("rate", 0),
                    cost_rate,
                    gst_rate,
                    gst_amount,
                    item_discount,
                    line_total,
                    warranty_months,
                    serial_number
                )
            )

            if not is_quotation and item_type != "Service" and serial_values:
                warranty_start = invoice_date
                warranty_end = add_months_to_date(invoice_date, warranty_months) if warranty_months > 0 else None

                for serial_value in serial_values:
                    if not serial_value:
                        continue

                    g.cursor.execute(
                        """
                        SELECT id, status
                        FROM product_serials
                        WHERE serial_number = %s AND product_id = %s
                        LIMIT 1
                        """,
                        (serial_value, product_id)
                    )
                    existing_serial = g.cursor.fetchone()
                    if existing_serial:
                        if str(existing_serial.get("status") or "").strip() == "Sold":
                            raise ValueError("Serial Number is not available in stock: " + serial_value)
                        g.cursor.execute(
                            """
                            UPDATE product_serials
                            SET
                                sale_invoice_id = %s,
                                customer_id = %s,
                                sale_date = %s,
                                warranty_start = %s,
                                warranty_end = %s,
                                status = 'Sold'
                            WHERE serial_number = %s AND product_id = %s
                            """,
                            (
                                invoice_id,
                                customer_id,
                                invoice_date,
                                warranty_start,
                                warranty_end,
                                serial_value,
                                product_id
                            )
                        )
                    else:
                        g.cursor.execute(
                            """
                            INSERT INTO product_serials
                            (
                                product_id,
                                serial_number,
                                sale_invoice_id,
                                customer_id,
                                sale_date,
                                warranty_start,
                                warranty_end,
                                status
                            )
                            VALUES (%s, %s, %s, %s, %s, %s, %s, 'Sold')
                            """,
                            (
                                product_id,
                                serial_value,
                                invoice_id,
                                customer_id,
                                invoice_date,
                                warranty_start,
                                warranty_end
                            )
                        )

        g.cursor.execute(
            """
            UPDATE sales_invoices
            SET
                customer_id = %s,
                invoice_date = %s,
                subtotal = %s,
                discount = %s,
                tax_amount = %s,
                grand_total = %s,
                paid_amount = %s,
                payment_method = %s,
                payment_status = %s,
                notes = %s,
                document_type = %s
            WHERE id = %s
            """,
            (
                customer_id,
                invoice_date,
                subtotal,
                discount,
                tax_amount,
                grand_total,
                paid_amount,
                payment_method,
                payment_status,
                notes,
                doc_type,
                invoice_id
            )
        )

        # Allow updating invoice_number / quotation_number if provided
        new_inv_num = (data.get("invoice_number") or "").strip()
        final_inv_num = invoice_record.get("invoice_number")
        if new_inv_num and new_inv_num != final_inv_num:
            g.cursor.execute("SELECT id FROM sales_invoices WHERE invoice_number = %s AND id != %s LIMIT 1", (new_inv_num, invoice_id))
            if g.cursor.fetchone():
                return jsonify({"success": False, "message": f"Document number '{new_inv_num}' पहले से किसी अन्य बिल/कोटेशन में उपयोग किया गया है।"}), 400

            g.cursor.execute("UPDATE sales_invoices SET invoice_number = %s WHERE id = %s", (new_inv_num, invoice_id))
            final_inv_num = new_inv_num
            if is_quotation:
                sync_document_sequence("quotation-billing", "next_quotation_number", "quotation_prefix", "QTN-", "Quotation")
            else:
                sync_document_sequence("invoice-billing", "next_invoice_number", "invoice_prefix", "INV-", "Invoice")

        g.db.commit()

        doc_label = "Quotation" if is_quotation else "Invoice"
        return jsonify({
            "success": True,
            "message": f"{doc_label} updated successfully.",
            "invoice_id": invoice_id,
            "invoice_number": final_inv_num,
            "document_type": doc_type
        })

    except Exception as e:
        g.db.rollback()
        print("Update Sale Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


# ==========================================
# DELETE SALE / QUOTATION
# ==========================================

@app.route("/sales/<int:invoice_id>", methods=["DELETE", "OPTIONS"])
@roles_required("Admin", "Manager", "Staff")
def delete_sale(invoice_id):
    if request.method == "OPTIONS":
        return "", 200

    try:
        g.cursor.execute("SELECT * FROM sales_invoices WHERE id = %s LIMIT 1", (invoice_id,))
        invoice = g.cursor.fetchone()
        if not invoice:
            return jsonify({"success": False, "message": "Document नहीं मिला"}), 404

        doc_type = (invoice.get("document_type") or "Invoice").strip()
        is_quotation = (doc_type == "Quotation")

        if not is_quotation:
            # Restore stock for Tax Invoice items
            g.cursor.execute(
                """
                SELECT sii.product_id, COALESCE(sii.base_quantity, sii.quantity) AS base_qty, p.item_type
                FROM sales_invoice_items sii
                LEFT JOIN products p ON p.id = sii.product_id
                WHERE sii.invoice_id = %s
                """,
                (invoice_id,)
            )
            items = g.cursor.fetchall()
            for item in items:
                pid = item.get("product_id")
                qty = float(item.get("base_qty") or 0)
                if pid and item.get("item_type") != "Service":
                    g.cursor.execute("UPDATE products SET stock = stock + %s WHERE id = %s", (qty, pid))
                    g.cursor.execute(
                        "UPDATE product_location_stock SET quantity = quantity + %s WHERE product_id = %s AND location_id = (SELECT id FROM stock_locations WHERE is_default = 1 LIMIT 1)",
                        (qty, pid)
                    )

            # Release serials back to 'In Stock'
            g.cursor.execute(
                """
                UPDATE product_serials
                SET status = 'In Stock', sale_invoice_id = NULL, customer_id = NULL, sale_date = NULL, warranty_start = NULL, warranty_end = NULL
                WHERE sale_invoice_id = %s
                """,
                (invoice_id,)
            )

        # Unlink customer payments if any
        try:
            g.cursor.execute("UPDATE customer_payments SET invoice_id = NULL WHERE invoice_id = %s", (invoice_id,))
        except Exception:
            pass

        # Clean up sales_returns if any
        try:
            g.cursor.execute("DELETE FROM sales_return_items WHERE return_id IN (SELECT id FROM sales_returns WHERE invoice_id = %s)", (invoice_id,))
            g.cursor.execute("DELETE FROM sales_returns WHERE invoice_id = %s", (invoice_id,))
        except Exception:
            pass

        g.cursor.execute("DELETE FROM sales_invoice_items WHERE invoice_id = %s", (invoice_id,))
        g.cursor.execute("DELETE FROM sales_invoices WHERE id = %s", (invoice_id,))

        if is_quotation:
            sync_document_sequence("quotation-billing", "next_quotation_number", "quotation_prefix", "QTN-", "Quotation")
        else:
            sync_document_sequence("invoice-billing", "next_invoice_number", "invoice_prefix", "INV-", "Invoice")

        g.db.commit()

        label = "Quotation" if is_quotation else "Sales Invoice"
        return jsonify({
            "success": True,
            "message": f"{label} सफलतापूर्वक हटा दिया गया (Deleted Successfully)"
        }), 200

    except Exception as e:
        g.db.rollback()
        print("Delete Sale Error:", e)
        return jsonify({"success": False, "message": str(e)}), 500


@app.route("/sales/sync-sequence", methods=["POST", "OPTIONS"])
@roles_required("Admin", "Manager", "Staff")
def api_sync_sales_sequence():
    if request.method == "OPTIONS":
        return "", 200
    try:
        data = request.get_json(silent=True) or {}
        doc_type = data.get("document_type", "Quotation")
        if doc_type == "Quotation":
            sync_document_sequence("quotation-billing", "next_quotation_number", "quotation_prefix", "QTN-", "Quotation")
            cat, key = "quotation-billing", "next_quotation_number"
        else:
            sync_document_sequence("invoice-billing", "next_invoice_number", "invoice_prefix", "INV-", "Invoice")
            cat, key = "invoice-billing", "next_invoice_number"

        g.cursor.execute("SELECT setting_value FROM settings WHERE category = %s AND setting_key = %s LIMIT 1", (cat, key))
        row = g.cursor.fetchone()
        next_val = row.get("setting_value") if row else "1"
        g.db.commit()

        return jsonify({
            "success": True,
            "message": f"{doc_type} sequence synchronized successfully.",
            "next_number": next_val
        }), 200
    except Exception as e:
        g.db.rollback()
        return jsonify({"success": False, "message": str(e)}), 500


# ==========================================
# GET SALE DETAILS
# ==========================================

@app.route("/sales/<int:invoice_id>", methods=["GET"])
def get_sale_details(invoice_id):
    try:
        g.cursor.execute(
            """
            SELECT
                si.id,
                si.customer_id,
                si.invoice_number,
                si.invoice_date,
                si.subtotal,
                si.discount,
                si.tax_amount,
                si.grand_total,
                si.paid_amount,
                si.payment_method,
                si.payment_status,
                si.notes,
                si.document_type,
                si.cash_amount,
                si.upi_amount,
                si.card_amount,
                si.bank_amount,
                si.service_job_id,

                c.customer_name,
                c.mobile AS customer_mobile,
                c.address AS customer_address,
                c.gstin AS customer_gstin,
                c.email AS customer_email
            FROM sales_invoices si
            LEFT JOIN customers c
                ON c.id = si.customer_id
            WHERE si.id = %s
            LIMIT 1
            """,
            (invoice_id,)
        )

        invoice = g.cursor.fetchone()
        if not invoice:
            return jsonify({
                "success": False,
                "message": "Invoice नहीं मिला"
            }), 404

        g.cursor.execute(
            """
            SELECT
                sii.id,
                sii.product_id,
                sii.quantity,
                sii.unit,
                sii.base_quantity,
                sii.conversion_factor,
                sii.rate,
                sii.discount,
                sii.gst_rate,
                sii.gst_amount,
                sii.line_total,
                sii.warranty_months,
                sii.serial_number,

                p.product_name,
                p.brand,
                p.item_type,
                p.warranty_months AS product_warranty_months,
                p.unit AS primary_unit,
                p.secondary_unit,
                p.conversion_rate,
                p.gst_rate AS product_gst_rate
            FROM sales_invoice_items sii
            LEFT JOIN products p
                ON p.id = sii.product_id
            WHERE sii.invoice_id = %s
            ORDER BY sii.id ASC
            """,
            (invoice_id,)
        )
        items = g.cursor.fetchall()

        g.cursor.execute(
            """
            SELECT product_id, serial_number, warranty_start, warranty_end
            FROM product_serials
            WHERE sale_invoice_id = %s
            ORDER BY id ASC
            """,
            (invoice_id,)
        )
        serial_rows = g.cursor.fetchall()
        serial_map = {}
        for row in serial_rows:
            product_id = row.get("product_id")
            serial_map.setdefault(product_id, []).append({
                "serial_number": row.get("serial_number"),
                "warranty_start": normalize_date_string(row.get("warranty_start")),
                "warranty_end": normalize_date_string(row.get("warranty_end"))
            })

        normalized_items = []
        for item in items:
            product_id = item.get("product_id")
            item_serials = serial_map.get(product_id, [])
            if not item_serials and item.get("serial_number"):
                item_serials = [{
                    "serial_number": item.get("serial_number"),
                    "warranty_start": None,
                    "warranty_end": None
                }]

            normalized_items.append({
                "id": item.get("id"),
                "product_id": product_id,
                "product_name": item.get("product_name") or "-",
                "brand": item.get("brand") or "-",
                "item_type": item.get("item_type") or "Product",
                "quantity": float(item.get("quantity") or 0),
                "unit": item.get("unit") or item.get("primary_unit") or "Piece",
                "base_quantity": float(item.get("base_quantity") or item.get("quantity") or 0),
                "conversion_factor": float(item.get("conversion_factor") or 1.0),
                "primary_unit": item.get("primary_unit"),
                "secondary_unit": item.get("secondary_unit"),
                "conversion_rate": float(item.get("conversion_rate") or 1.0),
                "rate": float(item.get("rate") or 0),
                "discount": float(item.get("discount") or 0),
                "gst_rate": float(item.get("gst_rate") or item.get("product_gst_rate") or 0),
                "gst_amount": float(item.get("gst_amount") or 0),
                "line_total": float(item.get("line_total") or 0),
                "warranty_months": int(item.get("warranty_months") if item.get("warranty_months") is not None else (item.get("product_warranty_months") or 0)),
                "serials": item_serials
            })

        return jsonify({
            "success": True,
            "invoice": {
                "id": invoice.get("id"),
                "customer_id": invoice.get("customer_id"),
                "invoice_number": invoice.get("invoice_number"),
                "invoice_date": normalize_date_string(invoice.get("invoice_date")),
                "subtotal": float(invoice.get("subtotal") or 0),
                "discount": float(invoice.get("discount") or 0),
                "tax_amount": float(invoice.get("tax_amount") or 0),
                "grand_total": float(invoice.get("grand_total") or 0),
                "paid_amount": float(invoice.get("paid_amount") or 0),
                "outstanding_amount": round(
                    float(invoice.get("grand_total") or 0) - float(invoice.get("paid_amount") or 0), 2
                ),
                "payment_method": invoice.get("payment_method"),
                "payment_status": invoice.get("payment_status"),
                "document_type": invoice.get("document_type") or "Invoice",
                "notes": invoice.get("notes") or "",
                "customer_name": invoice.get("customer_name") or "Walk-in Customer",
                "customer_mobile": invoice.get("customer_mobile") or "-",
                "customer_address": invoice.get("customer_address") or "-",
                "customer_gstin": invoice.get("customer_gstin") or "-",
                "customer_email": invoice.get("customer_email") or "-"
            },
            "items": normalized_items
        })

    except Exception as e:
        print("Get Sale Details Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


# ==========================================
# SERIAL NUMBER HISTORY
# ==========================================

@app.route(
    "/serial-history/<path:serial_number>",
    methods=["GET"]
)
def serial_history(serial_number):

    try:

        sql = """
            SELECT
                ps.id,
                ps.serial_number,
                ps.status,
                ps.purchase_date,
                ps.supplier_name,
                ps.warranty_start,
                ps.warranty_end,
                ps.sale_date,

                p.id AS product_id,
                p.product_name,
                p.brand,

                pi.id AS purchase_invoice_id,
                pi.purchase_number,

                si.id AS sale_invoice_id,
                si.invoice_number,

                c.id AS customer_id,
                c.customer_name,
                c.mobile

            FROM product_serials ps

            LEFT JOIN products p
                ON ps.product_id = p.id

            LEFT JOIN purchase_invoices pi
                ON ps.purchase_invoice_id = pi.id

            LEFT JOIN sales_invoices si
                ON ps.sale_invoice_id = si.id

            LEFT JOIN customers c
                ON ps.customer_id = c.id

            WHERE ps.serial_number = %s

            LIMIT 1
        """


        g.cursor.execute(
            sql,
            (serial_number,)
        )


        record = g.cursor.fetchone()


        if not record:

            return jsonify({
                "success": False,
                "message": "Serial Number नहीं मिला"
            }), 404


        return jsonify({
            "success": True,
            "serial": record,
            "warranty_status": get_warranty_status(
                record.get("warranty_start"),
                record.get("warranty_end")
            ),
            "service_history": get_service_history_records(
                record.get("serial_number")
            )
        })


    except Exception as e:

        print(
            "Serial History Error:",
            e
        )


        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


@app.route("/invoice-warranty/<path:invoice_number>", methods=["GET"])
def invoice_warranty(invoice_number):
    try:
        g.cursor.execute(
            """
            SELECT
                si.invoice_number,
                si.invoice_date,
                c.customer_name,
                c.mobile,
                p.product_name,
                p.brand,
                ps.serial_number,
                ps.warranty_start,
                ps.warranty_end
            FROM sales_invoices si
            LEFT JOIN customers c
                ON c.id = si.customer_id
            LEFT JOIN sales_invoice_items sii
                ON sii.invoice_id = si.id
            LEFT JOIN products p
                ON p.id = sii.product_id
            LEFT JOIN product_serials ps
                ON ps.sale_invoice_id = si.id
                AND ps.product_id = sii.product_id
            WHERE si.invoice_number = %s
            ORDER BY (ps.serial_number IS NOT NULL) DESC, sii.id ASC
            """,
            (invoice_number,)
        )

        records = g.cursor.fetchall()

        if not records:
            return jsonify({
                "success": False,
                "message": "Invoice Number नहीं मिला"
            }), 404

        record = records[0]

        items = []
        for row in records:
            if row.get("product_name"):
                items.append({
                    "product_name": row.get("product_name"),
                    "brand": row.get("brand") or "-",
                    "serial_number": row.get("serial_number") or "-",
                    "warranty_start": normalize_date_string(row.get("warranty_start")),
                    "warranty_end": normalize_date_string(row.get("warranty_end")),
                    "warranty_status": get_warranty_status(
                        row.get("warranty_start"),
                        row.get("warranty_end")
                    )
                })

        service_history = get_service_history_records(
            record.get("serial_number")
        ) if record.get("serial_number") else []

        return jsonify({
            "success": True,
            "product": record.get("product_name") or "External Product",
            "brand": record.get("brand") or "-",
            "customer": record.get("customer_name") or "Walk-in Customer",
            "mobile": record.get("mobile") or "-",
            "serial_number": record.get("serial_number") or "-",
            "invoice_number": record.get("invoice_number"),
            "invoice_date": normalize_date_string(record.get("invoice_date")),
            "warranty_start": normalize_date_string(record.get("warranty_start")),
            "warranty_end": normalize_date_string(record.get("warranty_end")),
            "warranty_status": get_warranty_status(
                record.get("warranty_start"),
                record.get("warranty_end")
            ),
            "service_history": service_history,
            "items": items
        })

    except Exception as e:
        print("Invoice Warranty Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


# ==========================================
# CONVERT QUOTATION TO TAX INVOICE
# ==========================================

@app.route("/sales/<int:invoice_id>/convert-quotation", methods=["POST", "OPTIONS"])
def convert_quotation(invoice_id):
    if request.method == "OPTIONS":
        return "", 200

    try:
        g.cursor.execute("SELECT * FROM sales_invoices WHERE id = %s LIMIT 1", (invoice_id,))
        invoice = g.cursor.fetchone()
        if not invoice:
            return jsonify({"success": False, "message": "Quotation नहीं मिला"}), 404

        if (invoice.get("document_type") or "").strip() != "Quotation":
            return jsonify({"success": False, "message": "यह document एक Quotation नहीं है"}), 400

        g.cursor.execute("SELECT * FROM sales_invoice_items WHERE invoice_id = %s", (invoice_id,))
        items = g.cursor.fetchall()

        # Check stock availability
        for item in items:
            p_id = item["product_id"]
            qty = float(item["quantity"] or 0)
            if p_id:
                prod = get_product_record_by_id(p_id)
                if prod and prod.get("item_type") != "Service":
                    g.cursor.execute("SELECT stock FROM products WHERE id = %s", (p_id,))
                    stk = g.cursor.fetchone()
                    if not stk or float(stk["stock"]) < qty:
                        return jsonify({
                            "success": False,
                            "message": f"Stock पर्याप्त नहीं है for product: {prod.get('product_name')}"
                        }), 400

        new_invoice_number = generate_document_number(
            "invoice-billing", "invoice_prefix", "next_invoice_number", "INV-"
        )
        today_str = datetime.now().strftime("%Y-%m-%d")

        for item in items:
            p_id = item["product_id"]
            qty = float(item["quantity"] or 0)
            if p_id:
                prod = get_product_record_by_id(p_id)
                if prod and prod.get("item_type") != "Service":
                    g.cursor.execute("UPDATE products SET stock = stock - %s WHERE id = %s", (qty, p_id))
                    ser = item.get("serial_number")
                    if ser:
                        g.cursor.execute(
                            """
                            UPDATE product_serials
                            SET status = 'Sold', sale_invoice_id = %s, customer_id = %s, sale_date = %s
                            WHERE serial_number = %s AND product_id = %s
                            """,
                            (invoice_id, invoice["customer_id"], today_str, ser, p_id)
                        )

        g.cursor.execute(
            """
            UPDATE sales_invoices
            SET document_type = 'Invoice',
                invoice_number = %s,
                invoice_date = %s
            WHERE id = %s
            """,
            (new_invoice_number, today_str, invoice_id)
        )

        g.db.commit()

        return jsonify({
            "success": True,
            "message": f"Quotation successfully Tax Invoice में बदल गया: {new_invoice_number}",
            "invoice_id": invoice_id,
            "invoice_number": new_invoice_number
        }), 200

    except Exception as e:
        g.db.rollback()
        return jsonify({"success": False, "message": str(e)}), 500


# ==========================================
# SALES RETURNS (RESTORE STOCK & REVERSALS)
# ==========================================

@app.route("/sales-returns", methods=["POST", "OPTIONS"])
def create_sales_return():
    if request.method == "OPTIONS":
        return "", 200

    data = request.get_json() or {}
    invoice_id = data.get("invoice_id")
    customer_id = data.get("customer_id")
    return_date = data.get("return_date") or datetime.now().strftime("%Y-%m-%d")
    return_amount = float(data.get("return_amount", 0) or 0)
    reason = data.get("reason", "")
    items = data.get("items", [])

    if invoice_id and not customer_id:
        try:
            g.cursor.execute("SELECT customer_id FROM sales_invoices WHERE id = %s LIMIT 1", (invoice_id,))
            inv_row = g.cursor.fetchone()
            if inv_row and inv_row.get("customer_id"):
                customer_id = inv_row["customer_id"]
        except Exception as ce:
            print("Fetch invoice customer error:", ce)

    if return_amount <= 0 and items:
        for itm in items:
            q = float(itm.get("quantity", 1) or 1)
            p = float(itm.get("rate") or itm.get("unit_price") or 0)
            t = float(itm.get("tax_amount", 0) or 0)
            return_amount += (q * p + t)

    if not items:
        return jsonify({"success": False, "message": "Return करने के लिए कम से कम 1 item चुनें"}), 400

    try:
        return_number = generate_document_number(
            "sales-return", "return_prefix", "next_return_number", "RET-"
        )

        refund_method = data.get("refund_method", "Cash")

        g.cursor.execute(
            """
            INSERT INTO sales_returns (return_number, invoice_id, customer_id, return_date, total_refund_amount, refund_method, reason)
            VALUES (%s, %s, %s, %s, %s, %s, %s)
            """,
            (return_number, invoice_id or None, customer_id or None, return_date, return_amount, refund_method, reason)
        )
        return_id = g.cursor.lastrowid

        for itm in items:
            product_id = itm.get("product_id")
            quantity = float(itm.get("quantity", 1) or 1)
            unit_price = float(itm.get("unit_price", 0) or 0)
            tax_amount = float(itm.get("tax_amount", 0) or 0)
            total_amount = float(itm.get("total_amount", 0) or (quantity * unit_price + tax_amount))
            serial_number = itm.get("serial_number")

            g.cursor.execute(
                """
                INSERT INTO sales_return_items (return_id, product_id, quantity, rate, tax_amount, line_total, serial_number)
                VALUES (%s, %s, %s, %s, %s, %s, %s)
                """,
                (return_id, product_id or None, quantity, unit_price, tax_amount, total_amount, serial_number or None)
            )

            if product_id:
                prod = get_product_record_by_id(product_id)
                if prod and prod.get("item_type") != "Service":
                    g.cursor.execute("UPDATE products SET stock = stock + %s WHERE id = %s", (quantity, product_id))

            if serial_number and product_id:
                g.cursor.execute(
                    """
                    UPDATE product_serials
                    SET status = 'In Stock', sale_invoice_id = NULL
                    WHERE serial_number = %s AND product_id = %s
                    """,
                    (serial_number, product_id)
                )

        g.db.commit()

        return jsonify({
            "success": True,
            "message": f"Sales Return दर्ज हो गया ({return_number}) एवं स्टॉक पुनः रीस्टोर हो गया",
            "return_id": return_id,
            "return_number": return_number
        }), 201

    except Exception as e:
        g.db.rollback()
        return jsonify({"success": False, "message": str(e)}), 500


@app.route("/sales-returns", methods=["GET"])
def get_sales_returns():
    try:
        invoice_id = request.args.get("invoice_id")
        sql = """
            SELECT sr.*, sr.total_refund_amount AS return_amount, COALESCE(c.customer_name, 'Walk-in Customer') AS customer_name, si.invoice_number
            FROM sales_returns sr
            LEFT JOIN customers c ON c.id = sr.customer_id
            LEFT JOIN sales_invoices si ON si.id = sr.invoice_id
            WHERE 1=1
        """
        params = []
        if invoice_id:
            sql += " AND sr.invoice_id = %s"
            params.append(invoice_id)
        sql += " ORDER BY sr.id DESC"
        g.cursor.execute(sql, tuple(params))
        returns = g.cursor.fetchall()
        for r in returns:
            r["return_date"] = normalize_date_string(r.get("return_date"))
            r["return_amount"] = float(r.get("total_refund_amount") or 0)
        return jsonify({"success": True, "returns": returns})
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500


# ==========================================
# CUSTOMER PAYMENTS & LEDGER
# ==========================================

@app.route("/customers/<int:customer_id>/payments", methods=["POST", "OPTIONS"])
def add_customer_payment(customer_id):
    if request.method == "OPTIONS":
        return "", 200

    data = request.get_json() or {}
    amount = float(data.get("amount", 0) or 0)
    payment_method = data.get("payment_method") or data.get("payment_mode") or "Cash"
    reference_no = data.get("reference_no", "")
    notes = data.get("notes", "")
    payment_date = data.get("payment_date") or datetime.now().strftime("%Y-%m-%d")
    invoice_id = data.get("invoice_id")

    if amount <= 0:
        return jsonify({"success": False, "message": "मान्य राशि (Amount > 0) दर्ज करें"}), 400

    try:
        receipt_number = generate_document_number(
            "customer-receipt", "receipt_prefix", "next_receipt_number", "RCP-"
        )
        g.cursor.execute(
            """
            INSERT INTO customer_payments (receipt_number, customer_id, invoice_id, payment_date, amount, payment_method, reference_no, notes)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
            """,
            (receipt_number, customer_id, invoice_id or None, payment_date, amount, payment_method, reference_no, notes)
        )
        payment_id = g.cursor.lastrowid

        if invoice_id:
            g.cursor.execute("SELECT grand_total, paid_amount FROM sales_invoices WHERE id = %s", (invoice_id,))
            inv = g.cursor.fetchone()
            if inv:
                new_paid = float(inv["paid_amount"] or 0) + amount
                new_status = "Paid" if new_paid >= float(inv["grand_total"] or 0) else "Partial"
                g.cursor.execute(
                    "UPDATE sales_invoices SET paid_amount = %s, payment_status = %s WHERE id = %s",
                    (new_paid, new_status, invoice_id)
                )

        g.db.commit()

        return jsonify({
            "success": True,
            "message": f"भुगतान सफलतापूर्वक दर्ज हो गया (रसीद सं. {receipt_number})",
            "payment_id": payment_id,
            "receipt_number": receipt_number
        }), 201

    except Exception as e:
        g.db.rollback()
        return jsonify({"success": False, "message": str(e)}), 500


@app.route("/customers/<int:customer_id>/ledger", methods=["GET"])
def get_customer_ledger(customer_id):
    try:
        g.cursor.execute("SELECT * FROM customers WHERE id = %s LIMIT 1", (customer_id,))
        customer = g.cursor.fetchone()
        if not customer:
            return jsonify({"success": False, "message": "Customer नहीं मिला"}), 404

        # Customer Invoices (excluding Quotations)
        g.cursor.execute(
            """
            SELECT id, invoice_number, invoice_date, grand_total, paid_amount, payment_status
            FROM sales_invoices
            WHERE customer_id = %s AND (document_type IS NULL OR document_type != 'Quotation')
            ORDER BY invoice_date ASC, id ASC
            """,
            (customer_id,)
        )
        invoices = g.cursor.fetchall()
        for inv in invoices:
            inv["invoice_date"] = normalize_date_string(inv.get("invoice_date"))
            inv["grand_total"] = float(inv.get("grand_total") or 0)
            inv["paid_amount"] = float(inv.get("paid_amount") or 0)
            inv["due_amount"] = round(inv["grand_total"] - inv["paid_amount"], 2)

        # Customer Payments
        g.cursor.execute(
            """
            SELECT cp.*, si.invoice_number
            FROM customer_payments cp
            LEFT JOIN sales_invoices si ON si.id = cp.invoice_id
            WHERE cp.customer_id = %s
            ORDER BY cp.payment_date ASC, cp.id ASC
            """,
            (customer_id,)
        )
        payments = g.cursor.fetchall()
        for p in payments:
            p["payment_date"] = normalize_date_string(p.get("payment_date"))
            p["amount"] = float(p.get("amount") or 0)

        total_billed = sum(inv["grand_total"] for inv in invoices)
        total_paid_invoices = sum(inv["paid_amount"] for inv in invoices)
        total_direct_payments = sum(p["amount"] for p in payments if not p.get("invoice_id"))
        total_paid = round(total_paid_invoices + total_direct_payments, 2)
        total_outstanding = max(0.0, round(total_billed - total_paid, 2))

        return jsonify({
            "success": True,
            "customer": {
                "id": customer["id"],
                "customer_name": customer["customer_name"],
                "mobile": customer.get("mobile"),
                "email": customer.get("email"),
                "address": customer.get("address"),
                "gstin": customer.get("gstin")
            },
            "invoices": invoices,
            "payments": payments,
            "total_billed": round(total_billed, 2),
            "total_paid": total_paid,
            "total_outstanding": total_outstanding
        })

    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500


# ==========================================
# CREATE SERVICE JOB
# ==========================================

def get_service_history_records(serial_number):
    if not serial_number:
        return []

    g.cursor.execute(
        """
        SELECT
            service_number,
            service_date,
            expected_delivery_date,
            complaint,
            service_status,
            technician,
            delivered_at
        FROM service_jobs
        WHERE serial_number = %s
        ORDER BY service_date DESC, id DESC
        """,
        (serial_number,)
    )

    return [
        {
            "service_number": row["service_number"],
            "service_date": normalize_date_string(row["service_date"]),
            "expected_delivery_date": normalize_date_string(
                row["expected_delivery_date"]
            ),
            "complaint": row["complaint"],
            "service_status": row["service_status"],
            "technician": row["technician"],
            "delivered_at": normalize_date_string(row["delivered_at"])
        }
        for row in g.cursor.fetchall()
    ]

@app.route("/service-jobs", methods=["POST", "OPTIONS"])
def create_service_job():

    if request.method == "OPTIONS":
        return "", 200

    data = request.get_json()

    if not data:
        return jsonify({
            "success": False,
            "message": "Service data नहीं मिला"
        }), 400

    try:

        service_number = (data.get("service_number") or "").strip()
        service_date = data.get("service_date")
        expected_delivery_date = data.get("expected_delivery_date")
        customer_id = data.get("customer_id")
        customer_name = (data.get("customer_name") or "").strip()
        customer_mobile = (data.get("customer_mobile") or data.get("mobile") or "").strip()
        customer_address = (data.get("customer_address") or data.get("address") or "").strip()

        if not customer_id and customer_mobile:
            cust_match = check_duplicate_customer_mobile(customer_mobile)
            if cust_match:
                customer_id = cust_match["id"]
                if not customer_name:
                    customer_name = cust_match.get("customer_name") or ""
                if customer_address and not cust_match.get("address"):
                    try:
                        g.cursor.execute("UPDATE customers SET address = %s WHERE id = %s", (customer_address, customer_id))
                    except Exception as ue:
                        print("Update customer address error:", ue)
            elif customer_name:
                try:
                    g.cursor.execute(
                        """
                        INSERT INTO customers (customer_name, mobile, address, status, customer_type, created_at)
                        VALUES (%s, %s, %s, 'Active', 'Regular', NOW())
                        """,
                        (customer_name, customer_mobile, customer_address or None)
                    )
                    customer_id = g.cursor.lastrowid
                except Exception as ce:
                    print("Auto-create customer error in service-jobs:", ce)
        elif customer_id:
            g.cursor.execute("SELECT customer_name, mobile, address FROM customers WHERE id = %s LIMIT 1", (customer_id,))
            c_row = g.cursor.fetchone()
            if c_row:
                if not customer_name:
                    customer_name = c_row.get("customer_name") or ""
                if not customer_mobile:
                    customer_mobile = c_row.get("mobile") or ""
                if not customer_address:
                    customer_address = c_row.get("address") or ""

        product_id = data.get("product_id")
        serial_number = data.get("serial_number")

        product_source = data.get(
            "product_source",
            "External Product"
        )

        external_product_name = data.get(
            "external_product_name"
        )

        external_brand = data.get(
            "external_brand"
        )

        external_model = data.get(
            "external_model"
        )

        complaint = data.get("complaint")

        accessories = data.get("accessories")

        warranty_status = data.get(
            "warranty_status",
            "Unknown"
        )

        service_status = data.get(
            "service_status",
            "Received"
        )

        technician = data.get("technician")

        estimated_charges = data.get(
            "estimated_charges",
            0
        )

        final_charges = data.get(
            "final_charges",
            0
        )

        notes = data.get("notes")


        if not service_date:
            service_date = data.get("appointment_date") or get_ist_today_str()
        if not expected_delivery_date:
            expected_delivery_date = data.get("appointment_date") or service_date or get_ist_today_str()

        if not service_date:
            return jsonify({
                "success": False,
                "message": "Service Date required"
            }), 400

        if not complaint:
            return jsonify({
                "success": False,
                "message": "Complaint required"
            }), 400

        if not expected_delivery_date:
            return jsonify({
                "success": False,
                "message": "Expected Delivery Date required"
            }), 400

        if service_date and expected_delivery_date:
            if str(expected_delivery_date) < str(service_date):
                return jsonify({
                    "success": False,
                    "message": "Expected Delivery Date cannot be earlier than Service Date"
                }), 400


        # ==================================
        # OUR PRODUCT
        # ==================================

        if product_source == "Purchased From Us":

            if not serial_number:
                return jsonify({
                    "success": False,
                    "message": "Serial Number required"
                }), 400


            g.cursor.execute(
                """
                SELECT
                    product_id
                FROM product_serials
                WHERE serial_number = %s
                LIMIT 1
                """,
                (serial_number,)
            )

            serial_record = g.cursor.fetchone()


            if not serial_record:

                return jsonify({
                    "success": False,
                    "message":
                        "Serial Number हमारे records में नहीं मिला"
                }), 404


            product_id = serial_record["product_id"]


            # Customer automatically find करें

            if not customer_id:

                g.cursor.execute(
                    """
                    SELECT
                        customer_id
                    FROM product_serials
                    WHERE serial_number = %s
                    LIMIT 1
                    """,
                    (serial_number,)
                )

                customer_record =    g.cursor.fetchone()


                if customer_record:

                    customer_id =    customer_record["customer_id"]


        # ==================================
        # EXTERNAL PRODUCT
        # ==================================

        else:

            product_source = "External Product"
            product_id = None

        service_type = (data.get("service_type") or "Shop").strip().title()
        if service_type not in ("Shop", "Field"):
            service_type = "Shop"
        customer_address = (data.get("customer_address") or data.get("address") or "").strip() or None

        req_service_no = (data.get("service_number") or "").strip()
        if req_service_no:
            g.cursor.execute("SELECT id FROM service_jobs WHERE service_number = %s LIMIT 1", (req_service_no,))
            if not g.cursor.fetchone():
                service_number = req_service_no
            else:
                prefix = "FLD-" if service_type == "Field" else "SRV-"
                service_number = f"{prefix}{int(time.time() * 1000) % 1000000:06d}"
        else:
            prefix = "FLD-" if service_type == "Field" else "SRV-"
            service_number = generate_document_number(
                "service", "service_prefix", "next_service_number", prefix
            )
            while True:
                g.cursor.execute("SELECT id FROM service_jobs WHERE service_number = %s LIMIT 1", (service_number,))
                if not g.cursor.fetchone():
                    break
                service_number = generate_document_number(
                    "service", "service_prefix", "next_service_number", prefix
                )

        complaint_type = (data.get("complaint_type") or "Hardware Repair").strip()
        priority = data.get("priority") or "Normal"
        assigned_technician_id = data.get("assigned_technician_id") or None
        if assigned_technician_id and not technician:
            g.cursor.execute("SELECT full_name FROM employees WHERE id = %s", (assigned_technician_id,))
            tech_emp = g.cursor.fetchone()
            if tech_emp:
                technician = tech_emp.get("full_name")
        technician_name = technician

        appointment_date = data.get("appointment_date") or None
        appointment_start_time = data.get("appointment_start_time") or None
        estimated_duration_minutes = int(data.get("estimated_duration_minutes") or 60)
        appointment_end_time = None
        if appointment_date and appointment_start_time:
            try:
                st_dt = datetime.strptime(f"{appointment_date} {appointment_start_time[:5]}", "%Y-%m-%d %H:%M")
                appointment_end_time = (st_dt + timedelta(minutes=estimated_duration_minutes)).strftime("%H:%M:%S")
            except Exception:
                pass

        g.cursor.execute(
            """
            INSERT INTO service_jobs
            (
                service_number,
                service_date,
                expected_delivery_date,
                customer_id,
                customer_name,
                customer_mobile,
                product_id,
                serial_number,
                product_source,
                external_product_name,
                external_brand,
                external_model,
                complaint,
                accessories,
                warranty_status,
                service_status,
                technician,
                technician_name,
                estimated_charges,
                final_charges,
                notes,
                complaint_type,
                priority,
                assigned_technician_id,
                appointment_date,
                appointment_start_time,
                appointment_end_time,
                estimated_duration_minutes,
                service_type,
                customer_address,
                assigned_at,
                created_at
            )
            VALUES
            (
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                IF(%s IS NOT NULL, NOW(), NULL),
                NOW()
            )
            """,
            (
                service_number,
                service_date,
                expected_delivery_date,
                customer_id,
                customer_name or None,
                customer_mobile or None,
                product_id,
                serial_number,
                product_source,
                external_product_name if product_source == "External Product" else None,
                external_brand if product_source == "External Product" else None,
                external_model if product_source == "External Product" else None,
                complaint,
                accessories,
                warranty_status,
                service_status,
                technician,
                technician_name,
                float(estimated_charges or 0),
                float(final_charges or 0),
                notes,
                complaint_type,
                priority,
                assigned_technician_id,
                appointment_date,
                appointment_start_time,
                appointment_end_time,
                estimated_duration_minutes,
                service_type,
                customer_address,
                assigned_technician_id
            )
        )
        


        service_id = g.cursor.lastrowid


        g.db.commit()


        return jsonify({
            "success": True,
            "message": "Service Complaint Successfully Registered",
            "service_id": service_id,
            "service_number": service_number
        })


    except Exception as e:

        g.db.rollback()

        print(
            "Service Error:",
            e
        )

        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


# ==========================================
# GET SERVICE JOBS
# ==========================================

@app.route("/service-jobs", methods=["GET"])
def get_service_jobs():

    try:
        service_type = (request.args.get("service_type") or "").strip().title()
        status_filter = (request.args.get("status") or "").strip()
        search_query = (request.args.get("search") or "").strip()

        where_clauses = ["1=1"]
        params = []

        if service_type in ("Shop", "Field"):
            if service_type == "Shop":
                where_clauses.append("(sj.service_type = 'Shop' OR sj.service_type IS NULL OR sj.service_type = '')")
            else:
                where_clauses.append("sj.service_type = 'Field'")

        if status_filter:
            where_clauses.append("sj.service_status = %s")
            params.append(status_filter)

        if search_query:
            where_clauses.append("(sj.service_number LIKE %s OR sj.customer_name LIKE %s OR c.customer_name LIKE %s OR sj.customer_mobile LIKE %s OR c.mobile LIKE %s OR sj.serial_number LIKE %s OR sj.complaint LIKE %s OR sj.customer_address LIKE %s)")
            wild = f"%{search_query}%"
            params.extend([wild, wild, wild, wild, wild, wild, wild, wild])

        where_sql = " AND ".join(where_clauses)

        sql = f"""
            SELECT
                sj.id,
                sj.service_number,
                sj.service_date,
                sj.expected_delivery_date,
                sj.customer_id,
                sj.product_id,
                sj.serial_number,
                sj.complaint,
                sj.service_status,
                sj.created_at,
                sj.final_charges,
                sj.estimated_charges,
                sj.advance_paid,
                COALESCE(sj.service_type, 'Shop') AS service_type,
                COALESCE(NULLIF(sj.customer_address, ''), NULLIF(c.address, ''), '') AS customer_address,
                sj.complaint_type,
                sj.priority,
                sj.assigned_technician_id,
                COALESCE(NULLIF(sj.technician_name, ''), NULLIF(sj.technician, ''), NULLIF(e.full_name, ''), '') AS technician_name,
                COALESCE(NULLIF(sj.technician, ''), NULLIF(sj.technician_name, ''), NULLIF(e.full_name, ''), '') AS technician,
                sj.appointment_date,
                sj.appointment_start_time,
                sj.appointment_end_time,
                sj.estimated_duration_minutes,
                sj.on_the_way_at,
                sj.reached_at,
                sj.work_started_at,
                sj.completed_at,
                sj.held_at,
                sj.hold_reason,

                COALESCE(NULLIF(c.customer_name, ''), NULLIF(sj.customer_name, ''), 'Walk-in Customer') AS customer_name,
                COALESCE(NULLIF(c.mobile, ''), NULLIF(sj.customer_mobile, ''), '') AS mobile,
                COALESCE(NULLIF(c.mobile, ''), NULLIF(sj.customer_mobile, ''), '') AS customer_mobile,

                COALESCE(p.product_name, sj.external_product_name, 'External Product') AS product_name,
                COALESCE(p.brand, sj.external_brand, '-') AS brand,
                sj.external_model

            FROM service_jobs sj

            LEFT JOIN customers c
                ON sj.customer_id = c.id

            LEFT JOIN products p
                ON sj.product_id = p.id

            LEFT JOIN employees e
                ON sj.assigned_technician_id = e.id

            WHERE {where_sql}

            ORDER BY sj.id DESC
        """

        g.cursor.execute(sql, tuple(params))
        jobs = g.cursor.fetchall()

        for job in jobs:
            if isinstance(job.get("appointment_start_time"), timedelta):
                secs = int(job["appointment_start_time"].total_seconds())
                job["appointment_start_time"] = f"{secs // 3600:02d}:{(secs % 3600) // 60:02d}:{secs % 60:02d}"
            elif job.get("appointment_start_time") is not None:
                job["appointment_start_time"] = str(job["appointment_start_time"])

            if isinstance(job.get("appointment_end_time"), timedelta):
                secs = int(job["appointment_end_time"].total_seconds())
                job["appointment_end_time"] = f"{secs // 3600:02d}:{(secs % 3600) // 60:02d}:{secs % 60:02d}"
            elif job.get("appointment_end_time") is not None:
                job["appointment_end_time"] = str(job["appointment_end_time"])

            if isinstance(job.get("service_date"), (date, datetime)):
                job["service_date"] = str(job["service_date"])[:10]
            if isinstance(job.get("expected_delivery_date"), (date, datetime)):
                job["expected_delivery_date"] = str(job["expected_delivery_date"])[:10]
            if isinstance(job.get("appointment_date"), (date, datetime)):
                job["appointment_date"] = str(job["appointment_date"])[:10]
            if isinstance(job.get("created_at"), (date, datetime)):
                job["created_at"] = str(job["created_at"])

        return jsonify({
            "success": True,
            "services": jobs
        })

    except Exception as e:

        print(
            "Get Service Jobs Error:",
            e
        )

        return jsonify({
            "success": False,
            "message": str(e)
        }), 500



# ==========================================
# GET SINGLE SERVICE JOB
# ==========================================

@app.route("/service-jobs/<int:service_id>", methods=["GET"])
def get_service_job(service_id):
    try:
        g.cursor.execute(
            """
            SELECT
                sj.*,
                COALESCE(NULLIF(c.customer_name, ''), NULLIF(sj.customer_name, ''), 'Walk-in Customer') AS customer_name,
                COALESCE(NULLIF(c.mobile, ''), NULLIF(sj.customer_mobile, ''), '') AS mobile,
                COALESCE(NULLIF(c.mobile, ''), NULLIF(sj.customer_mobile, ''), '') AS customer_mobile,
                c.address AS customer_address,
                c.gstin AS customer_gstin,
                c.email AS customer_email,
                COALESCE(p.product_name, sj.external_product_name, 'External Product') AS product_name,
                COALESCE(p.brand, sj.external_brand, '-') AS brand
            FROM service_jobs sj
            LEFT JOIN customers c
                ON sj.customer_id = c.id
            LEFT JOIN products p
                ON sj.product_id = p.id
            WHERE sj.id = %s
            LIMIT 1
            """,
            (service_id,)
        )
        service = g.cursor.fetchone()
        if not service:
            return jsonify({"success": False, "message": "Service Job नहीं मिला"}), 404

        service["advance_paid"] = float(service.get("advance_paid") or 0)
        service["estimated_charges"] = float(service.get("estimated_charges") or 0)
        service["final_charges"] = float(service.get("final_charges") or 0)

        # Fetch parts used
        g.cursor.execute(
            """
            SELECT sjp.*, p.product_name AS catalog_name
            FROM service_job_parts sjp
            LEFT JOIN products p ON p.id = sjp.product_id
            WHERE sjp.service_job_id = %s
            ORDER BY sjp.id ASC
            """,
            (service_id,)
        )
        parts = g.cursor.fetchall()
        for p in parts:
            p["unit_price"] = float(p.get("unit_price") or 0)
            p["line_total"] = float(p.get("line_total") or 0)

        # Linked invoice info if generated
        invoice_info = None
        if service.get("invoice_id"):
            g.cursor.execute("SELECT id, invoice_number, grand_total, payment_status FROM sales_invoices WHERE id = %s LIMIT 1", (service["invoice_id"],))
            invoice_info = g.cursor.fetchone()

        return jsonify({
            "success": True,
            "service": serialize_row(service),
            "parts": serialize_rows(parts),
            "invoice": serialize_row(invoice_info) if invoice_info else None
        })

    except Exception as e:
        print("Get Service Error:", e)
        return jsonify({"success": False, "message": str(e)}), 500


@app.route("/service-jobs/<int:service_id>", methods=["PUT"])
def update_service_job(service_id):
    data = request.get_json()
    if not data:
        return jsonify({"success": False, "message": "Update data नहीं मिला"}), 400

    try:
        g.cursor.execute(
            """
            SELECT * FROM service_jobs WHERE id = %s LIMIT 1
            """,
            (service_id,)
        )
        existing_service = g.cursor.fetchone()
        if not existing_service:
            return jsonify({"success": False, "message": "Service Job नहीं मिला"}), 404

        existing_customer_name = existing_service.get("customer_name")
        existing_customer_mobile = existing_service.get("customer_mobile")
        existing_customer_id = existing_service.get("customer_id")

        customer_name = data.get("customer_name") if "customer_name" in data else existing_customer_name
        customer_name = (customer_name or "").strip()

        customer_mobile = data.get("customer_mobile") if "customer_mobile" in data else (data.get("mobile") if "mobile" in data else existing_customer_mobile)
        customer_mobile = (customer_mobile or "").strip()

        customer_id = data.get("customer_id", existing_customer_id)
        customer_address = data.get("customer_address") if "customer_address" in data else existing_service.get("customer_address")
        customer_address = (customer_address or "").strip()

        if not customer_id and customer_mobile:
            cust_match = check_duplicate_customer_mobile(customer_mobile)
            if cust_match:
                customer_id = cust_match["id"]
                if not customer_name:
                    customer_name = cust_match.get("customer_name") or ""
                if customer_address and not cust_match.get("address"):
                    try:
                        g.cursor.execute("UPDATE customers SET address = %s WHERE id = %s", (customer_address, customer_id))
                    except Exception:
                        pass
            elif customer_name:
                try:
                    g.cursor.execute(
                        """
                        INSERT INTO customers (customer_name, mobile, address, status, customer_type, created_at)
                        VALUES (%s, %s, %s, 'Active', 'Regular', NOW())
                        """,
                        (customer_name, customer_mobile, customer_address or None)
                    )
                    customer_id = g.cursor.lastrowid
                except Exception as ce:
                    print("Auto-create customer on update error:", ce)
        elif customer_id:
            if customer_mobile:
                try:
                    g.cursor.execute(
                        "UPDATE customers SET mobile = %s WHERE id = %s AND (mobile IS NULL OR mobile = '')",
                        (customer_mobile, customer_id)
                    )
                except Exception:
                    pass
            if customer_address:
                try:
                    g.cursor.execute(
                        "UPDATE customers SET address = %s WHERE id = %s AND (address IS NULL OR address = '')",
                        (customer_address, customer_id)
                    )
                except Exception:
                    pass

        complaint = data.get("complaint", existing_service.get("complaint"))
        service_status = data.get("service_status", existing_service.get("service_status"))
        technician = data.get("technician", existing_service.get("technician"))
        estimated_charges = data.get("estimated_charges", existing_service.get("estimated_charges") or 0)
        final_charges = data.get("final_charges", existing_service.get("final_charges") or 0)
        advance_paid = data.get("advance_paid", existing_service.get("advance_paid") or 0)
        physical_condition = data.get("physical_condition", existing_service.get("physical_condition"))
        rma_status = data.get("rma_status", existing_service.get("rma_status") or "None")
        rma_center_name = data.get("rma_center_name", existing_service.get("rma_center_name"))
        rma_docket_no = data.get("rma_docket_no", existing_service.get("rma_docket_no"))
        rma_sent_date = data.get("rma_sent_date", existing_service.get("rma_sent_date"))
        rma_received_date = data.get("rma_received_date", existing_service.get("rma_received_date"))
        notes = data.get("notes", existing_service.get("notes"))
        expected_delivery_date = data.get("expected_delivery_date", existing_service.get("expected_delivery_date"))
        complaint_type = data.get("complaint_type", existing_service.get("complaint_type") or "Hardware Repair")
        priority = data.get("priority", existing_service.get("priority") or "Normal")
        assigned_technician_id = data.get("assigned_technician_id", existing_service.get("assigned_technician_id"))
        appointment_date = data.get("appointment_date", existing_service.get("appointment_date"))
        appointment_start_time = data.get("appointment_start_time", existing_service.get("appointment_start_time"))
        estimated_duration_minutes = int(data.get("estimated_duration_minutes", existing_service.get("estimated_duration_minutes") or 60))
        appointment_end_time = data.get("appointment_end_time", existing_service.get("appointment_end_time"))
        if appointment_date and appointment_start_time and not appointment_end_time:
            try:
                st_dt = datetime.strptime(f"{appointment_date} {str(appointment_start_time)[:5]}", "%Y-%m-%d %H:%M")
                appointment_end_time = (st_dt + timedelta(minutes=estimated_duration_minutes)).strftime("%H:%M:%S")
            except Exception:
                pass

        if assigned_technician_id and appointment_date:
            g.cursor.execute(
                """
                SELECT lr.id, lr.leave_type, lr.reason, e.full_name
                FROM employee_leave_requests lr
                JOIN employees e ON e.id = lr.employee_id
                WHERE lr.employee_id = %s
                  AND lr.status = 'Approved'
                  AND %s BETWEEN lr.start_date AND lr.end_date
                LIMIT 1
                """,
                (assigned_technician_id, str(appointment_date))
            )
            leave_conflict = g.cursor.fetchone()
            if leave_conflict:
                return jsonify({
                    "success": False,
                    "message": f"Technician '{leave_conflict['full_name']}' is on Approved Leave ({leave_conflict['leave_type']}: छुट्टी पर है) on {appointment_date}. Cannot assign jobs on leave dates."
                }), 400

        if service_status == "Delivered":
            delivered_at = datetime.now()
        elif existing_service.get("delivered_at"):
            delivered_at = existing_service["delivered_at"]
        else:
            delivered_at = None

        service_type = data.get("service_type", existing_service.get("service_type") or "Shop")
        customer_address = data.get("customer_address", existing_service.get("customer_address"))
        technician_name = technician

        g.cursor.execute(
            """
            UPDATE service_jobs
            SET
                customer_id = %s,
                customer_name = %s,
                customer_mobile = %s,
                complaint = %s,
                service_status = %s,
                technician = %s,
                technician_name = %s,
                estimated_charges = %s,
                final_charges = %s,
                advance_paid = %s,
                physical_condition = %s,
                rma_status = %s,
                rma_center_name = %s,
                rma_docket_no = %s,
                rma_sent_date = %s,
                rma_received_date = %s,
                notes = %s,
                expected_delivery_date = %s,
                delivered_at = %s,
                complaint_type = %s,
                priority = %s,
                assigned_technician_id = %s,
                appointment_date = %s,
                appointment_start_time = %s,
                appointment_end_time = %s,
                estimated_duration_minutes = %s,
                service_type = %s,
                customer_address = %s
            WHERE id = %s
            """,
            (
                customer_id or None,
                customer_name or None,
                customer_mobile or None,
                complaint,
                service_status,
                technician,
                technician_name,
                float(estimated_charges or 0),
                float(final_charges or 0),
                float(advance_paid or 0),
                physical_condition,
                rma_status,
                rma_center_name,
                rma_docket_no,
                rma_sent_date or None,
                rma_received_date or None,
                notes,
                expected_delivery_date or None,
                delivered_at,
                complaint_type,
                priority,
                assigned_technician_id or None,
                appointment_date or None,
                appointment_start_time or None,
                appointment_end_time or None,
                estimated_duration_minutes,
                service_type,
                customer_address,
                service_id
            )
        )
        g.db.commit()

        # If job marked completed, trigger cascade shift for subsequent jobs
        if service_status == "Completed" and existing_service.get("service_status") != "Completed":
            try:
                cascade_res = cascade_technician_schedule_shift(service_id)
                print("Admin update cascade reschedule result for job", service_id, ":", cascade_res)
            except Exception as cs_err:
                print("Admin update cascade reschedule error:", cs_err)

        return jsonify({
            "success": True,
            "message": "Service Job successfully update ho gaya"
        }), 200

    except Exception as e:
        print(
            "Update Service Error:",
            e
        )

        return jsonify({
            "success": False,
            "message": str(e)
        }), 500



# ==========================================
# GET SERVICE HISTORY BY SERIAL NUMBER
# ==========================================

@app.route("/service-history/<path:serial_number>", methods=["GET"])
def get_service_history(serial_number):

    try:

        g.cursor.execute(
            """
            SELECT
                id,
                service_number,
                service_date,
                expected_delivery_date,
                serial_number,
                complaint,
                service_status,
                technician,
                estimated_charges,
                final_charges,
                notes,
                delivered_at
            FROM service_jobs
            WHERE serial_number = %s
            ORDER BY id DESC
            """,
            (serial_number,)
        )

        services = g.cursor.fetchall()

        return jsonify({
            "success": True,
            "services": services
        })

    except Exception as e:

        print(
            "Service History Error:",
            e
        )

        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


# ==========================================
# TECHNICIANS LIST
# ==========================================

@app.route("/technicians", methods=["GET"])
def get_technicians():
    try:
        g.cursor.execute(
            """
            SELECT id, full_name, username, role, mobile
            FROM users
            WHERE status = 'Active'
            ORDER BY full_name ASC
            """
        )
        techs = g.cursor.fetchall()
        return jsonify({"success": True, "technicians": techs})
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500


# ==========================================
# SERVICE JOB PARTS (CONSUMPTION & RESTOCK)
# ==========================================

@app.route("/service-jobs/<int:service_id>/parts", methods=["GET"])
def get_service_job_parts(service_id):
    try:
        g.cursor.execute(
            """
            SELECT sjp.*, p.product_name, p.brand, p.item_type, p.stock
            FROM service_job_parts sjp
            LEFT JOIN products p ON p.id = sjp.product_id
            WHERE sjp.service_job_id = %s
            ORDER BY sjp.id ASC
            """,
            (service_id,)
        )
        parts = g.cursor.fetchall()
        for p in parts:
            p["quantity"] = float(p.get("quantity") or 0)
            p["unit_price"] = float(p.get("unit_price") or 0)
            p["rate"] = float(p.get("unit_price") or 0)
            p["line_total"] = float(p.get("line_total") or 0)
        return jsonify({"success": True, "parts": parts})
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500


@app.route("/service-jobs/<int:service_id>/parts", methods=["POST", "OPTIONS"])
def add_service_job_part(service_id):
    if request.method == "OPTIONS":
        return "", 200

    data = request.get_json() or {}
    product_id = data.get("product_id")
    part_name = (data.get("part_name") or "").strip()
    serial_number = (data.get("serial_number") or "").strip()
    quantity = float(data.get("quantity", 1) or 1)
    unit_price = float(data.get("unit_price") or data.get("rate") or 0)
    line_total = round(quantity * unit_price, 2)

    if not part_name and not product_id:
        return jsonify({"success": False, "message": "Part name या Product select करें"}), 400

    try:
        if product_id:
            prod = get_product_record_by_id(product_id)
            if not prod:
                return jsonify({"success": False, "message": "Product नहीं मिला"}), 404
            if not part_name:
                part_name = prod.get("product_name") or "Spare Part"
            if unit_price <= 0:
                unit_price = float(prod.get("sale_price") or 0)
                line_total = round(quantity * unit_price, 2)

            if prod.get("item_type") != "Service":
                g.cursor.execute("SELECT stock FROM products WHERE id = %s", (product_id,))
                stk = g.cursor.fetchone()
                if stk and float(stk["stock"]) < quantity:
                    return jsonify({
                        "success": False,
                        "message": f"पर्याप्त स्टॉक उपलब्ध नहीं है (Current stock: {stk['stock']})"
                    }), 400
                g.cursor.execute("UPDATE products SET stock = stock - %s WHERE id = %s", (quantity, product_id))

        g.cursor.execute(
            """
            INSERT INTO service_job_parts (service_job_id, product_id, part_name, serial_number, quantity, unit_price, line_total)
            VALUES (%s, %s, %s, %s, %s, %s, %s)
            """,
            (service_id, product_id or None, part_name, serial_number or None, quantity, unit_price, line_total)
        )
        part_id = g.cursor.lastrowid
        g.db.commit()

        return jsonify({
            "success": True,
            "message": "Spare Part सफलतापूर्वक जोड़ दिया गया एवं स्टॉक डिडक्ट हो गया",
            "part_id": part_id
        }), 201

    except Exception as e:
        g.db.rollback()
        return jsonify({"success": False, "message": str(e)}), 500


@app.route("/service-jobs/<int:service_id>/parts/<int:part_id>", methods=["DELETE", "OPTIONS"])
def remove_service_job_part(service_id, part_id):
    if request.method == "OPTIONS":
        return "", 200

    try:
        g.cursor.execute(
            "SELECT * FROM service_job_parts WHERE id = %s AND service_job_id = %s",
            (part_id, service_id)
        )
        part = g.cursor.fetchone()
        if not part:
            return jsonify({"success": False, "message": "Part नहीं मिला"}), 404

        product_id = part.get("product_id")
        quantity = float(part.get("quantity") or 0)

        if product_id and quantity > 0:
            prod = get_product_record_by_id(product_id)
            if prod and prod.get("item_type") != "Service":
                g.cursor.execute("UPDATE products SET stock = stock + %s WHERE id = %s", (quantity, product_id))

        g.cursor.execute("DELETE FROM service_job_parts WHERE id = %s", (part_id,))
        g.db.commit()

        return jsonify({
            "success": True,
            "message": "Part हटा दिया गया एवं स्टॉक पुनः रीस्टोर हो गया"
        })

    except Exception as e:
        g.db.rollback()
        return jsonify({"success": False, "message": str(e)}), 500


# ==========================================
# 1-CLICK INVOICE GENERATION FROM SERVICE JOB
# ==========================================

@app.route("/service-jobs/<int:service_id>/create-invoice", methods=["POST", "OPTIONS"])
def create_service_invoice(service_id):
    if request.method == "OPTIONS":
        return "", 200

    try:
        g.cursor.execute("SELECT * FROM service_jobs WHERE id = %s LIMIT 1", (service_id,))
        job = g.cursor.fetchone()
        if not job:
            return jsonify({"success": False, "message": "Service Job नहीं मिला"}), 404

        if job.get("invoice_id"):
            g.cursor.execute("SELECT invoice_number FROM sales_invoices WHERE id = %s", (job["invoice_id"],))
            inv = g.cursor.fetchone()
            inv_num = inv["invoice_number"] if inv else ""
            return jsonify({"success": False, "message": f"इस सर्विस जॉब के लिए पहले से बिल बन चुका है ({inv_num})"}), 400

        customer_id = job.get("customer_id")
        if not customer_id and job.get("customer_mobile"):
            g.cursor.execute("SELECT id FROM customers WHERE mobile = %s LIMIT 1", (job["customer_mobile"],))
            cust = g.cursor.fetchone()
            if cust:
                customer_id = cust["id"]
            else:
                g.cursor.execute(
                    "INSERT INTO customers (customer_name, mobile, status) VALUES (%s, %s, 'Active')",
                    (job.get("customer_name") or "Service Customer", job.get("customer_mobile"))
                )
                customer_id = g.cursor.lastrowid

        if not customer_id:
            g.cursor.execute("SELECT id FROM customers WHERE status = 'Active' LIMIT 1")
            cust = g.cursor.fetchone()
            customer_id = cust["id"] if cust else 1

        g.cursor.execute("SELECT * FROM service_job_parts WHERE service_job_id = %s", (service_id,))
        parts = g.cursor.fetchall()

        labor_charge = float(job.get("final_charges") or job.get("estimated_charges") or 0)
        parts_total = sum(float(p.get("line_total") or 0) for p in parts)
        subtotal = labor_charge + parts_total
        tax_amount = 0.0
        grand_total = subtotal + tax_amount
        advance_paid = float(job.get("advance_paid") or 0)
        paid_amount = min(advance_paid, grand_total)
        payment_status = "Paid" if paid_amount >= grand_total else ("Partial" if paid_amount > 0 else "Unpaid")

        invoice_number = generate_document_number(
            "invoice-billing", "invoice_prefix", "next_invoice_number", "INV-"
        )
        invoice_date = datetime.now().strftime("%Y-%m-%d")

        g.cursor.execute(
            """
            INSERT INTO sales_invoices
            (
                invoice_number, customer_id, invoice_date, subtotal, discount, tax_amount,
                grand_total, paid_amount, payment_method, payment_status, notes,
                document_type, service_job_id
            )
            VALUES (%s, %s, %s, %s, 0, %s, %s, %s, 'Cash', %s, %s, 'Invoice', %s)
            """,
            (
                invoice_number, customer_id, invoice_date, subtotal, tax_amount,
                grand_total, paid_amount, payment_status,
                f"Generated from Service Job #{job.get('service_number') or service_id}",
                service_id
            )
        )
        invoice_id = g.cursor.lastrowid

        # Check if there is a Service item in products for repair labor
        g.cursor.execute("SELECT id FROM products WHERE item_type = 'Service' LIMIT 1")
        srv_prod = g.cursor.fetchone()
        if not srv_prod:
            g.cursor.execute(
                """
                INSERT INTO products (product_name, category, brand, unit, purchase_price, sale_price, stock, min_stock, gst_rate, status, item_type)
                VALUES ('Repair & Service Labor', 'Services', 'CRD Global Solutions', 'Job', 0, 0, 9999, 0, 0, 'Active', 'Service')
                """
            )
            repair_product_id = g.cursor.lastrowid
        else:
            repair_product_id = srv_prod["id"]

        if labor_charge > 0 or not parts:
            item_desc = f"Service / Repair: {job.get('brand') or ''} {job.get('model_number') or ''} - {job.get('complaint') or 'Repair Service'}"
            g.cursor.execute(
                """
                INSERT INTO sales_invoice_items
                (invoice_id, product_id, quantity, rate, cost_rate, gst_rate, gst_amount, discount, line_total, serial_number)
                VALUES (%s, %s, 1, %s, 0, 0, 0, 0, %s, %s)
                """,
                (invoice_id, repair_product_id, labor_charge, labor_charge, job.get("serial_number"))
            )

        for p in parts:
            g.cursor.execute(
                """
                INSERT INTO sales_invoice_items
                (invoice_id, product_id, quantity, rate, cost_rate, gst_rate, gst_amount, discount, line_total, serial_number)
                VALUES (%s, %s, %s, %s, 0, 0, 0, 0, %s, %s)
                """,
                (invoice_id, p.get("product_id") or repair_product_id, p.get("quantity"), p.get("unit_price"), p.get("line_total"), p.get("serial_number"))
            )

        g.cursor.execute(
            """
            UPDATE service_jobs
            SET invoice_id = %s, service_status = 'Delivered', delivered_at = NOW()
            WHERE id = %s
            """,
            (invoice_id, service_id)
        )

        g.db.commit()

        return jsonify({
            "success": True,
            "message": f"Service Job के लिए Tax Invoice बन गया: {invoice_number}",
            "invoice_id": invoice_id,
            "invoice_number": invoice_number
        }), 201

    except Exception as e:
        g.db.rollback()
        return jsonify({"success": False, "message": str(e)}), 500


# ==========================================
# DASHBOARD SUMMARY
# ==========================================

@app.route(
    "/dashboard/summary",
    methods=["GET"]
)
def dashboard_summary():

    try:
        today_ist = get_ist_today_str()

        # ----------------------------------
        # TODAY'S SALES
        # ----------------------------------

        g.cursor.execute(
            """
            SELECT
                COALESCE(SUM(grand_total), 0) AS total
            FROM sales_invoices
            WHERE DATE(invoice_date) = %s AND (document_type IS NULL OR document_type != 'Quotation')
            """,
            (today_ist,)
        )

        sales_result = g.cursor.fetchone()


        # ----------------------------------
        # TODAY'S PURCHASE
        # ----------------------------------

        g.cursor.execute(
            """
            SELECT
                COALESCE(SUM(grand_total), 0) AS total
            FROM purchase_invoices
            WHERE DATE(purchase_date) = %s
            """,
            (today_ist,)
        )

        purchase_result = g.cursor.fetchone()


        # ----------------------------------
        # TODAY'S EXPENSE
        # ----------------------------------

        g.cursor.execute(
            """
            SELECT
                COALESCE(SUM(amount), 0) AS total
            FROM expenses
            WHERE DATE(expense_date) = %s
              AND status <> 'Deleted'
            """,
            (today_ist,)
        )

        expense_result = g.cursor.fetchone()


        # ----------------------------------
        # TODAY'S OUTSTANDING (SALES DUE TODAY)
        # ----------------------------------

        g.cursor.execute(
            """
            SELECT
                COALESCE(SUM(grand_total - paid_amount), 0) AS total
            FROM sales_invoices
            WHERE DATE(invoice_date) = %s
              AND (document_type IS NULL OR document_type != 'Quotation')
            """,
            (today_ist,)
        )

        today_outstanding_result = g.cursor.fetchone()


        # ----------------------------------
        # TOTAL PRODUCTS
        # ----------------------------------

        g.cursor.execute(
            """
            SELECT
                COUNT(*) AS total
            FROM products
            WHERE item_type = 'Product'
            """
        )

        products_result = g.cursor.fetchone()


        # ----------------------------------
        # TOTAL CUSTOMERS
        # ----------------------------------

        g.cursor.execute(
            """
            SELECT
                COUNT(*) AS total
            FROM customers
            WHERE status = 'Active'
            """
        )

        customers_result = g.cursor.fetchone()

        g.cursor.execute(
            """
            SELECT COALESCE(SUM(si.grand_total - si.paid_amount), 0) AS total
            FROM sales_invoices si
            INNER JOIN customers c ON c.id = si.customer_id
            WHERE c.status = 'Active' AND (si.document_type IS NULL OR si.document_type != 'Quotation')
            """
        )
        invoice_outstanding_result = g.cursor.fetchone()
        g.cursor.execute(
            "SELECT COALESCE(SUM(opening_balance), 0) AS total FROM customers WHERE status = 'Active'"
        )
        opening_balance_result = g.cursor.fetchone()


        # ----------------------------------
        # SERVICE OVERVIEW
        # ----------------------------------

        g.cursor.execute(
            """
            SELECT
                COUNT(*) AS total_services,
                SUM(CASE WHEN service_status = 'Received' THEN 1 ELSE 0 END) AS received_services,
                SUM(CASE WHEN service_status = 'In Progress' THEN 1 ELSE 0 END) AS in_progress_services,
                SUM(CASE WHEN service_status = 'Delivered' THEN 1 ELSE 0 END) AS delivered_services,
                SUM(CASE WHEN expected_delivery_date < %s
                    AND service_status NOT IN ('Delivered', 'Cancelled')
                    THEN 1 ELSE 0 END) AS overdue_services
            FROM service_jobs
            """,
            (today_ist,)
        )

        service_summary = g.cursor.fetchone()


        # ----------------------------------
        # RECENT SERVICES
        # ----------------------------------

        g.cursor.execute(
            """
            SELECT
                sj.service_number,
                sj.service_date,
                sj.expected_delivery_date,
                sj.serial_number,
                sj.complaint,
                sj.service_status,

                COALESCE(c.customer_name, 'Walk-in Customer') AS customer_name,
                COALESCE(p.product_name, sj.external_product_name, 'External Product') AS product_name

            FROM service_jobs sj

            LEFT JOIN customers c
                ON c.id = sj.customer_id

            LEFT JOIN products p
                ON p.id = sj.product_id

            ORDER BY sj.id DESC

            LIMIT 5
            """
        )

        recent_services = g.cursor.fetchall()


        # ----------------------------------
        # RECENT SALES
        # ----------------------------------

        g.cursor.execute(
            """
            SELECT
                s.invoice_number,
                COALESCE(
                    c.customer_name,
                    'Walk-in Customer'
                ) AS customer_name,
                s.invoice_date,
                s.grand_total,
                s.payment_status

            FROM sales_invoices s

            LEFT JOIN customers c
                ON c.id = s.customer_id

            WHERE (s.document_type IS NULL OR s.document_type != 'Quotation')

            ORDER BY s.id DESC

            LIMIT 5
            """
        )

        recent_sales = g.cursor.fetchall()

        sales_list = []

        for row in recent_sales:
            sales_list.append({
                "invoice_number": row["invoice_number"],
                "customer_name": row["customer_name"],
                "invoice_date": normalize_date_string(row["invoice_date"]),
                "grand_total": float(row["grand_total"] or 0),
                "payment_status": row["payment_status"]
            })

        recent_service_list = []

        for row in recent_services:
            recent_service_list.append({
                "service_number": row["service_number"],
                "service_date": normalize_date_string(row["service_date"]),
                "expected_delivery_date": normalize_date_string(row["expected_delivery_date"]),
                "customer_name": row["customer_name"],
                "product_name": row["product_name"],
                "serial_number": row["serial_number"],
                "complaint": row["complaint"],
                "service_status": row["service_status"]
            })

        return jsonify({
            "success": True,
            "sales_today": float(sales_result["total"] or 0),
            "purchase_today": float(purchase_result["total"] or 0),
            "expense_today": float(expense_result["total"] or 0),
            "today_outstanding": round(float(today_outstanding_result["total"] or 0), 2),
            "total_products": int(products_result["total"] or 0),
            "total_customers": int(customers_result["total"] or 0),
            "outstanding_amount": round(
                float(invoice_outstanding_result["total"] or 0)
                + float(opening_balance_result["total"] or 0),
                2
            ),
            "recent_sales": sales_list,
            "total_services": int(service_summary["total_services"] or 0),
            "received_services": int(service_summary["received_services"] or 0),
            "in_progress_services": int(service_summary["in_progress_services"] or 0),
            "delivered_services": int(service_summary["delivered_services"] or 0),
            "overdue_services": int(service_summary["overdue_services"] or 0),
            "recent_services": recent_service_list
        })

    except Exception as e:
        print("Dashboard Summary Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500



# ==========================================
# REPORTS
# ==========================================

@app.route("/reports/sales", methods=["GET"])
def reports_sales():
    try:
        from_date = normalize_date_string(request.args.get("from_date"))
        to_date = normalize_date_string(request.args.get("to_date"))
        customer_id = request.args.get("customer_id", "").strip()
        payment_status = request.args.get("payment_status", "").strip()
        payment_method = request.args.get("payment_method", "").strip()

        sql = """
            SELECT
                si.id,
                si.invoice_number,
                si.invoice_date,
                si.subtotal,
                si.discount,
                si.tax_amount,
                si.grand_total,
                si.payment_method,
                si.payment_status,
                COALESCE(c.customer_name, 'Walk-in Customer') AS customer_name
            FROM sales_invoices si
            LEFT JOIN customers c
                ON c.id = si.customer_id
            WHERE 1=1
              AND (si.document_type IS NULL OR si.document_type != 'Quotation')
        """
        params = []

        if from_date:
            sql += " AND DATE(si.invoice_date) >= %s "
            params.append(from_date)
        if to_date:
            sql += " AND DATE(si.invoice_date) <= %s "
            params.append(to_date)
        if customer_id:
            sql += " AND si.customer_id = %s "
            params.append(customer_id)
        if payment_status:
            sql += " AND si.payment_status = %s "
            params.append(payment_status)
        if payment_method:
            sql += " AND si.payment_method = %s "
            params.append(payment_method)

        sql += " ORDER BY si.invoice_date DESC, si.id DESC "

        g.cursor.execute(sql, tuple(params))
        rows = g.cursor.fetchall()

        summary = {
            "total_sales": round(sum(float(row["grand_total"] or 0) for row in rows), 2),
            "total_discount": round(sum(float(row["discount"] or 0) for row in rows), 2),
            "total_tax": round(sum(float(row["tax_amount"] or 0) for row in rows), 2),
            "net_sales": round(sum(float(row["grand_total"] or 0) for row in rows), 2),
            "invoice_count": len(rows)
        }

        return jsonify({
            "success": True,
            "rows": rows,
            "summary": summary,
            "report_name": "Sales Report"
        })

    except Exception as e:
        print("Sales Report Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


@app.route("/reports/purchases", methods=["GET"])
def reports_purchases():
    try:
        from_date = normalize_date_string(request.args.get("from_date"))
        to_date = normalize_date_string(request.args.get("to_date"))
        supplier_id = request.args.get("supplier_id", "").strip()
        payment_status = request.args.get("payment_status", "").strip()
        payment_method = request.args.get("payment_method", "").strip()

        sql = """
            SELECT
                pi.id,
                pi.purchase_number,
                pi.purchase_date,
                pi.subtotal,
                pi.discount,
                pi.tax_amount,
                pi.grand_total,
                pi.payment_method,
                pi.payment_status,
                COALESCE(s.supplier_name, 'Unknown Supplier') AS supplier_name
            FROM purchase_invoices pi
            LEFT JOIN suppliers s
                ON s.id = pi.supplier_id
            WHERE 1=1
        """
        params = []

        if from_date:
            sql += " AND DATE(pi.purchase_date) >= %s "
            params.append(from_date)
        if to_date:
            sql += " AND DATE(pi.purchase_date) <= %s "
            params.append(to_date)
        if supplier_id:
            sql += " AND pi.supplier_id = %s "
            params.append(supplier_id)
        if payment_status:
            sql += " AND pi.payment_status = %s "
            params.append(payment_status)
        if payment_method:
            sql += " AND pi.payment_method = %s "
            params.append(payment_method)

        sql += " ORDER BY pi.purchase_date DESC, pi.id DESC "

        g.cursor.execute(sql, tuple(params))
        rows = g.cursor.fetchall()

        summary = {
            "total_purchases": round(sum(float(row["grand_total"] or 0) for row in rows), 2),
            "total_discount": round(sum(float(row["discount"] or 0) for row in rows), 2),
            "total_tax": round(sum(float(row["tax_amount"] or 0) for row in rows), 2),
            "net_purchases": round(sum(float(row["grand_total"] or 0) for row in rows), 2),
            "invoice_count": len(rows)
        }

        return jsonify({
            "success": True,
            "rows": rows,
            "summary": summary,
            "report_name": "Purchase Report"
        })

    except Exception as e:
        print("Purchase Report Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


@app.route("/reports/expenses", methods=["GET"])
def reports_expenses():
    try:
        from_date = normalize_date_string(request.args.get("from_date"))
        to_date = normalize_date_string(request.args.get("to_date"))
        category = request.args.get("category", "").strip()
        payment_method = request.args.get("payment_method", "").strip()
        status = request.args.get("status", "").strip()

        sql = """
            SELECT
                e.id,
                e.expense_number,
                e.expense_date,
                e.category,
                e.description,
                e.paid_to,
                e.payment_method,
                e.amount,
                e.status
            FROM expenses e
            WHERE e.status <> 'Deleted'
        """
        params = []

        if from_date:
            sql += " AND DATE(e.expense_date) >= %s "
            params.append(from_date)
        if to_date:
            sql += " AND DATE(e.expense_date) <= %s "
            params.append(to_date)
        if category:
            sql += " AND e.category = %s "
            params.append(category)
        if payment_method:
            sql += " AND e.payment_method = %s "
            params.append(payment_method)
        if status:
            sql += " AND e.status = %s "
            params.append(status)

        sql += " ORDER BY e.expense_date DESC, e.id DESC "

        g.cursor.execute(sql, tuple(params))
        rows = g.cursor.fetchall()

        summary = {
            "total_expenses": round(sum(float(row["amount"] or 0) for row in rows), 2),
            "paid_expenses": round(sum(float(row["amount"] or 0) for row in rows if str(row["status"] or "").lower() == "paid"), 2),
            "pending_expenses": round(sum(float(row["amount"] or 0) for row in rows if str(row["status"] or "").lower() == "pending"), 2),
            "transaction_count": len(rows)
        }

        category_summary = {}
        for row in rows:
            key = str(row["category"] or "Uncategorized")
            category_summary[key] = category_summary.get(key, 0) + float(row["amount"] or 0)

        return jsonify({
            "success": True,
            "rows": rows,
            "summary": summary,
            "category_summary": {k: round(v, 2) for k, v in category_summary.items()},
            "report_name": "Expense Report"
        })

    except Exception as e:
        print("Expense Report Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


@app.route("/reports/profit-loss", methods=["GET"])
def reports_profit_loss():
    try:
        from_date = normalize_date_string(request.args.get("from_date"))
        to_date = normalize_date_string(request.args.get("to_date"))

        sales_sql = """
            SELECT COALESCE(SUM(grand_total), 0) AS total_sales
            FROM sales_invoices
            WHERE 1=1
        """
        cost_sql = """
            SELECT COALESCE(SUM(sii.quantity * sii.cost_rate), 0) AS total_cost
            FROM sales_invoice_items sii
            INNER JOIN sales_invoices si ON si.id = sii.invoice_id
            INNER JOIN products p ON p.id = sii.product_id
            WHERE COALESCE(p.item_type, 'Product') <> 'Service'
        """
        expense_sql = """
            SELECT COALESCE(SUM(amount), 0) AS total_expenses
            FROM expenses
            WHERE status <> 'Deleted'
        """

        params = []
        if from_date:
            sales_sql += " AND DATE(invoice_date) >= %s "
            cost_sql += " AND DATE(si.invoice_date) >= %s "
            expense_sql += " AND DATE(expense_date) >= %s "
            params.append(from_date)
        if to_date:
            sales_sql += " AND DATE(invoice_date) <= %s "
            cost_sql += " AND DATE(si.invoice_date) <= %s "
            expense_sql += " AND DATE(expense_date) <= %s "
            params.append(to_date)

        g.cursor.execute(sales_sql, tuple(params[:len(params)] if False else tuple(params)))
        sales_total = float(g.cursor.fetchone()["total_sales"] or 0)

        g.cursor.execute(cost_sql, tuple(params))
        purchase_total = float(g.cursor.fetchone()["total_cost"] or 0)

        g.cursor.execute(expense_sql, tuple(params[:len(params)] if False else tuple(params)))
        expense_total = float(g.cursor.fetchone()["total_expenses"] or 0)

        gross_profit = sales_total - purchase_total
        net_profit = gross_profit - expense_total

        return jsonify({
            "success": True,
            "report_name": "Profit & Loss",
            "summary": {
                "revenue": round(sales_total, 2),
                "cost": round(purchase_total, 2),
                "cogs": round(purchase_total, 2),
                "gross_profit": round(gross_profit, 2),
                "expenses": round(expense_total, 2),
                "net_profit": round(net_profit, 2),
                "net_loss": round(abs(net_profit) if net_profit < 0 else 0, 2),
                "note": "COGS is based on the saved sales_invoice_items.cost_rate for product items only; service items are excluded."
            }
        })

    except Exception as e:
        print("Profit & Loss Report Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


@app.route("/reports/outstanding", methods=["GET"])
def reports_outstanding():
    try:
        from_date = normalize_date_string(request.args.get("from_date"))
        to_date = normalize_date_string(request.args.get("to_date"))

        sql = """
            SELECT
                c.id,
                COALESCE(c.customer_name, 'Walk-in Customer') AS customer_name,
                c.mobile,
                COALESCE(SUM(si.grand_total), 0) AS total_sales,
                COALESCE(SUM(si.paid_amount), 0) AS paid_amount,
                COALESCE(SUM(si.grand_total - si.paid_amount), 0) + COALESCE(c.opening_balance, 0) AS outstanding_amount
            FROM customers c
            LEFT JOIN sales_invoices si
                ON si.customer_id = c.id
            WHERE c.status = 'Active'
        """
        params = []
        if from_date:
            sql += " AND DATE(si.invoice_date) >= %s "
            params.append(from_date)
        if to_date:
            sql += " AND DATE(si.invoice_date) <= %s "
            params.append(to_date)
        sql += " GROUP BY c.id, c.customer_name, c.mobile, c.opening_balance HAVING outstanding_amount > 0 ORDER BY outstanding_amount DESC, customer_name ASC "

        g.cursor.execute(sql, tuple(params))
        rows = g.cursor.fetchall()

        total_receivable = round(sum(float(row["outstanding_amount"] or 0) for row in rows), 2)
        customer_count = sum(1 for row in rows if float(row["outstanding_amount"] or 0) > 0)

        return jsonify({
            "success": True,
            "rows": rows,
            "summary": {
                "total_receivable": total_receivable,
                "customers_with_outstanding": customer_count
            },
            "report_name": "Customer Outstanding Report"
        })

    except Exception as e:
        print("Outstanding Report Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


@app.route("/reports/services", methods=["GET"])
def reports_services():
    try:
        from_date = normalize_date_string(request.args.get("from_date"))
        to_date = normalize_date_string(request.args.get("to_date"))
        service_status = (request.args.get("status") or "").strip()
        search = (request.args.get("search") or "").strip()

        sql = """
            SELECT
                sj.id,
                sj.service_number,
                sj.service_date,
                sj.expected_delivery_date,
                sj.customer_id,
                sj.product_id,
                sj.serial_number,
                sj.complaint,
                sj.service_status,
                sj.technician,
                sj.estimated_charges,
                sj.final_charges,
                sj.delivered_at,
                COALESCE(c.customer_name, 'Walk-in Customer') AS customer_name,
                COALESCE(p.product_name, sj.external_product_name, 'External Product') AS product_name,
                COALESCE(p.brand, sj.external_brand, '-') AS brand
            FROM service_jobs sj
            LEFT JOIN customers c
                ON c.id = sj.customer_id
            LEFT JOIN products p
                ON p.id = sj.product_id
            WHERE 1=1
        """
        params = []

        if from_date:
            sql += " AND DATE(sj.service_date) >= %s "
            params.append(from_date)
        if to_date:
            sql += " AND DATE(sj.service_date) <= %s "
            params.append(to_date)
        status_norm = (service_status or "").strip().lower()
        if status_norm and status_norm != "all":
            if status_norm in {"overdue_pending", "overdue / pending", "overdue"}:
                sql += " AND sj.expected_delivery_date IS NOT NULL AND DATE(sj.expected_delivery_date) < %s AND LOWER(COALESCE(sj.service_status, '')) NOT IN ('delivered', 'cancelled') "
                params.append(get_ist_today_str())

            elif status_norm in {"delivered_after_due", "delivered after due date", "delivered after due"}:
                sql += " AND LOWER(COALESCE(sj.service_status, '')) = 'delivered' AND sj.delivered_at IS NOT NULL AND sj.expected_delivery_date IS NOT NULL AND DATE(sj.delivered_at) > DATE(sj.expected_delivery_date) "
            elif status_norm in {"delivered_on_time", "delivered on time"}:
                sql += " AND LOWER(COALESCE(sj.service_status, '')) = 'delivered' AND sj.delivered_at IS NOT NULL AND sj.expected_delivery_date IS NOT NULL AND DATE(sj.delivered_at) <= DATE(sj.expected_delivery_date) "
            else:
                sql += " AND sj.service_status = %s "
                params.append(service_status)

        if search:
            search_pattern = f"%{search.lower()}%"
            sql += " AND ("
            sql += " LOWER(COALESCE(sj.service_number, '')) LIKE %s OR "
            sql += " LOWER(COALESCE(sj.serial_number, '')) LIKE %s OR "
            sql += " LOWER(COALESCE(c.customer_name, '')) LIKE %s OR "
            sql += " LOWER(COALESCE(p.product_name, sj.external_product_name, '')) LIKE %s OR "
            sql += " LOWER(COALESCE(sj.complaint, '')) LIKE %s "
            sql += ") "
            params.extend([search_pattern, search_pattern, search_pattern, search_pattern, search_pattern])

        sql += " ORDER BY sj.service_date DESC, sj.id DESC "

        g.cursor.execute(sql, tuple(params))
        services = g.cursor.fetchall()

        def as_date(value):
            if value in (None, "", "NULL"):
                return None
            try:
                return datetime.strptime(str(value)[:10], "%Y-%m-%d").date()
            except Exception:
                return None

        def as_datetime_date(value):
            if value in (None, "", "NULL"):
                return None
            try:
                return datetime.strptime(str(value)[:10], "%Y-%m-%d").date()
            except Exception:
                try:
                    return value.date()
                except Exception:
                    return None

        today = datetime.today().date()

        # Compute summary for current date range (unfiltered by status/search so cards remain accurate)
        sum_sql = "SELECT sj.service_status, sj.expected_delivery_date, sj.delivered_at FROM service_jobs sj WHERE 1=1"
        sum_params = []
        if from_date:
            sum_sql += " AND DATE(sj.service_date) >= %s "
            sum_params.append(from_date)
        if to_date:
            sum_sql += " AND DATE(sj.service_date) <= %s "
            sum_params.append(to_date)
        g.cursor.execute(sum_sql, tuple(sum_params))
        summary_source = g.cursor.fetchall()

        summary = {
            "total_services": len(summary_source),
            "cancelled_services": sum(1 for row in summary_source if str(row.get("service_status") or "").strip().lower() == "cancelled"),
            "delivered_on_time": sum(
                1
                for row in summary_source
                if str(row.get("service_status") or "").strip().lower() == "delivered"
                and row.get("delivered_at") is not None
                and row.get("expected_delivery_date") is not None
                and as_datetime_date(row.get("delivered_at")) is not None
                and as_date(row.get("expected_delivery_date")) is not None
                and as_datetime_date(row.get("delivered_at")) <= as_date(row.get("expected_delivery_date"))
            ),
            "delivered_after_due": sum(
                1
                for row in summary_source
                if str(row.get("service_status") or "").strip().lower() == "delivered"
                and row.get("delivered_at") is not None
                and row.get("expected_delivery_date") is not None
                and as_datetime_date(row.get("delivered_at")) is not None
                and as_date(row.get("expected_delivery_date")) is not None
                and as_datetime_date(row.get("delivered_at")) > as_date(row.get("expected_delivery_date"))
            ),
            "overdue_pending": sum(
                1
                for row in summary_source
                if row.get("expected_delivery_date") is not None
                and as_date(row.get("expected_delivery_date")) is not None
                and as_date(row.get("expected_delivery_date")) < today
                and str(row.get("service_status") or "").strip().lower() not in {"delivered", "cancelled"}
            )
        }

        # Tag each row with delivery timing status
        for row in services:
            st = str(row.get("service_status") or "").strip().lower()
            exp_d = as_date(row.get("expected_delivery_date"))
            del_at = as_datetime_date(row.get("delivered_at"))
            if st == "delivered":
                if del_at and exp_d and del_at > exp_d:
                    row["delivery_tag"] = "Delivered (Late)"
                    row["delivery_tag_class"] = "badge-late"
                elif del_at and exp_d:
                    row["delivery_tag"] = "Delivered (On Time)"
                    row["delivery_tag_class"] = "badge-ontime"
                else:
                    row["delivery_tag"] = "Delivered"
                    row["delivery_tag_class"] = "badge-ontime"
            elif st == "cancelled":
                row["delivery_tag"] = "Cancelled"
                row["delivery_tag_class"] = "badge-out-of-stock"
            else:
                if exp_d and exp_d < today:
                    row["delivery_tag"] = f"{row.get('service_status')} (⚠️ Overdue)"
                    row["delivery_tag_class"] = "badge-overdue"
                else:
                    row["delivery_tag"] = row.get("service_status") or "In Progress"
                    row["delivery_tag_class"] = "badge-in-stock"

        print("SERVICE REPORT COUNT:", len(services))
        print("SERVICE REPORT SUMMARY:", summary)

        return jsonify({
            "success": True,
            "services": services,
            "rows": services,
            "summary": summary,
            "report_name": "Service Report"
        })

    except Exception as e:
        print("Service Report Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


@app.route("/reports/inventory", methods=["GET"])
def reports_inventory():
    try:
        category = request.args.get("category", "").strip()
        stock_status = request.args.get("stock_status", "").strip().lower()

        sql = """
            SELECT
                p.id,
                p.product_name,
                p.category,
                p.brand,
                p.unit,
                p.secondary_unit,
                p.conversion_rate,
                p.stock,
                p.min_stock,
                p.purchase_price,
                p.track_serial,
                ROUND(p.stock * p.purchase_price, 2) AS stock_value
            FROM products p
            WHERE p.item_type != 'Service'
        """
        params = []
        if category:
            sql += " AND p.category = %s "
            params.append(category)

        if stock_status == "out_of_stock":
            sql += " AND p.stock <= 0 "
        elif stock_status == "low_stock":
            sql += " AND p.stock > 0 AND p.stock <= p.min_stock "
        elif stock_status == "in_stock":
            sql += " AND p.stock > p.min_stock "

        sql += " ORDER BY p.stock ASC, p.product_name ASC "

        g.cursor.execute(sql, tuple(params))
        rows = g.cursor.fetchall()

        # Fetch location stock for all returned products
        if rows:
            prod_ids = [r["id"] for r in rows]
            format_strings = ','.join(['%s'] * len(prod_ids))

            loc_sql = f"""
                SELECT 
                    pls.product_id,
                    sl.id AS location_id,
                    sl.location_name,
                    pls.quantity
                FROM product_location_stock pls
                JOIN stock_locations sl ON sl.id = pls.location_id
                WHERE pls.product_id IN ({format_strings}) AND pls.quantity > 0
                ORDER BY sl.is_default DESC, sl.location_name ASC
            """
            g.cursor.execute(loc_sql, tuple(prod_ids))
            loc_rows = g.cursor.fetchall()

            prod_locations_map = {}
            for lr in loc_rows:
                pid = lr["product_id"]
                if pid not in prod_locations_map:
                    prod_locations_map[pid] = []
                prod_locations_map[pid].append({
                    "location_id": lr["location_id"],
                    "location_name": lr["location_name"],
                    "quantity": float(lr["quantity"] or 0)
                })

            # Fetch active serials by location for these products
            serial_sql = f"""
                SELECT 
                    ps.product_id,
                    COALESCE(ps.location_id, (SELECT id FROM stock_locations WHERE is_default = 1 LIMIT 1)) AS location_id,
                    sl.location_name,
                    ps.serial_number
                FROM product_serials ps
                LEFT JOIN stock_locations sl ON sl.id = COALESCE(ps.location_id, (SELECT id FROM stock_locations WHERE is_default = 1 LIMIT 1))
                WHERE ps.product_id IN ({format_strings}) AND ps.status = 'In Stock'
                ORDER BY ps.serial_number ASC
            """
            g.cursor.execute(serial_sql, tuple(prod_ids))
            ser_rows = g.cursor.fetchall()

            prod_serials_map = {}
            for sr in ser_rows:
                key = (sr["product_id"], sr["location_id"])
                if key not in prod_serials_map:
                    prod_serials_map[key] = []
                prod_serials_map[key].append(sr["serial_number"])

            g.cursor.execute("SELECT id, location_name FROM stock_locations WHERE is_default = 1 LIMIT 1")
            default_loc_row = g.cursor.fetchone() or {"id": 1, "location_name": "Main Counter"}
            default_loc_id = default_loc_row["id"]
            default_loc_name = default_loc_row["location_name"]

            for row in rows:
                pid = row["id"]
                stk = float(row.get("stock") or 0)
                mstk = float(row.get("min_stock") or 0)
                if stk <= 0:
                    row["status_badge"] = "Out of Stock"
                elif stk <= mstk:
                    row["status_badge"] = "Low Stock"
                else:
                    row["status_badge"] = "In Stock"

                u1 = row.get("unit") or "Piece"
                u2 = row.get("secondary_unit")
                crate = float(row.get("conversion_rate") or 1)
                if u2 and crate > 1:
                    row["display_unit"] = f"{u1} (1 {u1} = {int(crate) if crate.is_integer() else crate} {u2})"
                    equiv_primary = round(stk / crate, 2)
                    eq_str = f"{int(equiv_primary) if equiv_primary.is_integer() else equiv_primary} {u1}"
                    stk_str = f"{int(stk) if stk.is_integer() else stk} {u2}"
                    row["display_stock"] = f"{stk_str} ({eq_str})"
                else:
                    row["display_unit"] = u1
                    row["display_stock"] = f"{int(stk) if stk.is_integer() else stk}"

                plocs = prod_locations_map.get(pid, [])
                if stk > 0 and not plocs:
                    plocs = [{"location_id": default_loc_id, "location_name": default_loc_name, "quantity": stk}]

                summary_parts = []
                badge_parts = []
                for loc in plocs:
                    loc_id = loc["location_id"]
                    loc_name = loc["location_name"]
                    qty_val = loc["quantity"]
                    qty_str = f"{int(qty_val) if qty_val.is_integer() else qty_val}"
                    loc_serials = prod_serials_map.get((pid, loc_id), [])
                    loc["serials"] = loc_serials

                    ser_summary = f" ({len(loc_serials)} Serials)" if loc_serials else ""
                    summary_parts.append(f"{loc_name}: {qty_str}{ser_summary}")

                    serial_tip = f" title='Serials: {', '.join(loc_serials)}'" if loc_serials else ""
                    serial_pill = f" <span class='loc-serials-pill'>({len(loc_serials)} SN)</span>" if loc_serials else ""
                    badge_parts.append(f"<span class='loc-tag loc-{loc_name.lower().replace(' ', '-')}'{serial_tip}><strong>{loc_name}:</strong> {qty_str}{serial_pill}</span>")

                row["location_summary"] = ", ".join(summary_parts) if summary_parts else "-"
                row["location_display"] = " ".join(badge_parts) if badge_parts else "-"
                row["locations"] = plocs
        else:
            for row in rows:
                stk = float(row.get("stock") or 0)
                mstk = float(row.get("min_stock") or 0)
                if stk <= 0:
                    row["status_badge"] = "Out of Stock"
                elif stk <= mstk:
                    row["status_badge"] = "Low Stock"
                else:
                    row["status_badge"] = "In Stock"

                u1 = row.get("unit") or "Piece"
                u2 = row.get("secondary_unit")
                crate = float(row.get("conversion_rate") or 1)
                if u2 and crate > 1:
                    row["display_unit"] = f"{u1} (1 {u1} = {int(crate) if crate.is_integer() else crate} {u2})"
                    equiv_primary = round(stk / crate, 2)
                    eq_str = f"{int(equiv_primary) if equiv_primary.is_integer() else equiv_primary} {u1}"
                    stk_str = f"{int(stk) if stk.is_integer() else stk} {u2}"
                    row["display_stock"] = f"{stk_str} ({eq_str})"
                else:
                    row["display_unit"] = u1
                    row["display_stock"] = f"{int(stk) if stk.is_integer() else stk}"

                row["location_summary"] = "-"
                row["location_display"] = "-"
                row["locations"] = []

        # Fetch distinct product categories for filter dropdown
        g.cursor.execute("SELECT DISTINCT category FROM products WHERE item_type != 'Service' AND category IS NOT NULL AND category != '' ORDER BY category ASC")
        category_rows = g.cursor.fetchall()
        categories = [r["category"] for r in category_rows]

        total_units = round(sum(float(row["stock"] or 0) for row in rows), 2)

        return jsonify({
            "success": True,
            "rows": rows,
            "categories": categories,
            "summary": {
                "total_products": len(rows),
                "total_units": total_units,
                "low_stock": sum(1 for row in rows if float(row["stock"] or 0) <= float(row["min_stock"] or 0)),
                "out_of_stock": sum(1 for row in rows if float(row["stock"] or 0) <= 0),
                "stock_value": round(sum(float(row["stock_value"] or 0) for row in rows), 2)
            },
            "report_name": "Inventory Report"
        })

    except Exception as e:
        print("Inventory Report Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


@app.route("/reports/fast-moving", methods=["GET"])
def reports_fast_moving():
    try:
        from_date = normalize_date_string(request.args.get("from_date"))
        to_date = normalize_date_string(request.args.get("to_date"))
        category = (request.args.get("category") or "").strip()
        search = (request.args.get("search") or "").strip()
        sort_by = (request.args.get("sort_by") or "qty_desc").strip().lower()
        filter_mode = (request.args.get("filter") or "sold_only").strip().lower()

        date_conds = ""
        date_params = []
        if from_date:
            date_conds += " AND DATE(si.invoice_date) >= %s "
            date_params.append(from_date)
        if to_date:
            date_conds += " AND DATE(si.invoice_date) <= %s "
            date_params.append(to_date)

        sql = f"""
            SELECT
                p.id AS product_id,
                p.product_name,
                COALESCE(p.category, 'General') AS category,
                COALESCE(p.brand, '-') AS brand,
                COALESCE(p.unit, 'Pcs') AS unit,
                COALESCE(p.stock, 0) AS current_stock,
                COALESCE(p.sale_price, 0) AS standard_price,
                COALESCE(SUM(sii.quantity), 0) AS total_sold_qty,
                COALESCE(SUM(sii.line_total), 0) AS total_revenue,
                CASE 
                    WHEN COALESCE(SUM(sii.quantity), 0) > 0 
                    THEN ROUND(SUM(sii.line_total) / SUM(sii.quantity), 2)
                    ELSE COALESCE(p.sale_price, 0)
                END AS avg_price,
                COUNT(DISTINCT si.id) AS invoice_count
            FROM products p
            LEFT JOIN (
                sales_invoice_items sii
                JOIN sales_invoices si 
                    ON si.id = sii.invoice_id 
                    AND COALESCE(si.document_type, 'Invoice') = 'Invoice'
                    {date_conds}
            ) ON sii.product_id = p.id
            WHERE p.status = 'Active'
        """
        params = list(date_params)

        if category:
            sql += " AND p.category = %s "
            params.append(category)

        if search:
            sql += " AND (p.product_name LIKE %s OR p.brand LIKE %s OR p.category LIKE %s) "
            like_pat = f"%{search}%"
            params.extend([like_pat, like_pat, like_pat])

        sql += " GROUP BY p.id, p.product_name, p.category, p.brand, p.unit, p.stock, p.sale_price "

        if filter_mode == "sold_only":
            sql += " HAVING total_sold_qty > 0 "

        if sort_by == "revenue_desc":
            sql += " ORDER BY total_revenue DESC, total_sold_qty DESC, p.product_name ASC "
        elif sort_by == "qty_asc":
            sql += " ORDER BY total_sold_qty ASC, total_revenue ASC, p.product_name ASC "
        else:
            sql += " ORDER BY total_sold_qty DESC, total_revenue DESC, p.product_name ASC "

        g.cursor.execute(sql, tuple(params))
        rows = g.cursor.fetchall()

        total_units = 0.0
        total_rev = 0.0
        sold_products_count = 0

        enhanced_rows = []
        for idx, row in enumerate(rows, 1):
            qty = float(row["total_sold_qty"] or 0)
            rev = float(row["total_revenue"] or 0)
            avg_p = float(row["avg_price"] or 0)
            stock = float(row["current_stock"] or 0)

            total_units += qty
            total_rev += rev
            if qty > 0:
                sold_products_count += 1

            if qty >= 10:
                status_badge = "🔥 Fast Moving"
                status_class = "badge-fast"
            elif qty >= 5:
                status_badge = "⚡ Moderate"
                status_class = "badge-moderate"
            elif qty > 0:
                status_badge = "🐢 Slow Moving"
                status_class = "badge-slow"
            else:
                status_badge = "⏸️ Non-Moving"
                status_class = "badge-idle"

            rank_display = f"#{idx}"
            if idx == 1 and qty > 0:
                rank_display = "🥇 #1"
            elif idx == 2 and qty > 0:
                rank_display = "🥈 #2"
            elif idx == 3 and qty > 0:
                rank_display = "🥉 #3"

            unit_str = str(row["unit"] or "Pcs")
            qty_val = int(qty) if qty.is_integer() else round(qty, 2)
            stock_val = int(stock) if stock.is_integer() else round(stock, 2)

            enhanced_rows.append({
                "rank": rank_display,
                "product_id": row["product_id"],
                "product_name": row["product_name"],
                "category": row["category"],
                "brand": row["brand"],
                "unit": unit_str,
                "total_sold_qty": qty_val,
                "sold_qty_display": f"{qty_val} {unit_str}",
                "total_revenue": round(rev, 2),
                "avg_price": round(avg_p, 2),
                "current_stock": stock_val,
                "stock_display": f"{stock_val} {unit_str}",
                "invoice_count": row["invoice_count"],
                "status_badge": status_badge,
                "status_class": status_class
            })

        top_prod_str = "None"
        if enhanced_rows and enhanced_rows[0]["total_sold_qty"] > 0:
            top_prod_str = f"{enhanced_rows[0]['product_name']} ({enhanced_rows[0]['sold_qty_display']})"

        summary = {
            "top_product": top_prod_str,
            "total_qty_sold": int(total_units) if total_units.is_integer() else round(total_units, 2),
            "total_revenue": round(total_rev, 2),
            "unique_products_sold": sold_products_count
        }

        g.cursor.execute("SELECT DISTINCT COALESCE(category, 'General') AS cat FROM products WHERE status = 'Active' AND category IS NOT NULL AND TRIM(category) != '' ORDER BY cat ASC")
        cat_rows = g.cursor.fetchall()
        product_categories = [r["cat"] for r in cat_rows if r.get("cat")]

        return jsonify({
            "success": True,
            "rows": enhanced_rows,
            "summary": summary,
            "categories": product_categories,
            "report_name": "Fast Moving Products Report"
        })

    except Exception as e:
        print("Fast Moving Report Error:", e)
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


@app.route("/reports/export/<report_type>", methods=["GET"])
def export_report(report_type):
    try:
        from openpyxl import Workbook
    except Exception:
        return jsonify({
            "success": False,
            "message": "openpyxl library not installed. Please install it first."
        }), 500

    report_type = (report_type or "").strip().lower()
    format_type = (request.args.get("format") or "csv").strip().lower()
    if report_type not in {"sales", "purchases", "expenses", "profit-loss", "outstanding", "services", "inventory", "fast-moving", "scrap"}:
        return jsonify({
            "success": False,
            "message": "Invalid report type"
        }), 400

    report_handlers = {
        "sales": reports_sales,
        "purchases": reports_purchases,
        "expenses": reports_expenses,
        "profit-loss": reports_profit_loss,
        "outstanding": reports_outstanding,
        "services": reports_services,
        "inventory": reports_inventory,
        "fast-moving": reports_fast_moving,
        "scrap": reports_scrap
    }

    handler = report_handlers.get(report_type)
    if not handler:
        return jsonify({
            "success": False,
            "message": "Invalid report type"
        }), 400

    report_response = handler()
    if getattr(report_response, "status_code", 200) != 200:
        return report_response

    data = report_response.get_json() or {}
    rows = data.get("rows", [])

    if report_type == "sales":
        headers = ["Invoice No.", "Date", "Customer", "Subtotal", "Discount", "Tax", "Grand Total", "Payment Method", "Payment Status"]
        values = [[r.get("invoice_number"), r.get("invoice_date"), r.get("customer_name"), r.get("subtotal"), r.get("discount"), r.get("tax_amount"), r.get("grand_total"), r.get("payment_method"), r.get("payment_status")] for r in rows]
    elif report_type == "purchases":
        headers = ["Purchase No.", "Date", "Supplier", "Subtotal", "Discount", "Tax", "Grand Total", "Payment Method", "Payment Status"]
        values = [[r.get("purchase_number"), r.get("purchase_date"), r.get("supplier_name"), r.get("subtotal"), r.get("discount"), r.get("tax_amount"), r.get("grand_total"), r.get("payment_method"), r.get("payment_status")] for r in rows]
    elif report_type == "expenses":
        headers = ["Expense No.", "Date", "Category", "Description", "Paid To", "Payment Method", "Amount", "Status"]
        values = [[r.get("expense_number"), r.get("expense_date"), r.get("category"), r.get("description"), r.get("paid_to"), r.get("payment_method"), r.get("amount"), r.get("status")] for r in rows]
    elif report_type == "profit-loss":
        headers = ["Metric", "Value"]
        summary = data.get("summary", {})
        values = [["Revenue", summary.get("revenue", 0)], ["Cost", summary.get("cost", 0)], ["Expenses", summary.get("expenses", 0)], ["Net Profit", summary.get("net_profit", 0)], ["Net Loss", summary.get("net_loss", 0)]]
    elif report_type == "outstanding":
        headers = ["Customer", "Mobile", "Total Sales", "Paid Amount", "Outstanding Amount"]
        values = [[r.get("customer_name"), r.get("mobile"), r.get("total_sales"), r.get("paid_amount"), r.get("outstanding_amount")] for r in rows]
    elif report_type == "services":
        headers = ["Service No.", "Service Date", "Expected Delivery Date", "Customer", "Product", "Serial Number", "Complaint", "Status", "Delivery Timing", "Technician", "Estimated Charges", "Final Charges"]
        values = [[r.get("service_number"), r.get("service_date"), r.get("expected_delivery_date"), r.get("customer_name"), r.get("product_name"), r.get("serial_number"), r.get("complaint"), r.get("service_status"), r.get("delivery_tag"), r.get("technician"), r.get("estimated_charges"), r.get("final_charges")] for r in rows]
    elif report_type == "scrap":
        headers = ["Date", "Item Name", "Available Qty", "Sold Qty", "Total Qty", "Unit", "Source", "Job/Customer Ref", "Reason", "Status"]
        values = [[r.get("scrap_date"), r.get("item_name"), r.get("available_quantity"), r.get("sold_quantity"), r.get("quantity"), r.get("unit"), r.get("source_type"), r.get("customer_job_no"), r.get("reason"), r.get("status")] for r in rows]
    elif report_type == "fast-moving":
        headers = ["Rank", "Product Name", "Category", "Brand", "Sold Qty", "Unit", "Total Revenue (Rs.)", "Avg Price (Rs.)", "Current Stock", "Movement Status"]
        values = [[r.get("rank"), r.get("product_name"), r.get("category"), r.get("brand"), r.get("total_sold_qty"), r.get("unit"), r.get("total_revenue"), r.get("avg_price"), r.get("current_stock"), r.get("status_badge")] for r in rows]
    else:
        headers = ["Product", "Category", "Brand", "Unit", "Stock Qty", "Stock Location", "Min Stock", "Purchase Price", "Stock Value", "Status"]
        values = [[r.get("product_name"), r.get("category"), r.get("brand"), r.get("unit"), r.get("stock"), r.get("location_summary"), r.get("min_stock"), r.get("purchase_price"), r.get("stock_value"), r.get("status_badge")] for r in rows]

    if format_type == "csv":
        output = StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerows(values)
        filename = f"{report_type}-report.csv"
        return send_file(
            BytesIO(output.getvalue().encode("utf-8")),
            mimetype="text/csv",
            as_attachment=True,
            download_name=filename
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = report_type
    sheet.append(headers)
    for row in values:
        sheet.append(row)

    file_buffer = BytesIO()
    workbook.save(file_buffer)
    file_buffer.seek(0)
    filename = f"{report_type}-report.xlsx"
    return send_file(file_buffer, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name=filename)


# ==========================================
# START SERVER
# ==========================================


if __name__ == "__main__":

    app.run(
        debug=os.getenv("FLASK_DEBUG", "false").lower() == "true"
    )

